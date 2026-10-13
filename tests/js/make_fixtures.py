"""e2e_dashboard.js 용 샘플 엑셀 3종(세로형 LIMS·가로 2줄 머리글+단위 줄·모르는 열 이름). 값은 임의 난수."""
import datetime as dt
import random
import sys
from pathlib import Path

import openpyxl

out = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/fx"); out.mkdir(parents=True, exist_ok=True)
random.seed(1)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Result"
ws.append(["품질시험 결과 조회"]); ws.append(["조회기간: 2026-10-01 ~ 2026-10-10"]); ws.append([])
ws.append(["시료채취일시", "시료명", "시험항목", "결과값", "단위", "판정"])
d0 = dt.datetime(2026, 10, 1, 8, 0)
for i in range(20):
    t = d0 + dt.timedelta(hours=8 * i)
    ws.append([t, "클링커", "f-CaO", round(random.uniform(.5, 1.6), 2), "%", "적합"])
    ws.append([t, "클링커", "리터중량", round(random.uniform(1250, 1400)), "g/L", "적합"])
    ws.append([t, "시멘트(OPC)", "분말도", round(random.uniform(3200, 3600)), "cm2/g", "적합"])
wb.save(out / "long_lims.xlsx")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "클링커"
ws.append(["클링커 성분 일보"]); ws.append(["", "", "화학성분", "", "", "", "모듈러스", ""])
ws.append(["분석일자", "시간", "CaO", "SiO2", "Al2O3", "Fe2O3", "LSF", "SM"]); ws.append(["", "", "%", "%", "%", "%", "", ""])
for i in range(15):
    ws.append([dt.datetime(2026, 10, 1 + i // 3), f"{8 + (i % 3) * 8:02d}:00", round(random.uniform(64, 66), 2), round(random.uniform(21, 22.5), 2),
               round(random.uniform(5, 5.6), 2), round(random.uniform(3, 3.5), 2), round(random.uniform(93, 97), 1), round(random.uniform(2.3, 2.7), 2)])
wb.save(out / "wide_two_row.xlsx")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
ws.append(["일자", "X101", "X102", "비고"])
for i in range(12):
    ws.append([dt.datetime(2026, 10, 1 + i), round(random.uniform(.4, 1.8), 2), round(random.uniform(1250, 1400)), "-"])
wb.save(out / "unknown_cols.xlsx")
print("만든 위치:", out)
