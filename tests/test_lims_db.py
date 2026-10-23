"""LIMS DB 간편 연결: 비밀번호 암호화 저장 · 열 역할/항목 자동 추천 · 표 고르기 → 소스별 동기화(긴 형식·가로형·출하시멘트) · 화면."""

import json
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from qms import lims, lims_db as L, secret_store as SS
from qms import shipment as S
from qms.store import load_raw

APP = str(Path(__file__).resolve().parent.parent / "streamlit_app.py")
PW = "Qm$#@!pw:1/한;x"


def test_secret_store_encrypts_and_never_keeps_plaintext(tmp_path):
    p = tmp_path / "secrets.enc.json"
    assert SS.status("x", p)["saved"] is False and SS.get_secret("x", p) == ""
    SS.set_secret("x", PW, p)
    raw = p.read_text(encoding="utf-8")
    assert PW not in raw and json.loads(raw)["x"]["m"] == SS.method()
    assert SS.get_secret("x", p) == PW
    st = SS.status("x", p)
    assert st["saved"] and st["usable"] and "v" not in st
    data = json.loads(raw)                       # 다른 PC·계정 값처럼 풀 수 없는 값 → 빈 값(오류 없이)
    data["x"]["v"] = "AAAA" + data["x"]["v"][4:]
    p.write_text(json.dumps(data), encoding="utf-8")
    assert SS.get_secret("x", p) == "" and SS.status("x", p) == {**SS.status("x", p), "saved": True, "usable": False}
    SS.delete_secret("x", p)
    assert not SS.status("x", p)["saved"]
    with pytest.raises(ValueError):
        SS.set_secret("x", "", p)


def test_name_rules_suggest_roles_items_and_formats():
    cols = ["SMP_DATE", "SMP_TIME", "SMP_PT", "ITEM_NM", "RSLT", "PROD", "STAT", "APPR_DT"]
    assert L.suggest_roles(cols, "long") == {"sampled_at": "SMP_DATE", "time": "SMP_TIME", "sample_point": "SMP_PT",
                                             "test_code": "ITEM_NM", "value": "RSLT", "product": "PROD", "status": "STAT",
                                             "updated_at": "APPR_DT"}
    assert L.guess_layout(cols) == "long"
    assert L.guess_layout(["채취일자", "CaO(%)", "C3S", "f-CaO"]) == "wide"
    assert L.guess_layout(["채취일", "첨가재명칭", "비고", "압축강도28일(MPa)", "Blaine", "SO3(XRF)"]) == "ship"
    assert [L.suggest_item(t, n) for t, n in [("clinker", "CaO(%)"), ("clinker", "f-CaO"), ("xrd", "C3S"), ("cement", "분말도"),
                                              ("physical", "압축강도28일"), ("physical", "초결"), ("cement", "SO3(XRF)"), ("xrd", "f-CaO")]] == \
        ["clk_cao", "clk_fcao", "xrd_alite", "cem_blaine", "phy_s28", "phy_ist", "cem_so3", "xrd_fcao"]
    assert L.suggest_item("clinker", "알수없는항목") is None
    assert L.guess_date_format(["20261015", "20261014"]) == "%Y%m%d"
    assert L.guess_date_format(["2026-10-15 08:30:00"]) == "%Y-%m-%d %H:%M:%S"
    assert L.guess_date_format([20261015, 20261014]) == "int:%Y%m%d"
    assert L.guess_date_format([pd.Timestamp("2026-10-15")]) == "datetime"
    assert L.guess_type(1433) == "mssql" and L.guess_type(1521) == "oracle" and L.guess_type(9999) is None
    from qms.ui_lims import guess_table
    assert [guess_table(x) for x in ["KILN_FEED", "CLINKER_DAILY", "CEMENT_LOT", "클링커", "CM"]] == \
        ["raw_meal", "xrd", "physical", "clinker", "cement"]


def test_urls_hide_password_and_error_messages_are_sanitized():
    base = {"ip": "10.0.0.5", "host": "LIMS-DB", "port": 1433, "database": "LIMS", "user": "qms_reader"}
    u = L.build_url({**base, "type": "mssql", "driver": "pymssql"}, PW)
    assert u.drivername == "mssql+pymssql" and u.host == "10.0.0.5" and u.password == PW and PW not in str(u)
    assert f"DRIVER={{{L.ODBC_DEFAULT}}}" in L.build_url({**base, "type": "mssql", "driver": "pyodbc"}, PW).query["odbc_connect"]
    o = L.build_url({**base, "type": "oracle", "driver": "oracledb", "port": 1521, "database": "LIMSPDB"}, PW)
    assert o.drivername == "oracle+oracledb" and o.query["service_name"] == "LIMSPDB"
    assert L.build_url({**base, "type": "oracle", "oracle_mode": "sid", "database": "ORCL"}, PW).database == "ORCL"
    assert L.build_url({**base, "type": "postgresql", "database": ""}, PW).database == "postgres"
    assert L.build_url({**base, "type": "mysql", "host": "", "ip": ""}, PW).query["charset"] == "utf8mb4"
    assert L.server_of({"ip": "", "host": "LIMS-DB"}) == "LIMS-DB" and L.server_of({"ip": "1.2.3.4", "host": "x"}) == "1.2.3.4"
    msg = L.sanitize(f"login failed for mssql+pymssql://u:{PW}@h/db password={PW}", PW)
    assert PW not in msg and "****" in msg
    assert "DB 이름" in L.explain("(4060, b'Cannot open database \"NOPE\" requested by the login. The login failed.')")
    assert "비밀번호" in L.explain("(18456, b'Login failed for user')")
    assert "방화벽" in L.explain("timed out") or "연결하지 못했습니다" in L.explain("timed out")
    ok, m = L.tcp_check("127.0.0.1", 1, 1)
    assert not ok and "포트" in m


@pytest.fixture()
def lims_sqlite(tmp_path):
    db = tmp_path / "lims.db"
    now = datetime.now().replace(microsecond=0)
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE QC_RESULT (SMP_DT TEXT, SMP_PT TEXT, ITEM_CD TEXT, RSLT_VAL TEXT, PROD_CD TEXT, STAT TEXT, APPR_DT TEXT)")
    rows = []
    for i in range(12):
        t = now - timedelta(days=i, hours=2)
        a = (t + timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S")
        ts = t.strftime("%Y-%m-%d %H:%M:%S")
        rows += [(ts, "CLK", "CAO", f"{65 + i * 0.01:.2f}", None, "승인", a), (ts, "CLK", "FCAO", "1.2", None, "승인", a),
                 (ts, "CM", "BLAINE", f"{3400 + i}", "OPC", "승인", a), (ts, "CM", "BLAINE", "9999", "OPC", "대기", a),
                 (ts, "CM", "WATER", "26", "OPC", "승인", a)]
    con.executemany("INSERT INTO QC_RESULT VALUES (?,?,?,?,?,?,?)", rows)
    con.execute('CREATE TABLE CLK_DAILY ("채취일자" TEXT, "채취시각" TEXT, "C3S" REAL, "f-CaO" REAL, "비고" TEXT)')
    con.executemany("INSERT INTO CLK_DAILY VALUES (?,?,?,?,?)",
                    [((now - timedelta(days=i)).strftime("%Y%m%d"), "0830", 60 + i * 0.1, 1.1, "") for i in range(10)])
    con.execute('CREATE TABLE SHIP ("채취일" TEXT, "첨가재명칭" TEXT, "비고" TEXT, "석회석비율" REAL, "첨가재비율" REAL, '
                '"압축강도28일(MPa)" REAL, "Blaine" REAL)')
    con.executemany("INSERT INTO SHIP VALUES (?,?,?,?,?,?,?)",
                    [((now - timedelta(days=i)).strftime("%Y-%m-%d"), "슬래그", "리나호, 수성6호" if i % 2 else "삼표1호(저염)",
                      5.0, 3.0, 48.0 - i * 0.1, 3450 + i) for i in range(8)])
    con.commit()
    con.close()
    return {"type": "sqlite", "database": str(db)}


def _sources(conn):
    cols = list(L.get_columns(conn, "", None, "QC_RESULT")["열"])
    long = {"name": "결과", "table": "QC_RESULT", "layout": "long", "enabled": True, "cols": L.suggest_roles(cols, "long"),
            "date_format": "%Y-%m-%d %H:%M:%S", "updated_format": "%Y-%m-%d %H:%M:%S",
            "mapping": [{"sample_point": "CLK", "test_code": "CAO", "table": "clinker", "column": "clk_cao", "factor": 1, "offset": 0},
                        {"sample_point": "CM", "test_code": "BLAINE", "table": "cement", "column": "cem_blaine", "factor": 1, "offset": 0}]}
    wide = {"name": "클링커 일일", "table": "CLK_DAILY", "layout": "wide", "enabled": True,
            "cols": {"sampled_at": "채취일자", "time": "채취시각"}, "date_format": "%Y%m%d", "fixed_sample_point": "CLINKER_DAILY",
            "item_cols": ["C3S"], "mapping": [{"sample_point": "CLINKER_DAILY", "test_code": "C3S", "table": "xrd",
                                              "column": "xrd_alite", "factor": 1, "offset": 0}]}
    scols = list(L.get_columns(conn, "", None, "SHIP")["열"])
    ship = {"name": "출하", "table": "SHIP", "layout": "ship", "enabled": True, "date_format": "%Y-%m-%d %H:%M:%S",
            "ship_map": {k: scols[j] for k, j in S.match_headers(scols).items()}}
    return long, wide, ship


def test_browse_and_sync_three_layouts(lims_sqlite, tmp_path):
    conn = lims_sqlite
    assert [x["ok"] for x in L.diagnose(conn, "")] == [True] * 4
    assert set(L.list_objects(conn, "", None)["이름"]) == {"QC_RESULT", "CLK_DAILY", "SHIP"}
    assert len(L.preview(conn, "", None, "QC_RESULT", 5)) == 5
    long, wide, ship = _sources(conn)
    codes = L.distinct_codes(conn, "", long, datetime.now() - timedelta(days=35))
    assert {("CLK", "CAO"), ("CM", "BLAINE"), ("CM", "WATER")} <= set(zip(codes["sample_point"], codes["test_code"]))
    d, capped = L.fetch_source(conn, "", wide, datetime.now() - timedelta(days=5))
    assert not capped and list(d.columns) == ["sampled_at", "C3S", "sample_point"] and d["sampled_at"].dt.hour.eq(8).all()
    assert len(d) == 6                                                   # 문자 날짜(YYYYMMDD)로 최근 5일만

    S.clear_shipment()
    cfg = lims.LimsConfig(enabled=True, mode="db", db=conn, sources=[long, wide, ship])
    qdb = tmp_path / "qms.db"
    dry = lims.sync(cfg, db_path=qdb, dry_run=True)
    assert not dry.errors and not dry.saved and dry.rows_by_table == {"clinker": 12, "cement": 12, "xrd": 10, "shipment": 8}
    assert dry.n_status_skipped == 12 and dry.n_unmapped == 24          # '대기' 결과 제외, WATER·FCAO 는 매칭 안 함
    rep = lims.sync(cfg, db_path=qdb)
    assert rep.saved and not rep.errors
    raw = load_raw(qdb)
    assert raw["cement"]["cem_blaine"].max() < 9999 and set(raw["cement"]["product"]) == {"1종"}
    assert len(raw["xrd"]) == 10 and len(S.load_shipment()) == 8
    sh = S.load_shipment()
    assert set(sh["product_kind"]) == {"일반", "저염"} and sh["n_ship"].max() == 2
    assert cfg.sources[0]["watermark"] and "watermark" not in cfg.sources[1]   # 확정일시 열이 있는 소스만 워터마크
    again = lims.sync(cfg, db_path=qdb, dry_run=True)
    assert again.rows_by_table.get("clinker", 0) == 0                     # 워터마크 이후 새 결과 없음
    bad = lims.LimsConfig(enabled=True, mode="db", db=conn, sources=[{**long, "name": "틀린 표", "table": "NOPE"}, wide])
    r2 = lims.sync(bad, db_path=qdb, dry_run=True)
    assert len(r2.errors) == 1 and "틀린 표" in r2.errors[0] and r2.rows_by_table.get("xrd") == 10   # 한 소스 실패해도 나머지 진행
    S.clear_shipment()


def test_password_is_never_written_to_config_or_log(lims_sqlite, monkeypatch, tmp_path):
    monkeypatch.setattr(SS, "STORE_PATH", tmp_path / "s.json")
    SS.set_secret(lims.STORE_NAMES["password"], PW)
    assert lims._secret("password") == PW and lims.secret_source("password").startswith("화면 입력")
    cfg = lims.LimsConfig(enabled=True, mode="db", db={**lims_sqlite, "password": PW}, sources=list(_sources(lims_sqlite)))
    lims.save_lims_config(cfg, tmp_path / "lims.json")
    t = (tmp_path / "lims.json").read_text(encoding="utf-8")
    assert PW not in t and '"password"' not in t
    rep = lims.sync(cfg, db_path=tmp_path / "q.db", dry_run=True, log_path=tmp_path / "log.jsonl")
    assert PW not in (tmp_path / "log.jsonl").read_text(encoding="utf-8")
    monkeypatch.setenv("QMS_LIMS_PASSWORD", "env-pw")
    assert lims._secret("password") == "env-pw" and lims.secret_source("password") == "환경변수 QMS_LIMS_PASSWORD"
    assert rep.sources
def test_partial_items_from_lims_do_not_break_derived_values():
    """LIMS에서 일부 항목(예: 클링커 CaO만)만 들어와도 LSF·SM 등 파생값 계산과 화면이 멈추지 않는다."""
    from qms.store import build_store
    raw = {"clinker": pd.DataFrame({"timestamp": pd.to_datetime(["2026-10-01 08:00"]), "clk_cao": [65.2]}),
           "raw_meal": pd.DataFrame({"timestamp": pd.to_datetime(["2026-10-01 08:00"]), "rm_cao": [43.1]})}
    st = build_store(raw)
    assert st.tables["clinker"]["clk_cao"].iloc[0] == 65.2 and pd.isna(st.tables["clinker"]["clk_lsf"].iloc[0])


def test_detect_db_type_from_first_handshake():
    """아이디·비밀번호 없이 첫 인사(핸드셰이크)로 DB 종류 판별 — 실제 서버에서 받은 응답 바이트로 확인."""
    tds = bytes.fromhex("0401001a0000010000000b00060100110001ff100010d1000002")          # SQL Server 2022 PRELOGIN 응답
    assert L._classify(tds) == {"type": "mssql", "version": "SQL Server 2022 (16.0.4305)",
                                "encryption": "강제 아님(암호화 없이도 접속 허용)"}
    maria = b"Z\x00\x00\x00\n11.8.9-MariaDB-ubu2404\x00\x05\x00\x00\x00tz)[l0FO\x00"
    assert L._classify(maria) == {"type": "mysql", "version": "MariaDB 11.8.9"}
    assert L._classify(b"N")["type"] == "postgresql" and L._classify(b"S")["type"] == "postgresql"
    ora = (b'\x00]\x00\x00\x04\x00\x00\x00"\x00\x00Q(DESCRIPTION=(TMP=)(VSNNUM=0)(ERR=1189)'
           b'(ERROR_STACK=(ERROR=(CODE=1189)(EMFI=4))))')
    assert L._classify(ora) == {"type": "oracle", "version": ""}
    assert L._classify(b"") is None and L._classify(b"HTTP/1.1 400 Bad Request\r\n") is None
    assert len(L._tds_prelogin()) == 26 and len(L._tns_version()) == int.from_bytes(L._tns_version()[:2], "big")

    import socket
    import threading
    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(5)
    port = srv.getsockname()[1]

    def serve():                                   # 가짜 MySQL: 접속하면 먼저 인사
        for _ in range(2):
            try:
                c, _ = srv.accept()
            except OSError:
                return
            c.sendall(maria)
            c.close()
    threading.Thread(target=serve, daemon=True).start()
    det = L.detect_type("127.0.0.1", port, timeout=2)
    srv.close()
    assert det["type"] == "mysql" and "MariaDB 11.8.9" in det["msg"]
    closed = L.detect_type("127.0.0.1", 1, timeout=1)
    assert closed["type"] is None and "포트" in closed["msg"]


def test_explore_ranks_result_tables_and_finds_terms(lims_sqlite, tmp_path):
    """DB 구조를 몰라도: 표·열 목록을 훑어 결과 표를 위로(잡표는 0점), 이름·값 찾기, 구조 엑셀."""
    db = Path(lims_sqlite["database"])
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE USER_INFO (USER_ID TEXT, USER_NM TEXT, REG_DT TEXT)")
    con.execute("CREATE TABLE SYS_LOG (LOG_DT TEXT, MSG TEXT, USER_ID TEXT)")
    con.commit()
    con.close()
    res = L.explore(lims_sqlite, "")
    cand = res["candidates"]
    order = list(cand["table"])
    assert set(order[:3]) == {"QC_RESULT", "CLK_DAILY", "SHIP"} and not res["errors"]
    lay = dict(zip(cand["table"], cand["layout"]))
    assert lay["QC_RESULT"] == "long" and lay["CLK_DAILY"] == "wide" and lay["SHIP"] == "ship"
    sc = dict(zip(cand["table"], cand["score"]))
    assert sc["USER_INFO"] == 0 and sc["SYS_LOG"] == 0
    hit = L.find_term(lims_sqlite, "", res["catalog"], cand, "BLAINE")
    assert ("값", "QC_RESULT", "ITEM_CD") in set(zip(hit["찾은 곳"], hit["표"], hit["열"]))
    assert ("열 이름", "SHIP", "Blaine") in set(zip(hit["찾은 곳"], hit["표"], hit["열"]))
    assert L.find_term(lims_sqlite, "", res["catalog"], cand, "  ").empty
    xls = L.structure_workbook(res, {"QC_RESULT": L.preview(lims_sqlite, "", None, "QC_RESULT", 3)})
    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO(xls))
    assert wb.sheetnames == ["① 추천 표", "② 표 목록", "③ 열 목록", "④ 예시 데이터"]
    assert L.score_table("X", [("a", "text"), ("b", "text")])["score"] == 0
    assert L.score_table("CLK", [("측정일자", "varchar"), ("CaO", "float"), ("SiO2", "float")])["layout"] == "wide"


def test_demo_lims_view_is_found_first(demo_raw, tmp_path):
    p = tmp_path / "demo.db"
    lims.create_demo_lims(demo_raw, p)
    res = L.explore({"type": "sqlite", "database": str(p)}, "")
    top = res["candidates"].iloc[0]
    assert top["table"] == "V_QC_RESULT" and top["layout"] == "long"


def test_missing_driver_is_explained_and_detected():
    """드라이버가 없는 실행 환경(PC 파이썬으로 코드 실행 등) → 정확한 안내(드라이버 바꾸라는 엉뚱한 안내 금지)·설치 버튼용 상태."""
    msg = L.explain("ModuleNotFoundError: No module named 'pymssql'")
    assert "'pymssql'" in msg and "드라이버 설치" in msg and "바꿔" not in msg
    assert "드라이버 이름" in L.explain("('IM002', '[IM002] Data source name not found')")
    assert L.driver_ok("sqlite") and not L.driver_ok("no_such_driver_mod")
    assert L.DRIVER_PIP["psycopg2"].startswith("psycopg2-binary")
    assert L.b64_ip("MTcyLjE3LjIyLjEx") == "172.17.22.11" and L.b64_ip("LIMS-DB01") is None and L.b64_ip("") is None
    req = (Path(__file__).resolve().parent.parent / "requirements-lims.txt").read_text(encoding="utf-8")
    assert all(p in req for p in ("pymssql", "oracledb", "psycopg2-binary", "PyMySQL"))


def test_connection_plans_and_tls_hint():
    """로그인 전 TLS 단계 끊김(20002) → 정확한 안내, 연결 방법 후보(암호화 없는 방법은 동의할 때만), pymssql 설정 전달."""
    m = "OperationalError: (20002, b'DB-Lib error message 20002, severity 9:\\nTDS server connection failed (10.50.11.71)')"
    assert "암호화(TLS)" in L.explain(m) and "비밀번호 문제 아님" in L.explain(m)
    base = {"type": "mssql", "ip": "10.0.0.1", "port": 1433, "user": "u", "driver": "pymssql"}
    labels = [p[0] for p in L.connection_plans(base)]
    assert labels[:3] == ["pymssql(기본)", "QMS 보안 중계 · TLS 1.2 고정", "QMS 보안 중계 · 오래된 TLS·인증서 허용"]
    assert "pymssql · TDS 7.4" in labels and not any("암호화 없이" in x for x in labels)
    plain = dict(L.connection_plans(base, allow_plain=True))
    assert plain["pymssql · 암호화 없이(사내망 전용)"]["encryption"] == "off"
    assert L.connect_args({**base, "tds_version": "7.0", "encryption": "off"})["tds_version"] == "7.0"
    assert "tds_version" not in L.connect_args(base)
    cs = L.build_url({**base, "driver": "pyodbc", "odbc_driver": "ODBC Driver 18 for SQL Server"}, "pw").query["odbc_connect"]
    assert "Encrypt=yes" in cs and "TrustServerCertificate=yes" in cs
    assert "Encrypt" not in L.build_url({**base, "driver": "pyodbc", "odbc_driver": "SQL Server"}, "pw").query["odbc_connect"]
    assert L.connection_plans({"type": "postgresql"}) == [("기본", {})]
    tds = bytes.fromhex("0401001a0000010000000b00060100110001ff100010d1000003")
    assert L._classify(tds)["encryption"].startswith("필수")


def test_odbc_without_database_is_not_treated_as_dsn():
    """DB 이름을 비우면 SQLAlchemy 가 IP 를 ODBC 'DSN 이름'으로 써서(IM002) 서버에 가지도 못하던 문제."""
    from sqlalchemy.dialects.mssql import pyodbc as P

    c = {"type": "mssql", "driver": "pyodbc", "odbc_driver": "SQL Server", "ip": "10.50.11.71", "port": 3718, "user": "sampyo",
         "database": ""}
    (cs,), _ = P.dialect().create_connect_args(L.build_url(c, PW))
    assert "dsn=" not in cs.lower() and "DRIVER={SQL Server}" in cs and "SERVER=10.50.11.71,3718" in cs and "Network=DBMSSOCN" in cs
    assert "PWD={" + PW.replace("}", "}}") + "}" in cs and "DATABASE" not in cs
    cs17 = L.odbc_string({**c, "odbc_driver": "ODBC Driver 17 for SQL Server", "database": "TEST"}, "a}b")
    assert "SERVER=tcp:10.50.11.71,3718" in cs17 and "DATABASE={TEST}" in cs17 and "PWD={a}}b}" in cs17
    assert "a}}b" not in L.sanitize("x PWD={a}}b} y", "a}b")
    assert "비밀번호" in L.explain("(18456, b\"Login failed for user 'sa'. 20002 TDS server connection failed\")")
    assert L.plan_label({"type": "mssql", "tls_relay": "tls12"}).startswith("QMS 보안 중계(TLS 1.2")
    assert L.plan_label({"type": "mssql", "driver": "pyodbc", "odbc_driver": "SQL Server"}) == "Windows ODBC · SQL Server"


def _fake_tds_tls_server(ctx, enc=0x03):
    """PRELOGIN 에 '암호화 필수'로 답하고 TDS 방식 TLS 를 맺은 뒤, 받은 내용을 'ECHO:' 를 붙여 돌려주는 가짜 SQL Server."""
    import socket
    import ssl
    import threading

    from qms import tds_relay as R

    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(4)

    def serve():
        while True:
            try:
                c, _ = srv.accept()
            except OSError:
                return
            try:
                R.read_message(c)
                c.sendall(R.packets(R.RESPONSE, R.set_enc(R.prelogin_body(0x02), enc)))
                inc, out = ssl.MemoryBIO(), ssl.MemoryBIO()
                t = ctx.wrap_bio(inc, out, server_side=True)
                while True:
                    try:
                        t.do_handshake()
                        break
                    except ssl.SSLWantReadError:
                        d = out.read()
                        if d:
                            c.sendall(R.packets(R.PRELOGIN, d))
                        inc.write(R.read_message(c)[1])
                d = out.read()
                if d:
                    c.sendall(R.packets(R.PRELOGIN, d))
                while True:
                    try:
                        msg = t.read(4096)
                        break
                    except ssl.SSLWantReadError:
                        d = c.recv(65536)
                        if not d:                             # 핸드셰이크만 하고 끊는 진단 접속
                            raise EOFError from None
                        inc.write(d)
                t.write(b"ECHO:" + msg)
                c.sendall(out.read())
            except Exception:  # noqa: BLE001
                pass
            finally:
                c.close()

    threading.Thread(target=serve, daemon=True).start()
    return srv


def _legacy_server_ctx(tmp_path, bits, sha, tls_max):
    import shutil
    import ssl
    import subprocess
    import warnings

    if not shutil.which("openssl"):
        pytest.skip("openssl 명령 없음")
    key, crt = tmp_path / "k.pem", tmp_path / "c.pem"
    subprocess.run(["openssl", "req", "-x509", "-nodes", "-newkey", f"rsa:{bits}", f"-{sha}", "-subj", "/CN=lims", "-keyout", str(key),
                    "-out", str(crt), "-days", "30"], check=True, capture_output=True)
    ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", DeprecationWarning)
        ctx.minimum_version = ssl.TLSVersion.TLSv1
        ctx.maximum_version = tls_max
    ctx.set_ciphers("ALL:@SECLEVEL=0")
    try:
        ctx.load_cert_chain(str(crt), str(key))
    except ssl.SSLError:
        pytest.skip("이 OpenSSL 은 오래된 인증서를 서버 쪽에서도 거부")
    return ctx


def _through_relay(port, profile):
    import socket

    from qms import tds_relay as R

    r = R.Relay("127.0.0.1", port, profile)
    try:
        with socket.create_connection(("127.0.0.1", r.port), timeout=10) as s:
            s.sendall(R.packets(R.PRELOGIN, R.prelogin_body(0x00)))
            _, body = R.read_message(s)
            assert R.get_enc(body) == R.ENC_NOT_SUP          # pymssql 쪽에는 '암호화 지원 안 함'으로 보임
            s.sendall(b"login7-packet")
            data = b""
            while not data.endswith(b"packet"):
                try:
                    d = s.recv(4096)
                except ConnectionResetError:                  # 중계가 TLS 실패 후 연결을 닫음
                    d = b""
                if not d:
                    break
                data += d
            return data, r.last_tls, r.last_error
    finally:
        r.close()


def test_tls_relay_old_server_and_tls_version_pinning(tmp_path):
    """QMS 보안 중계: 오래된 서버(TLS 1.0·1024비트 SHA1 인증서)는 'legacy' 로만, 일반 서버는 TLS 1.2 고정으로 이어 준다."""
    import ssl

    from qms import tds_relay as R

    old = _fake_tds_tls_server(_legacy_server_ctx(tmp_path, 1024, "sha1", ssl.TLSVersion.TLSv1))
    data, tls, err = _through_relay(old.getsockname()[1], "legacy")
    assert data == b"ECHO:login7-packet" and tls.startswith("TLSv1") and not err
    data, _, err = _through_relay(old.getsockname()[1], "tls12")
    assert data == b"" and "TLS 연결 실패" in err                 # TLS 1.2 고정은 TLS 1.0 서버와 맞지 않음 → 이유를 남김
    old.close()
    new = _fake_tds_tls_server(_legacy_server_ctx(tmp_path, 2048, "sha256", ssl.TLSVersion.TLSv1_3))
    data, tls, _ = _through_relay(new.getsockname()[1], "tls12")
    assert data == b"ECHO:login7-packet" and tls.startswith("TLSv1.2")   # 서버가 TLS 1.3 을 허용해도 1.2 로 맺음
    pr = R.probe_tls("127.0.0.1", new.getsockname()[1])
    assert dict((a, b) for a, b, _ in pr["versions"])["TLS 1.2"] and pr["cert"]["키"] == "RSA 2048비트"
    assert pr["cert"]["서명 방식"] == "SHA256-RSA" and pr["encryption"].startswith("필수")
    new.close()
    m = L.relay_conn({"type": "mssql", "ip": "10.9.9.9", "port": 1433, "tls_relay": "tls12", "driver": "pymssql"})
    assert m["ip"] == "127.0.0.1" and m["encryption"] == "off" and m["port"] != 1433
    assert L.relay_conn({"type": "mssql", "ip": "10.9.9.9"})["ip"] == "10.9.9.9"
