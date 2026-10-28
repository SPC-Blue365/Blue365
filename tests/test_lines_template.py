"""공정구분(호기·밀) 저장·판정과 통합 업로드 양식 시험."""

import io

import numpy as np
import openpyxl
import pandas as pd

from qms import workbook
from qms.alerts import detect_events
from qms.standards import DEFAULT_SETTINGS, SpecRegistry
from qms.store import DataStore, enrich, normalize_line, parse_upload, save_raw, load_raw
from qms.upload_template import SHEETS, build_upload_template


def test_normalize_line_codes():
    assert [normalize_line(v) for v in ("4K", "4 k", "4호기", "#5", "6", "6RM", "6 R/M", "2CM", "3C/M", None, 0, "")] == \
        ["4K", "4K", "4K", "5K", "6K", "6R/M", "6R/M", "2C/M", "3C/M", "", "", ""]


def test_same_time_different_kilns_are_kept_apart(tmp_path):
    """같은 채취 시각의 6K·7K 클링커가 서로 덮어쓰지 않는다(이전: 시각만으로 병합돼 절반 유실)."""
    db = tmp_path / "q.db"
    t = pd.to_datetime(["2026-10-01 00:00", "2026-10-01 00:00", "2026-10-01 04:00"])
    legacy = pd.DataFrame({"timestamp": t[:1], "clk_cao": [64.0]})                  # 예전 구분 없는 행
    save_raw({"clinker": legacy}, path=db)
    new = pd.DataFrame({"timestamp": t, "line": ["6K", "7K", "6K"], "clk_cao": [64.07, 64.05, 64.06]})
    save_raw({"clinker": new}, path=db)
    df = load_raw(db)["clinker"]
    assert len(df) == 3 and sorted(df["line"]) == ["6K", "6K", "7K"]             # 같은 시각의 예전 행은 정리
    st = DataStore(enrich(load_raw(db)))
    reg = SpecRegistry()
    assert st.lines("clinker") == ["6K", "7K"]
    assert len(st.series("clk_cao", reg, line="6K")) == 2 and len(st.series("clk_cao", reg, line="7K")) == 1


def test_events_are_detected_per_line():
    """호기별로 따로 판정: 7K만 f-CaO가 높으면 7K 이벤트만 생긴다."""
    idx = pd.date_range("2026-09-01", periods=240, freq="2h")
    rng = np.random.default_rng(1)
    base = pd.DataFrame({"timestamp": idx, "clk_cao": 64.5, "clk_sio2": 21.2, "clk_al2o3": 5.3, "clk_fe2o3": 3.2,
                         "clk_fcao": rng.normal(1.0, 0.1, len(idx))})
    k6 = base.assign(line="6K")
    k7 = base.assign(line="7K")
    k7.loc[k7.index[-12:], "clk_fcao"] = 3.5
    st = DataStore(enrich({"clinker": pd.concat([k6, k7], ignore_index=True)}))
    reg = SpecRegistry()
    ev = [e for e in detect_events(st, reg, dict(DEFAULT_SETTINGS)) if e.item_key == "clk_fcao" and e.rule in ("KS", "SPEC")]
    assert ev and {e.line for e in ev} == {"7K"} and all("7K" in e.item_name for e in ev)


def _fill_template() -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(build_upload_template()))
    put = {
        "원료": [["2026-10-01", "02:00", "6K", "1/2종조합원료", 42.6, 12.3, 3.6, 2.2, 2.6, 0.4, 0.7, 0.09, 35.4],
                ["2026-10-01", "02:00", "7K", "1/2종조합원료", 42.7, 12.2, 3.6, 2.2, 2.5, 0.4, 0.7, 0.09, 35.5],
                ["2026-10-01", "00:00", "6R/M", "1/2종조합원료", 42.3, 12.5, 3.6, 2.3, 2.6, 0.3, 0.7, 0.09, 35.6]],
        "원료 잔사": [["2026-10-01", "00:00", "6R/M", "1/2종조합원료", 13.2, 1.7]],          # 원료 시트 6R/M 행과 합쳐짐
        "클링커(6K,7K)": [["2026-10-01", "00:00", "6K", "내수", "6-4호", 64.07, 20.99, 5.49, 3.07, 3.82, 0.49, 1.02, 0.22, 0.56, 214, 1466, 135],
                         ["2026-10-01", "00:00", "7K", "수출클링커", "6-3호", 64.05, 21.01, 5.5, 2.95, 3.84, 0.57, 1.17, 0.23, 1.04]],
        "클링커(4K,5K)": [["2026-10-01", "01:00", "5K", "내수", "5호", 64.2, 21.3, 5.3, 3.2, 4.8, 0.3, 0.9, 0.2, 0.73]],
        "시멘트": [["2026-10-01", "00:00", "2C/M", "1종시멘트 벌크", "구내9호", 7.2, 3720, 2.48]],
    }
    for sn, rows in put.items():
        ws = wb[sn]
        for i, r in enumerate(rows, 5):
            for j, v in enumerate(r, 1):
                ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_upload_template_structure():
    wb = openpyxl.load_workbook(io.BytesIO(build_upload_template()))
    names = wb.sheetnames
    assert names[:2] == ["사용법", "공정구분"] and names[-1] == "항목사전"
    assert [s.name for s in SHEETS] == names[2:-1]
    ws = wb["원료"]
    assert ws["C3"].value == "공정구분" and ws.data_validations.dataValidation     # 목록 선택
    assert any("4K" in str(c.value) for row in wb["공정구분"].iter_rows() for c in row)


def test_empty_template_reads_nothing_and_no_errors():
    parts = workbook.analyze(build_upload_template(), "양식.xlsx")
    assert parts.kinds == [] and parts.errors == [] and parts.is_template


def test_filled_template_roundtrip_keeps_lines():
    parts = workbook.analyze(_fill_template(), "양식.xlsx")
    assert parts.kinds == ["공정 데이터"] and not parts.errors and not parts.ships            # 공정 시트를 출하시멘트로 오인하지 않음
    t = parts.process.tables
    raw = t["raw_meal"]
    assert sorted(raw["line"]) == ["6K", "6R/M", "6R/M", "7K"] and set(raw["grade"]) == {"1/2종조합원료"}   # 원료 + 원료 잔사
    assert raw.loc[raw["rm_r90"].notna(), "line"].tolist() == ["6R/M"]
    clk = t["clinker"]                                                                          # 클링커 시트 2개 합침
    assert sorted(clk["line"]) == ["5K", "6K", "7K"] and set(clk["silo"]) == {"6-4호", "6-3호", "5호"}
    assert clk.loc[clk["line"] == "6K", ["clk_cl", "clk_f", "clk_temp"]].iloc[0].tolist() == [214, 1466, 135]
    cem = t["cement"]
    assert cem["line"].tolist() == ["2C/M"] and cem["product"].tolist() == ["1종"] and cem["silo"].tolist() == ["구내9호"]
    assert cem["cem_r45"].tolist() == [7.2] and cem["cem_blaine"].tolist() == [3720]             # LIMS 순서: 45㎛ → Blaine
    assert raw["timestamp"].dt.hour.tolist() == [2, 2, 0, 0]                                    # 채취일+채취시간


def test_swapped_blaine_r45_columns_are_fixed():
    """Blaine 열에 45㎛ 값, 45㎛ 열에 Blaine 값이 통째로 들어오면 바로잡고 알린다."""
    from qms.store import _fix_swapped_fineness
    df = pd.DataFrame({"cem_blaine": [5.6, 6.2, 5.9, 0.9, 6.0], "cem_r45": [3560, 3560, 3600, 4770, 3700]})
    out, swapped = _fix_swapped_fineness(df)
    assert swapped and out["cem_blaine"].tolist() == [3560, 3560, 3600, 4770, 3700]
    assert not _fix_swapped_fineness(out)[1]


def test_sheet_names_with_suffix_map_to_tables():
    from qms.store import template_table
    assert [template_table(n) for n in ("원료 잔사", "클링커(6K,7K)", "클링커일일", "시멘트", "출하시멘트(1종)", "인수검사(원료)")] == \
        ["raw_meal", "clinker", "xrd", "cement", None, None]


def test_lims_raw_header_names_map_to_line():
    """LIMS 원본 머리글: 크링커공정·시멘트밀공정·공정구분 → 공정구분(line)."""
    rows = [["크링커성적"], ["검사일자"], ["로트번호", "검사일", "채취일", "품명", "채취시간", "크링커공정", "CaO", "SiO2", "F-CaO(XRF)"],
            [1, "2026-10-01", "2026-10-01", "1/2종크링커", "00:00", "6K", 64.0, 21.0, 0.5],
            [2, "2026-10-01", "2026-10-01", "1/2종크링커", "00:00", "7K", 64.1, 21.1, 1.0]]
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    res = parse_upload(buf.getvalue(), "크링커.xlsx")
    df = res.tables["clinker"]
    assert df["line"].tolist() == ["6K", "7K"] and df["grade"].tolist() == ["1/2종크링커"] * 2


def test_every_process_template_header_is_recognized():
    """양식의 공정 시트 머리글이 전부 QMS 항목으로 연결된다(무시되는 열 없음)."""
    from qms.store import RAW_COLUMNS, SHEET_NAMES, apply_column_aliases

    by_sheet = {v: k for k, v in SHEET_NAMES.items()}
    for spec in SHEETS:
        if spec.name not in by_sheet:
            continue
        tbl = by_sheet[spec.name]
        df = apply_column_aliases(pd.DataFrame(columns=[c.header for c in spec.cols]), tbl)
        valid = set(RAW_COLUMNS[tbl]) | {"timestamp"}
        left = [c for c in df.columns if c not in valid and c != "채취시간"]
        assert not left, (spec.name, left)


def test_future_year_typo_is_corrected():
    """2029-09-25 같은 연도 오타는 올해로, 검사일이 있으면 검사일 연도로 고쳐 읽는다(고쳐도 미래면 제외)."""
    from qms.store import fix_future_dates
    now = pd.Timestamp.now().normalize()
    past = now - pd.Timedelta(days=30)
    df = pd.DataFrame({"d": [past.replace(year=now.year + 3), past, now.replace(year=now.year + 1) + pd.Timedelta(days=10)],
                       "ref": [pd.NaT, pd.NaT, pd.NaT]})
    out, fixed, dropped = fix_future_dates(df, "d")
    assert out["d"].iloc[0] == past and len(out) == 2 and fixed and dropped
    inc = pd.DataFrame({"채취일": [pd.Timestamp("2027-08-05")], "검사일": [pd.Timestamp("2025-08-06")]})
    out2, fixed2, _ = fix_future_dates(inc, "채취일", ref_col="검사일")
    assert out2["채취일"].iloc[0] == pd.Timestamp("2025-08-05") and "2025-08-05" in fixed2


def test_plant_specific_line_code():
    assert normalize_line("SA1KL3") == "3K" and normalize_line("5_1R/M") == "5-1R/M"
