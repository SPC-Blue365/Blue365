import math
from datetime import datetime

import pandas as pd
import pytest

from qms import lims
from qms.store import load_raw, save_raw

NOW = datetime(2026, 10, 8, 1, 0)


@pytest.fixture()
def demo_lims(demo_raw, tmp_path, monkeypatch):
    src = tmp_path / "src.db"
    save_raw(demo_raw, path=src, replace=True)
    path = tmp_path / "lims_demo.db"
    lims.create_demo_lims(load_raw(src), path)
    monkeypatch.setattr(lims, "LIMS_DEMO_DB", path)
    monkeypatch.setattr(lims, "LIMS_CONFIG_PATH", tmp_path / "lims.json")
    return tmp_path


def test_parse_value_variants():
    assert lims.parse_value("<0.1") == (0.1, True)
    assert lims.parse_value("1,234") == (1234.0, False)
    assert lims.parse_value("12,5") == (12.5, False)
    assert lims.parse_value(3) == (3.0, False)
    v, c = lims.parse_value("N.D.")
    assert math.isnan(v) and not c


def test_demo_sync_into_empty_db_reproduces_values(demo_lims, demo_raw):
    cfg = lims.LimsConfig(enabled=True, mode="demo")
    db = demo_lims / "qms.db"
    rep = lims.sync(cfg, db_path=db, now=NOW, log_path=demo_lims / "log.jsonl")
    assert not rep.errors, rep.errors
    assert rep.saved and rep.n_status_skipped > 0                       # 미승인(PENDING) 제외
    assert ("CEMENT_LOT", "WATER_DEMAND") in {(sp, tc) for sp, tc, _ in rep.unmapped}
    got = load_raw(db)
    a = got["clinker"].set_index("timestamp")["clk_fcao"]
    b = demo_raw["clinker"].set_index("timestamp")["clk_fcao"]
    j = pd.concat([a, b], axis=1, join="inner").dropna()
    assert len(j) > 50 and (j.iloc[:, 0] - j.iloc[:, 1]).abs().max() < 1e-6
    assert set(got["physical"]["product"]) <= {"1종", "3종"}           # 품종 코드 변환(OPC/HES)
    assert cfg.watermark == rep.watermark_after                         # 워터마크 갱신·저장
    rep2 = lims.sync(cfg, db_path=db, now=NOW, log_path=demo_lims / "log.jsonl")
    assert rep2.n_fetched == 0                                          # 증분: 새 결과 없음


def test_late_results_fill_existing_rows(demo_lims, tmp_path):
    db = tmp_path / "late.db"
    base = pd.DataFrame({"timestamp": [pd.Timestamp("2026-10-01 12:00")], "product": ["1종"], "phy_s3": [29.0],
                         "sand_lot": ["ISO-1"]})
    save_raw({"physical": base}, path=db, replace=True)
    late = pd.DataFrame([{"sampled_at": "2026-10-01 12:00", "sample_point": "CEMENT_LOT", "test_code": "CS28",
                          "value": "52.4", "product": "OPC", "status": "APPROVED", "updated_at": "2026-10-29 13:00"}])
    path = tmp_path / "exp"
    path.mkdir()
    late.to_csv(path / "lims_export.csv", index=False)
    cfg = lims.LimsConfig(enabled=True, mode="file", file_dir=str(path), watermark="2026-10-20 00:00:00")
    rep = lims.sync(cfg, db_path=db, log_path=tmp_path / "log.jsonl", persist_watermark=False)
    assert rep.cells_new == 1 and rep.cells_changed == 0
    got = load_raw(db)["physical"].iloc[0]
    assert got["phy_s28"] == pytest.approx(52.4) and got["phy_s3"] == pytest.approx(29.0) and got["sand_lot"] == "ISO-1"


def test_rest_mode_with_column_map(tmp_path):
    class Resp:
        def raise_for_status(self):
            return None

        def json(self):
            return {"result": {"items": [{"DT": "2026-10-07 10:00", "PT": "CLINKER", "ITEM": "FCAO", "RES": "1.25",
                                          "ST": "A"}]}}

    class Sess:
        def __init__(self):
            self.kw = None

        def get(self, url, **kw):
            self.kw = kw
            return Resp()

    sess = Sess()
    cfg = lims.LimsConfig(enabled=True, mode="rest", rest_url="https://lims.local/api", rest_records_path="result.items",
                          column_map={"DT": "sampled_at", "PT": "sample_point", "ITEM": "test_code", "RES": "value",
                                      "ST": "status"})
    rep = lims.sync(cfg, db_path=tmp_path / "r.db", session=sess, log_path=tmp_path / "log.jsonl",
                    persist_watermark=False)
    assert not rep.errors and rep.rows_by_table == {"clinker": 1}
    assert "updated_since" in sess.kw["params"]
    assert load_raw(tmp_path / "r.db")["clinker"]["clk_fcao"].iloc[0] == pytest.approx(1.25)


def test_password_never_saved_in_url(tmp_path):
    cfg = lims.LimsConfig(sql_url="mssql+pyodbc://user:secret@host/db")
    with pytest.raises(ValueError):
        lims.save_lims_config(cfg, tmp_path / "lims.json")
    cfg.sql_url = "mssql+pyodbc://user:{password}@host/db"
    lims.save_lims_config(cfg, tmp_path / "lims.json")
    assert "secret" not in (tmp_path / "lims.json").read_text(encoding="utf-8")


def test_file_mode_reads_cp949_csv(tmp_path):
    exp = tmp_path / "exp"
    exp.mkdir()
    pd.DataFrame([{"sampled_at": "2026-10-07 10:00", "sample_point": "CLINKER", "test_code": "FCAO", "value": "1.31",
                   "status": "APPROVED", "비고": "재시험"}]).to_csv(exp / "lims.csv", index=False, encoding="cp949")
    cfg = lims.LimsConfig(enabled=True, mode="file", file_dir=str(exp))
    assert lims.fetch_files(cfg)["비고"].iloc[0] == "재시험"
    rep = lims.sync(cfg, db_path=tmp_path / "f.db", log_path=tmp_path / "log.jsonl", persist_watermark=False)
    assert not rep.errors and rep.rows_by_table == {"clinker": 1}


def test_wide_export_file_sample_point_from_file_name(tmp_path):
    """가로형(시료당 1행) 엑셀: 파일 이름의 지점 코드로 채취 지점을 채우고, '-' 로 지정한 열은 버린다."""
    exp = tmp_path / "exp"
    exp.mkdir()
    pd.DataFrame([{"채취일시": "2026-10-07 10:00", "FCAO": 1.31, "LW": 1250, "비고": "재시험", "시료번호": "C-10"},
                  {"채취일시": "2026-10-07 12:00", "FCAO": "<0.5", "LW": 1262, "비고": None, "시료번호": "C-12"}]
                 ).to_excel(exp / "CLINKER_20261007.xlsx", index=False)
    (exp / "~$CLINKER_20261007.xlsx").write_bytes(b"lock")                   # 엑셀이 열어 둔 잠금 파일은 무시
    cfg = lims.LimsConfig(enabled=True, mode="file", file_dir=str(exp), column_map={"채취일시": "sampled_at", "시료번호": "-"})
    ok, msg, _prev = lims.test_connection(cfg)
    assert ok and "가로형" in msg, msg
    rep = lims.sync(cfg, db_path=tmp_path / "w.db", log_path=tmp_path / "log.jsonl", persist_watermark=False)
    assert not rep.errors, rep.errors
    assert rep.rows_by_table == {"clinker": 2} and rep.n_censored == 1
    assert ("CLINKER", "비고", 1) in rep.unmapped and all(tc != "시료번호" for _, tc, _ in rep.unmapped)
    got = load_raw(tmp_path / "w.db")["clinker"].sort_values("timestamp")
    assert got["clk_fcao"].tolist() == pytest.approx([1.31, 0.5]) and got["clk_lw"].tolist() == pytest.approx([1250, 1262])


def test_missing_sample_point_needs_column_file_name_or_default(tmp_path):
    exp = tmp_path / "exp"
    exp.mkdir()
    pd.DataFrame([{"sampled_at": "2026-10-07 10:00", "FCAO": 1.2}]).to_csv(exp / "export_1007.csv", index=False)
    cfg = lims.LimsConfig(enabled=True, mode="file", file_dir=str(exp))
    rep = lims.sync(cfg, db_path=tmp_path / "x.db", log_path=tmp_path / "log.jsonl", persist_watermark=False)
    assert rep.errors and "채취 지점" in rep.errors[0]
    ok, msg, _ = lims.test_connection(cfg)
    assert not ok and "채취 지점" in msg
    cfg.default_sample_point = "clinker"
    rep = lims.sync(cfg, db_path=tmp_path / "x.db", log_path=tmp_path / "log.jsonl", persist_watermark=False)
    assert not rep.errors and rep.rows_by_table == {"clinker": 1}


def test_prep_workbook_mapping_sheet_roundtrip():
    """준비서 ⑤ 시트에 회사 코드를 적어 올리면 그 행만 매핑으로 등록된다(단위 환산 계수 포함)."""
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(lims.prep_workbook()))
    assert wb.sheetnames == ["① 연동 절차", "② 연동 방식", "③ IT 요청서", "④ 표준 뷰 정의", "⑤ 시험항목 매핑표",
                             "⑥ 코드 변환", "⑦ 검증 체크리스트"]
    ws = wb["⑤ 시험항목 매핑표"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    filled = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, col["QMS 항목 코드"]).value
        if key in ("clk_fcao", "phy_s28", "cem_blaine"):
            ws.cell(r, col["[입력] 회사 채취지점 코드"], {"clk_fcao": "K-CLK", "phy_s28": "C-LOT",
                                                         "cem_blaine": "C-MILL"}[key])
            ws.cell(r, col["[입력] 회사 시험코드"], {"clk_fcao": "F_CAO", "phy_s28": "comp28",
                                                    "cem_blaine": "SSA"}[key])
            if key == "cem_blaine":
                ws.cell(r, col["계수"], 10)          # 예: m²/kg → cm²/g
            filled[key] = r
    assert len(filled) == 3
    buf = io.BytesIO()
    wb.save(buf)
    rows, warns = lims.parse_mapping_table(buf.getvalue(), "LIMS_연동_준비서.xlsx")
    assert not warns
    got = {(m["sample_point"], m["test_code"]): (m["table"], m["column"], m["factor"]) for m in rows}
    assert got == {("K-CLK", "F_CAO"): ("clinker", "clk_fcao", 1.0), ("C-LOT", "COMP28"): ("physical", "phy_s28", 1.0),
                   ("C-MILL", "SSA"): ("cement", "cem_blaine", 10.0)}


def test_mapping_csv_roundtrip_and_warnings():
    csv = pd.DataFrame(lims.default_mapping()).to_csv(index=False).encode("utf-8-sig")
    rows, warns = lims.parse_mapping_table(csv, "LIMS_매핑표.csv")
    assert not warns and rows == lims.default_mapping()
    bad = pd.DataFrame([{"채취 지점": "CLINKER", "시험코드": "FCAO", "QMS 항목(테이블.열)": "clinker.clk_fcao"},
                        {"채취 지점": "CLINKER", "시험코드": "XX", "QMS 항목(테이블.열)": "clinker.nope"},
                        {"채취 지점": "CLINKER", "시험코드": "FCAO", "QMS 항목(테이블.열)": "clinker.clk_fcao", "계수": "2"}])
    rows, warns = lims.parse_mapping_table(bad.to_csv(index=False).encode("cp949"), "map.csv")
    assert len(rows) == 1 and rows[0]["factor"] == 2.0
    assert any("알 수 없는 QMS 항목" in w for w in warns) and any("중복" in w for w in warns)
    with pytest.raises(ValueError):
        lims.parse_mapping_table(pd.DataFrame({"a": [1]}).to_csv(index=False).encode(), "x.csv")
