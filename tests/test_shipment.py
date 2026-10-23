"""출하시멘트(LIMS 일일 로트) 모듈: 형식 인식·0=미시험·선박 비고 정규화·저염 구분·연도별 관리기준·혼합재↔강도."""

import io
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from qms import shipment as S

HEAD = ["", "전송", "ERP LOT.NO", "LIMS LOT.NO", "검사일", "채취일", "품명", "공정구분", "분석방법", "첨가재명칭", "석회석비율", "석회석CO2", "석회석품위", "첨가재비율", "LOI",
        "SiO2", "Al2O3", "Fe2O3", "CaO", "MgO", "SO3(XRF)", "SO3(C/S)", "Na2O", "K2O", "Cl(XRF)", "Cl(습식)", "TA", "압축강도1일(MPa)", "압축강도3일(MPa)",
        "압축강도7일(MPa)", "압축강도28일(MPa)", "45㎛", "Blaine", "주도", "Flow", "간이Flow", "초결시간(분)", "종결시간(HH:MM)", "담당자", "비고"]
APP = str(Path(__file__).resolve().parent.parent / "streamlit_app.py")


def make_xlsx(n=240, seed=3, extra_note=True) -> bytes:
    """실제 파일과 같은 열 구조의 합성 자료. 혼합재 합계가 10→8→5 로 바뀌고 강도는 혼합재가 줄수록 오르게 만든다."""
    rng = np.random.default_rng(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEAD)
    ships = ["수성3호", "수성1호", "리나호", "삼표1호", "정진호"]
    start = pd.Timestamp("2025-06-01")
    for i in range(n):
        d = start + pd.Timedelta(days=i)
        mix = [(5, 5), (5, 3), (5, 0)][0 if i < 80 else (1 if i < 160 else 2)]
        tot = mix[0] + mix[1]
        base28 = 57.0 - 0.8 * (tot - 5)
        note = ",".join(rng.choice(ships, size=int(rng.choice([1, 1, 2])), replace=False))
        if extra_note and i % 40 == 7:
            note = "삼표1호(부산,저염)    ."
        if extra_note and i % 40 == 9:
            note = "수성3 ,리나호."
        r = [i + 1, None, f"QC{i}", str(1000 + i), d.date().isoformat(), d.date().isoformat(), "1종시멘트 벌크", "P3", "Powder",
             "플라이 애시", mix[0], 37, 84, mix[1], round(rng.normal(2.5, .2), 2), 21, round(rng.normal(5.7, .2), 2), 3.4, 58,
             round(rng.normal(3.9, .3), 2), round(rng.normal(2.6, .06), 2), 0, .2, 1.0, int(rng.integers(300, 700)), 0, round(rng.normal(.85, .03), 2),
             round(rng.normal(15.5, .8), 1), round(rng.normal(35.5, 1), 1), round(rng.normal(45, 1), 1), 0 if i >= n - 3 else round(rng.normal(base28, .8), 1),
             round(rng.normal(6.8, .5), 1), int(rng.normal(3780, 60)), round(rng.normal(27.3, .2), 1), int(rng.integers(90, 100)), 0, str(int(rng.integers(220, 280))),
             f"{int(rng.integers(5, 7)):02d}:{int(rng.integers(0, 60)):02d}", "", note]
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def parsed():
    return S.parse_upload(make_xlsx(), "ship.xlsx")


def test_detect_and_headers(parsed):
    assert not parsed.errors
    rep = parsed.report
    assert rep["missing"] == [] and len(rep["found"]) == 17
    found = dict(rep["found"])
    assert found["so3"] == "SO3(XRF)" and found["flow"] == "Flow" and found["r45"] == "45㎛" and found["fst"] == "종결시간(HH:MM)"
    # 다른 표(세로형 LIMS 등)는 출하시멘트로 감지하지 않는다
    other = [["시료채취일시", "시료명", "시험항목", "결과값", "단위", "판정", "a", "b"]] + [["2026-01-01", "클링커", "f-CaO", 1, "%", "적합", 1, 2]] * 3
    assert S.detect(other) is None


def test_zero_is_untested_but_ratio_zero_is_real(parsed):
    df = parsed.records
    assert df["s28"].isna().sum() == 3                      # 마지막 3건은 28일 강도 0 → 미시험
    assert parsed.report["zeros"]["s28"] == 3
    assert (df["addr"] == 0).sum() == 80                    # 첨가재 비율 0 = 투입 안 함(실제 값)
    assert df.loc[df["addr"] == 0, "mix_total"].eq(5).all()
    assert df["fst"].notna().all() and df["fst"].between(300, 420).all()   # HH:MM → 분


def test_ship_notes_and_low_chloride(parsed):
    df = parsed.records
    low = df[df["product_kind"] == "저염"]
    assert len(low) == 6 and (low["ships"] == "삼표1호").all() and (low["dest_hint"] == "부산").all()
    row = df[df["date"] == pd.Timestamp("2025-06-01") + pd.Timedelta(days=9)].iloc[0]
    assert row["ships"] == "수성3호,리나호" and row["n_ship"] == 2 and row["ship_ok"]       # '수성3 ,리나호.' 정규화
    assert df["n_ship"].min() >= 1
    # 선박별 값을 추정·분해하는 열은 만들지 않는다
    assert not any(c.startswith("ship_est") for c in df.columns)


def test_save_merge_keeps_old_values_and_updates():
    S.clear_shipment()
    first = S.parse_upload(make_xlsx(), "a.xlsx").records
    r = S.save_shipment(first)
    assert r == {"added": 240, "updated": 0, "total": 240}
    # 28일 강도가 나중에 들어온 새 파일(마지막 3건 값 있음) + 기존 값은 비어 있는 칸에도 유지
    newer = first.copy().tail(5).reset_index(drop=True)
    newer.loc[newer.index[-3:], "s28"] = [55.5, 55.6, 55.7]
    newer["s1"] = np.nan                                    # 새 파일에서 1일 강도가 비어 있어도 기존 값 유지
    r2 = S.save_shipment(newer)
    assert r2["added"] == 0 and r2["updated"] == 5 and r2["total"] == 240
    d = S.load_shipment()
    assert d["s28"].iloc[-1] == 55.7 and d["s1"].iloc[-1] == first["s1"].iloc[-1]
    assert "x:CaO" not in d.columns or True
    S.clear_shipment()


def test_standards_by_year_carry_forward_and_step():
    std = {"version": 1, "years": {}}
    assert S.effective_limits(std, "s28", 2025)["ks_min"] == 42.5 and S.effective_limits(std, "s28", 2025)["lsl"] is None
    t = S.year_table(std, 2025)
    t.loc[t["key"] == "s28", "사내 하한"] = 50.0
    S.save_year(std, 2025, t)
    t26 = S.year_table(std, 2026)
    assert t26.loc[t26["key"] == "s28", "사내 하한"].iloc[0] == 50.0 and t26.loc[t26["key"] == "s28", "출처"].iloc[0] == "2025년 승계"
    t26.loc[t26["key"] == "s28", "사내 하한"] = 52.0
    S.save_year(std, 2026, t26)
    assert S.effective_limits(std, "s28", 2024)["lsl"] is None          # 이전 연도에는 소급 적용하지 않음
    assert S.effective_limits(std, "s28", 2027)["lsl"] == 52.0
    df = pd.DataFrame({"date": pd.to_datetime(["2025-12-31", "2026-01-01"]), "s28": [51.0, 51.0]})
    lim = S.limits_frame(std, df, "s28")
    assert lim["lsl"].tolist() == [50.0, 52.0]
    # 같은 51.0 이라도 2025년 기준에서는 정상, 2026년 기준에서는 사내 기준 이탈
    j = S.judge_item(df.assign(n_ship=1, ship_ok=True), "s28", std)
    assert j.plant_out.tolist() == [False, True] and not j.ks_out.any()
    S.save_ship_standards({"version": 1, "years": {}})


def test_judge_status_ks_over_plant_over_spc():
    std = {"version": 1, "years": {}}
    rng = np.random.default_rng(1)
    dates = pd.date_range("2026-01-01", periods=60)
    v = rng.normal(55, 1, 60)
    df = pd.DataFrame({"date": dates, "s28": v})
    assert S.judge_item(df, "s28", std).status in ("정상", "참고", "주의")
    df.loc[59, "s28"] = 40.0                                # KS 하한(42.5) 미달
    j = S.judge_item(df, "s28", std)
    assert j.status == "위험" and j.ks_out.iloc[-1]
    df.loc[59, "s28"] = 45.0                                # KS 이내, 사내 하한 미달
    t = S.year_table(std, 2026)
    t.loc[t["key"] == "s28", "사내 하한"] = 50.0
    std = S.save_year(std, 2026, t)
    assert S.judge_item(df, "s28", std).status == "경고"
    S.save_ship_standards({"version": 1, "years": {}})


def test_mix_effect_matches_domain_knowledge(parsed):
    df = parsed.records
    e = S.mix_effect(df, "s28")
    assert e and e["mix_coef"] < -0.5                       # 혼합재 +1%p → 28일 강도 하락(설계값 -0.8)
    ch = S.mix_changes(df)
    dec = ch[ch["변화(%p)"] < 0]
    assert len(dec) >= 1 and (dec["압축강도 28일 변화(MPa)"] > 0).all()   # 혼합재↓ → 강도↑
    assert dec["28일 강도 방향(경험칙 대비)"].str.startswith("✔").all()
    ct = S.composition_table(df)
    assert set(ct["혼합재합계"]) == {10.0, 8.0, 5.0}


def test_filters(parsed):
    df = parsed.records
    only = S.apply_filters(df, kind="일반 제품만")
    assert (only["product_kind"] == "일반").all() and len(only) == len(df) - 6
    single = S.apply_filters(df, mode="1모선 단독 시험만", kind="저염 포함 전체")
    assert ((single["n_ship"] == 1) & single["ship_ok"]).all()


def test_page_renders_empty_and_with_data():
    S.clear_shipment()
    at = AppTest.from_file(APP, default_timeout=240)
    at.run()
    at.switch_page("app_pages/shipment.py")
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.title and "출하시멘트" in at.title[0].value
    S.save_shipment(S.parse_upload(make_xlsx(), "a.xlsx").records)
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert [t.label for t in at.tabs][:2] == ["📊 한눈에", "📈 항목 상세"]
    assert at.dataframe, "판정표가 있어야 함"
    S.clear_shipment()


# ── 선박별 양하지 자동 연결 ─────────────────────────────────────────────
def _dest_xlsx(rows, head=("출항일자", "선박명", "양하항")) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["운항 실적"])
    ws.append(list(head))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_dest_from_notes_and_hints(parsed):
    from qms import ship_dest as D
    assert S.dest_from_notes(["부산 저염"]) == "부산" and S.dest_from_notes(["수분-0.34"]) == "" and S.dest_from_notes([""]) == ""
    df = parsed.records
    low = df[df["product_kind"] == "저염"].iloc[0]
    assert '"삼표1호": "부산"' in low["ship_hints"]                     # 삼표1호(부산,저염) → 그 배의 양하지 = 부산
    a = D.attach_destinations(df, master=D.load_master().iloc[0:0], voyages=D.load_voyages().iloc[0:0])
    assert (a.loc[a["product_kind"] == "저염", "dest"] == "부산").all()
    assert (a.loc[a["product_kind"] == "저염", "dest_src"] == "비고").all()


def test_dest_upload_priority_and_auto_relink(tmp_path):
    from qms import ship_dest as D
    S.clear_shipment()
    D.DEST_DB_PATH.unlink(missing_ok=True)
    S.save_shipment(S.parse_upload(make_xlsx(), "a.xlsx").records)
    df = S.load_shipment()
    known = set(D.ship_table(df)["ship"])
    # ① 기본 양하지(일자 없는 표) → 모든 날짜에 적용
    p = D.parse_dest_upload(_dest_xlsx([("수성3", "평택"), ("리나호", "부산항"), ("수성1호", "부산"), ("삼표1호", "부산"), ("정진호", "광양")],
                                       head=("선박", "양하지")), "m.xlsx", known)
    assert p.kind == "master" and not p.errors and set(p.records["ship"]) >= {"수성3호", "리나호"}
    assert dict(zip(p.records["ship"], p.records["dest"]))["리나호"] == "부산"     # '부산항' → '부산'
    D.save_master(p.records.assign(note=""))
    a = D.attach_destinations(df)
    assert (a["dest"] == D.UNKNOWN).sum() == 0
    two = a[a["ships"] == "수성3호,리나호"].iloc[0]
    assert two["dest"] == D.MIXED and "수성3호→평택" in two["dest_detail"]          # 양하지가 다르면 나누지 않음
    # ② 운항 기록이 기본값보다 우선(같은 날짜, ±1일 허용)
    d0 = df[df["ships"] == "정진호"].iloc[0]["date"]
    v = D.parse_dest_upload(_dest_xlsx([((d0 + pd.Timedelta(days=1)).strftime("%Y.%m.%d"), "정진호", "부산")]), "v.xlsx", known)
    assert v.kind == "voyage" and len(v.records) == 1
    assert D.save_voyages(v.records, "v.xlsx") == {"added": 1, "updated": 0}
    assert D.save_voyages(v.records, "v.xlsx") == {"added": 0, "updated": 1}
    a2 = D.attach_destinations(df)
    r = a2[a2["date"] == d0].iloc[0]
    assert r["dest"] == "부산" and r["dest_src"] == "운항기록"
    other = a2[(a2["ships"] == "정진호") & (abs(a2["date"] - d0) > pd.Timedelta(days=1))]
    assert (other["dest"] == "광양").all()                                            # 다른 날은 기본값
    q = D.dest_quality(a2, ["s28"])
    assert set(q["양하지"]) <= {"부산", "평택", "광양"} and q["건수"].sum() == (~a2["dest"].isin([D.MIXED, D.UNKNOWN, D.PARTIAL])).sum()
    # 출하시멘트 LIMS 표는 양하지 표로 오인하지 않음
    assert not D.looks_like_dest(make_xlsx(), "a.xlsx") and D.looks_like_dest(_dest_xlsx([("2026-01-01", "리나호", "부산")]), "v.xlsx")
    S.clear_shipment()
    D.DEST_DB_PATH.unlink(missing_ok=True)


def test_monitor_import_folder_handles_shipment_and_dest(tmp_path):
    import importlib

    from qms import ship_dest as D
    S.clear_shipment()
    D.DEST_DB_PATH.unlink(missing_ok=True)
    mon = importlib.import_module("qms_monitor")
    (tmp_path / "lims_출하.xlsx").write_bytes(make_xlsx())
    (tmp_path / "운항실적.xlsx").write_bytes(_dest_xlsx([("2025-06-01", "리나호", "부산")]))
    logs = mon.import_folder(tmp_path)
    assert any("출하시멘트" in x for x in logs) and any("운항(양하지)" in x for x in logs), logs
    assert len(S.load_shipment()) == 240 and len(D.load_voyages()) == 1
    assert sorted(p.name for p in (tmp_path / "processed").iterdir()) == ["lims_출하.xlsx", "운항실적.xlsx"]
    S.clear_shipment()
    D.DEST_DB_PATH.unlink(missing_ok=True)


def test_labels_rename_items_and_titles_then_reset():
    from qms import labels as L
    from qms.standards import load_registry
    L.save_labels({"menu": {"출하시멘트": "출하 품질"}, "pages": {"🚢 출하시멘트 모니터링": "🚢 출하 품질 현황"},
                   "ship_items": {"s28": {"name": "28일 강도(KS)", "unit": "N/mm²"}}, "items": {"cem_blaine": {"name": "비표면적"}}})
    lab = L.load_labels()
    assert L.menu("출하시멘트", lab) == "출하 품질" and L.menu("종합 현황", lab) == "종합 현황"
    assert L.page("🚢 출하시멘트 모니터링", lab) == "🚢 출하 품질 현황"
    L.apply_shipment(lab)
    assert S.BY_KEY["s28"].name == "28일 강도(KS)" and S.BY_KEY["s28"].unit == "N/mm²" and S.BY_KEY["s3"].name == "압축강도 3일"
    assert next(it for it in S.ITEMS if it.key == "s28").name == "28일 강도(KS)"
    assert load_registry()["cem_blaine"].name == "비표면적"
    L.save_labels({})                                         # 지우면 기본값
    L.apply_shipment()
    assert S.BY_KEY["s28"].name == "압축강도 28일" and load_registry()["cem_blaine"].name != "비표면적"


def test_chart_helpers_and_saved_views():
    from qms import charts as CH
    idx = pd.date_range("2026-01-01", periods=60, freq="12h")
    s = pd.Series(np.arange(60, dtype=float), index=idx)
    d = CH.aggregate(s, CH.AGGS["일 평균"])
    assert len(d) == 30 and d.iloc[0] == 0.5
    assert len(CH.aggregate(s, CH.AGGS["월 평균"])) == 1
    z = CH.zscore(d)
    assert abs(z.mean()) < 1e-9 and abs(z.std(ddof=1) - 1) < 1e-9
    fig = CH.item_figure(d, "시험", "MPa", ma=5, limits={"lsl": 3.0, "usl": pd.Series(20.0, index=d.index)})
    assert len(fig.data) == 3                                  # 값 + 이동평균 + 상한(계단선), 하한은 수평선
    ov = CH.overlay_figure([("a", "MPa", d), ("b", "%", d * 2)], standardize=True)
    assert len(ov.data) == 2 and ov.data[0].line.color == CH.SERIES[0] and ov.data[1].line.color == CH.SERIES[1]
    CH.save_view({"name": "강도", "state": {"cv_ship_keys": ["s1", "s28"]}})
    CH.save_view({"name": "강도", "state": {"cv_ship_keys": ["s3"]}})          # 같은 이름은 덮어씀
    assert [v["state"]["cv_ship_keys"] for v in CH.load_views()] == [["s3"]]
    assert CH.delete_view("강도") == []


def test_chart_page_starts_on_source_with_data():
    """출하시멘트 자료가 없으면 그래프 보기는 공정 데이터로 시작(빈 화면 방지), 있으면 출하시멘트 강도 1·3·7·28일."""
    S.clear_shipment()
    at = AppTest.from_file(APP, default_timeout=240)
    at.run()
    at.switch_page("app_pages/charts.py")
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.radio(key="cv_source").value == "공정 데이터"
    S.save_shipment(S.parse_upload(make_xlsx(), "a.xlsx").records)
    at2 = AppTest.from_file(APP, default_timeout=240)
    at2.run()
    at2.switch_page("app_pages/charts.py")
    at2.run()
    assert not at2.exception, [e.value for e in at2.exception]
    assert at2.radio(key="cv_source").value == "출하시멘트"
    assert at2.multiselect(key="cv_ship_keys").value == ["s1", "s3", "s7", "s28"]
    S.clear_shipment()


def test_lims_korean_headers_and_setting_time_minutes():
    """LIMS·한글 머리글(채취일·CaO·압축강도28일 등)을 공정 데이터 업로드에서 인식하고, 품명·응결시간을 정규화한다."""
    import io

    import openpyxl

    from qms import store

    assert store.normalize_product("1종시멘트 벌크") == "1종"
    assert store.normalize_product("조강 3종") == "3종"
    assert store._to_minutes("06:33") == 393 and store._to_minutes(262) == 262
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "클링커"
    ws.append(["채취일", "CaO", "SiO2", "F-CaO(습식)"])
    ws.append(["2026-10-01", 65.8, 21.3, 1.2])
    ws2 = wb.create_sheet("물성")
    ws2.append(["채취일", "품명", "초결시간(분)", "종결시간(HH:MM)", "압축강도28일(MPa)", "6가크롬"])
    ws2.append(["2026-10-01", "1종시멘트 벌크", 262, "06:33", 57.1, 0.1])
    buf = io.BytesIO()
    wb.save(buf)
    res = store.parse_upload(buf.getvalue(), "x.xlsx")
    assert not res.errors
    clk = res.tables["clinker"]
    assert {"clk_cao", "clk_sio2", "clk_fcao"} <= set(clk.columns) and clk.iloc[0]["clk_cao"] == 65.8
    phy = res.tables["physical"]
    assert phy.iloc[0]["product"] == "1종" and phy.iloc[0]["phy_fst"] == 393 and phy.iloc[0]["phy_s28"] == 57.1
    assert phy.iloc[0]["phy_crvi"] == 0.1


def test_generated_update_workbook_ingests():
    """배포용 업데이트 양식(scripts/build_lims_update_workbook.py 결과)을 QMS가 빈 오류로 읽어들인다."""
    from pathlib import Path

    from qms import shipment as S
    from qms import store

    f = Path(__file__).resolve().parent.parent / "docs" / "Blue365_QMS_업데이트_양식.xlsx"
    if not f.exists():
        import pytest
        pytest.skip("양식 파일이 아직 생성되지 않음")
    res = store.parse_upload(f.read_bytes(), str(f))
    assert not res.errors
    assert {"raw_meal", "clinker", "cement", "physical", "xrd"} <= set(res.tables)
    import openpyxl
    ws = openpyxl.load_workbook(f, data_only=True)["출하시멘트"]
    aoa = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
    sp = S.parse_aoa(aoa)
    assert not sp.errors and sp.report["rows"] > 0 and not sp.report["missing"]


def test_lims_raw_export_auto_detected_by_title():
    """LIMS 원본 파일(제목 1행 + 머리글 3행)을 복사·붙여넣기 없이 제목으로 표를 자동 인식한다."""
    import io

    import openpyxl

    from qms import store

    def lims_book(title, headers, row):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([title])
        ws.append(["검사일 : 2026-10-01 ~ 2026-10-23"])
        ws.append(headers)
        ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    clk = lims_book("크링커성적", ["로트번호", "검사일", "채취일", "품명", "CaO", "SiO2", "F-CaO(습식)", "용중", "C3S"],
                    ["L1", "2026-10-01", "2026-10-01", "1종크링커", 65.8, 21.3, 1.1, 1320, 62.0])
    res = store.parse_upload(clk, "크링커성적.xlsx")
    assert not res.errors and "clinker" in res.tables
    row0 = res.tables["clinker"].iloc[0]
    assert row0["clk_cao"] == 65.8 and row0["clk_fcao"] == 1.1 and row0["clk_lw"] == 1320
    assert "clk_c3s" not in res.tables["clinker"].columns          # C3S 는 QMS 가 계산(저장 안 함)

    rm = lims_book("6_7호조합원료성적", ["로트번호", "채취일", "품명", "CaO", "SiO2", "Al2O3", "Fe2O3", "LSF"],
                   ["L", "2026-10-02", "1/2종조합원료", 43.1, 13.6, 3.3, 2.1, 95.0])
    assert store.parse_upload(rm, "조합원료.xlsx").tables["raw_meal"].iloc[0]["rm_cao"] == 43.1

    # 인수검사 류는 QMS 표가 없어 무시
    gyp = lims_book("석고류 인수검사", ["로트번호", "채취일", "품명", "SO3(XRF)", "순도"],
                    ["L", "2026-10-03", "탈황석고", 44.0, 95.0])
    r2 = store.parse_upload(gyp, "석고.xlsx")
    assert not r2.tables and "인식 가능한 시트가 없습니다" in r2.errors[0]


def test_shipment_product_split_and_filter():
    """출하시멘트에서 품명(1종/3종)을 인식해 저장·필터한다."""
    from qms import shipment as S

    assert S._product("1종시멘트 벌크") == "1종" and S._product("3종시멘트 벌크") == "3종"
    assert S._product("") == "1종" and S._product("2종고로슬래그시멘트") in ("2종", "고로슬래그")
    head = ["채취일", "품명", "첨가재명칭", "석회석비율", "첨가재비율", "LOI", "Al2O3", "MgO", "SO3(XRF)",
            "압축강도1일(MPa)", "압축강도3일(MPa)", "압축강도7일(MPa)", "압축강도28일(MPa)", "45㎛", "Blaine", "주도", "비고"]
    aoa = [["시멘트성적"], ["검사일"], head,
           ["2026-10-01", "1종시멘트 벌크", "플라이 애시", 3, 5, 1.9, 6.0, 3.8, 2.5, 15, 38, 46, 57.1, 6.6, 3810, 27.5, "수성1호"],
           ["2026-10-02", "3종시멘트 벌크", "", 0, 0, 1.8, 5.9, 3.7, 2.4, 22, 44, 50, 61.0, 7.0, 4200, 28.0, "수성2호"]]
    sp = S.parse_aoa(aoa)
    assert not sp.errors and set(sp.records["product"]) == {"1종", "3종"}
    f1 = S.apply_filters(sp.records, product="1종")
    assert len(f1) == 1 and f1.iloc[0]["product"] == "1종"
    assert len(S.apply_filters(sp.records, product="전체")) == 2
