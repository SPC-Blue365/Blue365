import io

import pandas as pd
from openpyxl import load_workbook
from pptx import Presentation

from qms.alerts import events_frame
from qms.diagnosis import diagnose
from qms.notify import NotifyConfig, dispatch, format_message, send_email
from qms.reports import (
    build_excel_report,
    build_issue_ppt,
    build_management_ppt,
    capability_table,
)
from qms.standards import SpecRegistry
from qms.store import (
    load_alert_status,
    load_raw,
    parse_upload,
    save_raw,
    template_workbook,
    update_alert_status,
)


def test_save_load_roundtrip_and_upsert(demo_raw, tmp_path):
    db = tmp_path / "qms.db"
    small = {k: v.head(20) for k, v in demo_raw.items()}
    counts = save_raw(small, path=db, replace=True)
    assert counts["clinker"] == 20
    changed = {"clinker": small["clinker"].head(3).assign(clk_fcao=9.99)}
    save_raw(changed, path=db)
    raw = load_raw(db)
    assert len(raw["clinker"]) == 20                       # 중복은 덮어쓰기
    assert (raw["clinker"].head(3)["clk_fcao"] == 9.99).all()


def test_upsert_is_column_wise_so_late_results_keep_earlier_columns(demo_raw, tmp_path):
    db = tmp_path / "qms.db"
    lots = demo_raw["physical"].head(4).copy()
    first = lots.drop(columns=["phy_s28"])                       # 1~7일 결과 먼저 입력
    save_raw({"physical": first}, path=db, replace=True)
    late = lots[["timestamp", "product"]].assign(phy_s28=[50.1, 50.2, 50.3, 50.4])   # 28일 결과만 나중에
    save_raw({"physical": late}, path=db)
    got = load_raw(db)["physical"]
    assert len(got) == 4
    assert got["phy_s28"].tolist() == [50.1, 50.2, 50.3, 50.4]
    assert got["phy_s3"].round(3).tolist() == lots["phy_s3"].round(3).tolist()   # 먼저 들어온 값 유지
    assert got["sand_lot"].tolist() == lots["sand_lot"].tolist()


def test_template_upload_roundtrip(demo_raw):
    data = template_workbook(SpecRegistry(), {k: v.head(5) for k, v in demo_raw.items()})
    res = parse_upload(data, "template.xlsx")
    assert not res.errors
    assert set(res.tables) == {"raw_meal", "kiln", "clinker", "cement", "physical", "xrd"}
    assert len(res.tables["cement"]) == 5 and "product" in res.tables["cement"]


def test_upload_rejects_missing_product_column():
    df = pd.DataFrame({"timestamp": ["2026-10-01 00:00"], "cem_blaine": [3400]})
    res = parse_upload(df.to_csv(index=False).encode(), "cement_2026.csv")
    assert res.errors and "product" in res.errors[0]


def test_alert_status_persistence(tmp_path):
    db = tmp_path / "alerts.db"
    update_alert_status("abc123", path=db, status="조치중", assignee="품질팀", note="재시험")
    st = load_alert_status(db)
    assert st.loc[st["event_id"] == "abc123", "status"].iloc[0] == "조치중"


def test_notify_format_and_dispatch_filters(demo):
    store, reg, settings, _, events = demo
    e = next(x for x in events if x.severity == "위험")
    subj, body = format_message(e, "요약", ["조치1"], "http://dash")
    assert "[QMS 위험]" in subj and "조치1" in body and "http://dash" in body
    cfg = NotifyConfig(min_severity="경고")
    status = pd.DataFrame({"event_id": [e.event_id], "notified": [1]})
    res = dispatch(events, status, cfg, dry_run=True)
    ids = {r["event_id"] for r in res}
    assert e.event_id not in ids                                  # 이미 발송
    assert all(x.severity in ("위험", "경고") for x in events if x.event_id in ids)
    ok, msg = send_email(NotifyConfig(), "s", "b")
    assert not ok and "설정" in msg


def test_reports_build_valid_files(demo):
    store, reg, settings, _, events = demo
    start, end = pd.Timestamp("2026-09-08"), pd.Timestamp("2026-10-07 23:59")
    ev = [e for e in events if e.end >= start]
    edf = events_frame(ev)
    dxs = [diagnose(e, store, reg, settings) for e in ev if e.severity in ("위험", "경고")]
    wb = load_workbook(io.BytesIO(build_excel_report(store, reg, edf, dxs, start, end, include_raw=False)))
    assert wb.sheetnames[:4] == ["요약", "공정능력", "알림목록", "원인진단"]
    prs = Presentation(io.BytesIO(build_management_ppt(store, reg, edf, dxs, start, end)))
    assert len(prs.slides) >= 8
    issue = Presentation(io.BytesIO(build_issue_ppt(dxs[0], store, reg)))
    assert len(issue.slides) == 5
    cap = capability_table(store, reg, start, end, key_only=True)
    assert {"Cpk", "등급"} <= set(cap.columns) and len(cap) > 5


def test_csv_upload_reads_korean_excel_cp949():
    """한글 Windows 엑셀의 'CSV(쉼표로 분리)' 저장 파일(CP949)과 UTF-8(BOM) 파일을 모두 읽는다."""
    df = pd.DataFrame({"timestamp": ["2026-10-01 08:00", "2026-10-01 16:00"], "product": ["1종", "3종"],
                       "phy_s3": [29.1, 33.0]})
    res = parse_upload(df.to_csv(index=False).encode("cp949"), "physical_202610.csv")
    assert not res.errors, res.errors
    assert list(res.tables["physical"]["product"]) == ["1종", "3종"]
    res2 = parse_upload(df.to_csv(index=False).encode("utf-8-sig"), "physical.csv")
    assert res2.tables["physical"]["phy_s3"].iloc[1] == 33.0
