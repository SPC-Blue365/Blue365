"""무인 감시 스크립트: 데이터 자동 반영 → 이상 감지 → 알림 발송.

사용 예 (10~30분 주기로 cron / Windows 작업 스케줄러에 등록)
  python qms_monitor.py --since-hours 24
  python qms_monitor.py --import-dir ./inbox --since-hours 24
  python qms_monitor.py --dry-run          # 발송하지 않고 대상만 출력
  python qms_monitor.py --sync-lims        # LIMS 결과 증분 동기화 후 감시(🔗 LIMS 연동 화면에서 설정)

--import-dir: LIMS·DCS가 내보낸 엑셀/CSV를 넣어두는 폴더. 처리한 파일은 processed/ 로 옮긴다.
--sync-lims : data/lims.json 설정(SQL·REST·파일)으로 승인된 시험 결과를 가져와 열 단위로 병합 저장한다.
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

import pandas as pd

from qms import incoming, ship_dest, shipment
from qms.alerts import detect_events
from qms.diagnosis import diagnose
from qms.lims import load_lims_config, sync
from qms.notify import dispatch, load_config
from qms.prediction import attach_predictions
from qms.standards import load_registry, load_settings
from qms.store import (
    load_alert_status,
    load_store,
    parse_upload,
    save_raw,
    update_alert_status,
)


def import_folder(folder: Path) -> list[str]:
    logs = []
    done = folder / "processed"
    done.mkdir(parents=True, exist_ok=True)
    for f in sorted(folder.iterdir()):
        if f.suffix.lower() not in (".xlsx", ".xls", ".csv") or not f.is_file():
            continue
        content = f.read_bytes()
        if shipment.looks_like_shipment(content, f.name):          # LIMS 출하시멘트 일일 로트 표
            p = shipment.parse_upload(content, f.name)
            if p.errors:
                logs.append(f"[오류] {f.name}: {'; '.join(p.errors)}")
                continue
            r = shipment.save_shipment(p.records)
            shutil.move(str(f), done / f.name)
            logs.append(f"[반영] {f.name}: 출하시멘트 추가 {r['added']}건·갱신 {r['updated']}건")
            continue
        if ship_dest.looks_like_dest(content, f.name):              # 선박 운항(양하지) 기록·기본 양하지
            p = ship_dest.parse_dest_upload(content, f.name)
            if p.errors:
                logs.append(f"[오류] {f.name}: {'; '.join(p.errors)}")
                continue
            if p.kind == "voyage":
                r = ship_dest.save_voyages(p.records, f.name)
                msg = f"운항(양하지) 기록 추가 {r['added']}건·갱신 {r['updated']}건"
            else:
                cur = ship_dest.load_master()
                merged = pd.concat([cur[["ship", "dest", "note"]], p.records.assign(note="")[["ship", "dest", "note"]]]).drop_duplicates("ship", keep="last")
                msg = f"선박 기본 양하지 {ship_dest.save_master(merged)}척"
            shutil.move(str(f), done / f.name)
            logs.append(f"[반영] {f.name}: {msg}")
            continue
        if incoming.looks_like_incoming(content, f.name):           # 원료·석고·고령토 인수검사
            p = incoming.parse_incoming(content, f.name)
            if p.errors:
                logs.append(f"[오류] {f.name}: {'; '.join(p.errors)}")
                continue
            n = incoming.save_incoming(p.rows)
            shutil.move(str(f), done / f.name)
            logs.append(f"[반영] {f.name}: {p.kind} 인수검사 누계 {n}행")
            continue
        res = parse_upload(content, f.name)
        if res.errors:
            logs.append(f"[오류] {f.name}: {'; '.join(res.errors)}")
            continue
        counts = save_raw(res.tables, replace=False)
        shutil.move(str(f), done / f.name)
        logs.append(f"[반영] {f.name}: " + ", ".join(f"{k} {v}행" for k, v in counts.items()))
    return logs


def sync_file(path: Path) -> list[str]:
    """고정된 한 개 워크북(공정·출하시멘트·인수검사 시트 혼재 가능)을 제자리에서 읽어 반영한다(파일을 옮기지 않음).

    매일 같은 파일(예: Blue365_QMS.xlsx)을 덮어써 두면 작업 스케줄러가 이 함수로 다시 읽어 반영한다.
    """
    import io

    import openpyxl

    logs: list[str] = []
    if not path.is_file():
        return [f"[오류] 파일이 없습니다: {path}"]
    content = path.read_bytes()

    res = parse_upload(content, path.name)                       # 공정 데이터(템플릿·LIMS 원본 시트)
    if res.tables:
        counts = save_raw(res.tables, replace=False)
        logs.append("공정 데이터: " + ", ".join(f"{k} {v}행" for k, v in counts.items()))

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        logs.append(f"[경고] 출하·인수검사 시트 확인 실패: {exc}")
        return logs
    for ws in wb.worksheets:
        aoa = [list(r) for r in ws.iter_rows(values_only=True)]
        if not aoa:
            continue
        if shipment.detect(aoa):                                 # 출하시멘트 시트
            sp = shipment.parse_aoa(aoa)
            if not sp.errors:
                r = shipment.save_shipment(sp.records)
                logs.append(f"출하시멘트({ws.title}): 추가 {r['added']}·갱신 {r['updated']}")
                continue
        ip = incoming.parse_incoming_df(pd.DataFrame(aoa))        # 인수검사 시트
        if not ip.errors and not ip.rows.empty:
            n = incoming.save_incoming(ip.rows)
            logs.append(f"인수검사({ws.title}·{ip.kind}): 누계 {n}행")
    if not logs:
        logs.append("반영할 데이터를 찾지 못했습니다(시트·머리글 확인).")
    return logs


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Blue365 QMS 무인 감시")
    ap.add_argument("--import-dir", type=Path, help="자동 반영할 데이터 파일 폴더")
    ap.add_argument("--sync-file", type=Path, help="고정 워크북 1개를 제자리에서 반영(매일 같은 파일 갱신용)")
    ap.add_argument("--since-hours", type=float, default=24, help="최근 N시간 내 종료된 이벤트만 발송(기본 24)")
    ap.add_argument("--dry-run", action="store_true", help="발송하지 않고 대상만 출력")
    ap.add_argument("--sync-lims", action="store_true", help="LIMS 결과 증분 동기화 후 감시")
    args = ap.parse_args(argv)

    if args.sync_lims:
        cfg = load_lims_config()
        if not cfg.enabled:
            print("[LIMS] 연동이 꺼져 있습니다(🔗 LIMS 연동 화면에서 사용 설정).")
        else:
            rep = sync(cfg, dry_run=args.dry_run)
            print(f"[LIMS] {rep.summary_text()}")
            for code in rep.unmapped[:10]:
                print(f"    미매핑 시험코드: {code[0]}/{code[1]} {code[2]}건")

    if args.import_dir:
        for line in import_folder(args.import_dir):
            print(line)

    if args.sync_file:
        for line in sync_file(args.sync_file):
            print(line)

    store = load_store()
    if store.is_empty():
        print("데이터가 없습니다.")
        return 0
    registry, settings = load_registry(), load_settings()
    attach_predictions(store, registry)
    events = detect_events(store, registry, settings)
    period = store.period()
    since = period[1] - pd.Timedelta(hours=args.since_hours)
    cfg = load_config()
    results = dispatch(events, load_alert_status(), cfg, diagnose_fn=lambda e: diagnose(e, store, registry, settings),
                       since=since, dry_run=args.dry_run)
    print(f"이벤트 {len(events)}건 중 발송 대상 {len(results)}건 (기준 {cfg.min_severity} 이상, {since:%Y-%m-%d %H:%M} 이후)")
    for r in results:
        print("-", r["subject"])
        for ch, ok, msg in r["channels"]:
            print(f"    {ch}: {'성공' if ok else '실패'} - {msg}")
        if r["sent"]:
            update_alert_status(r["event_id"], notified=1)
    if not args.dry_run and results and not (cfg.email_enabled or cfg.webhook_enabled):
        print("※ 이메일·웹훅이 모두 꺼져 있어 실제 발송은 되지 않았습니다(⚙️ 기준·알림 설정).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
