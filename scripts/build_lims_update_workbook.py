"""방법 A용 업데이트 양식(엑셀) 생성 — LIMS에서 내려받은 결과를 시트별로 붙여넣어 지정 폴더에 두면
QMS(5_MONITOR_TASK / 📥 데이터 관리)가 자동으로 읽어 반영한다.

한 파일 안의 시트:
  사용법 · 원료 · 클링커 · 시멘트 · 물성 · 클링커일일 · 출하시멘트 · LIMS매핑
머리글은 LIMS에서 쓰는 한글 이름 그대로라 QMS가 알아서 내부 항목으로 연결한다(qms.store.COL_ALIASES).
공정 데이터(원료·클링커·시멘트·물성·클링커일일) 시트는 '채취일' 1건 = 1행.

사용:  python scripts/build_lims_update_workbook.py [출력경로]
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "Blue365_QMS_업데이트_양식.xlsx"

NAVY = "1F3A5F"
HEAD = PatternFill("solid", fgColor=NAVY)
HEADF = Font(color="FFFFFF", bold=True, size=10)
NOTEF = Font(color="7A7A7A", italic=True, size=9)
TITLEF = Font(bold=True, size=14, color=NAVY)
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _lims_rows() -> list[dict]:
    p = ROOT.parent / "scratchpad_lims_rows.json"
    cand = [Path("/tmp/claude-0/-home-user-Blue365/a64ee9b0-81fc-5d30-b033-f3935f2bf740/scratchpad/lims_rows.json"), p]
    for c in cand:
        if c.exists():
            return json.loads(c.read_text(encoding="utf-8"))
    return []


def _d(v):
    return str(v)[:10] if v else ""


def _header(ws, headers, row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, HEADF, CENTER, BORDER
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 3)
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(16, len(str(h)) + 3))


def _note(ws, text, row, span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row, 1, text)
    cell.font, cell.alignment = NOTEF, LEFT


def _rows(ws, data, start=2):
    for r, rec in enumerate(data, start):
        for c, v in enumerate(rec, 1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            cell.alignment = CENTER if c > 1 else Alignment(horizontal="center", vertical="center")


def build():
    rows = _lims_rows()
    wb = Workbook()

    # ── 사용법 ──
    ws = wb.active
    ws.title = "사용법"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 112
    ws["B2"] = "Blue365 QMS — LIMS 결과 업데이트 양식 (방법 A: 폴더 자동 반영)"
    ws["B2"].font = TITLEF
    guide = [
        "",
        "■ 쓰는 순서",
        "  1) LIMS(품질관리)에서 공정별 결과를 엑셀로 내려받습니다 — 원료, 클링커, 시멘트(분쇄), 제품(출하).",
        "  2) 내려받은 값을 이 파일의 '맞는 시트'에 붙여넣습니다(머리글은 그대로 두고, 그 아래 데이터만).",
        "  3) 이 파일을 지정 폴더(예: C:\\Blue365_QMS\\inbox)에 저장합니다. 파일 이름은 바꿔도 됩니다.",
        "  4) QMS가 정해진 시각(작업 스케줄러)에 자동으로 읽어 반영합니다. 즉시 반영하려면 [📥 데이터 관리]에서 올려도 됩니다.",
        "",
        "■ 시트와 화면의 관계 (한 번 올리면 여러 화면이 같이 채워집니다)",
        "  · 원료 / 클링커 / 시멘트 / 물성 / 클링커일일  →  공정 모니터링 · 그래프 보기 · SPC · 원인 진단 · 강도 예측 · 배합 설계 · 6가크롬",
        "  · 출하시멘트  →  출하시멘트 화면 (선박·양하지 포함)",
        "  ※ '그래프 보기'·'공정 모니터링'·'SPC'는 별도 시트가 아니라, 위 공정 데이터를 보는 화면입니다.",
        "",
        "■ 꼭 지킬 것",
        "  · '채취일' 열은 비우지 마세요(이 값이 1행=1건의 기준입니다). YYYY-MM-DD 형식.",
        "  · '시멘트'·'물성'·'출하시멘트' 시트는 '품명'(1종/3종 등)이 필요합니다. LIMS '1종시멘트 벌크'는 자동으로 '1종'으로 인식합니다.",
        "  · 같은 날짜(+품종)가 이미 있으면 '값이 있는 칸'만 덮어씁니다. 늦게 나온 28일 강도를 나중에 채워도 기존 값은 지워지지 않습니다.",
        "  · 빈 칸은 그대로 두세요(0을 넣지 마세요 — 0은 '측정값 0'으로 처리됩니다).",
        "  · 머리글 이름은 바꾸지 마세요. LIMS 이름 그대로여야 QMS가 자동으로 연결합니다('LIMS매핑' 시트 참고).",
        "",
        "■ 데이터가 없는 공정은?",
        "  · 해당 시트를 비워 두거나 삭제해도 됩니다. 있는 시트만 반영됩니다.",
        "",
        "※ 아래 각 시트에는 이해를 돕기 위한 예시 행이 들어 있습니다. '시멘트·물성·출하시멘트'는 실제로 주신 LIMS 샘플 값이고,",
        "   '원료·클링커·클링커일일'은 형태를 보여주기 위한 예시(추정)입니다. 실제 값으로 바꿔 쓰세요.",
    ]
    for i, line in enumerate(guide, 4):
        ws.cell(i, 2, line)
        if line.startswith("■"):
            ws.cell(i, 2).font = Font(bold=True, size=11, color=NAVY)

    # ── 원료 (raw_meal) — 예시(추정) ──
    ws = wb.create_sheet("원료")
    h = ["채취일", "CaO", "SiO2", "Al2O3", "Fe2O3", "MgO", "SO3(XRF)", "K2O", "Na2O", "90㎛", "200㎛", "LOI"]
    _header(ws, h)
    base = datetime(2026, 10, 1)  # noqa: DTZ001 - 예시 날짜(시간대 불필요)
    ex = [[_d(base + timedelta(days=i)), 43.1 + i * 0.05, 13.6, 3.3, 2.1, 1.4, 0.3, 0.9, 0.12, 14.5, 1.2, 35.2]
          for i in range(5)]
    _rows(ws, ex)
    _note(ws, "※ 예시(추정) — 원료(킬른 피드) XRF. 실제 LIMS 원료 결과로 바꿔 쓰세요. 90㎛·200㎛는 체 잔사(%).", len(ex) + 3, len(h))

    # ── 클링커 (clinker) — 예시(추정) ──
    ws = wb.create_sheet("클링커")
    h = ["채취일", "CaO", "SiO2", "Al2O3", "Fe2O3", "MgO", "SO3(XRF)", "K2O", "Na2O", "F-CaO(습식)", "용적비중"]
    _header(ws, h)
    ex = [[_d(base + timedelta(days=i)), 65.8, 21.3, 5.4, 3.3, 1.5, 0.9, 0.85, 0.18, 1.1 + i * 0.05, 1320]
          for i in range(5)]
    _rows(ws, ex)
    _note(ws, "※ 예시(추정) — 클링커 XRF + 자유석회(F-CaO)·용적비중(litre weight). 실제 LIMS 클링커 결과로 바꿔 쓰세요.",
          len(ex) + 3, len(h))

    # ── 시멘트 (cement) — 실제 LIMS 샘플 값 ──
    ws = wb.create_sheet("시멘트")
    h = ["채취일", "품명", "Blaine", "45㎛", "SO3(XRF)", "LOI", "MgO", "석회석비율"]
    _header(ws, h)
    data = [[_d(r.get("채취일")), r.get("품명"), r.get("Blaine"), r.get("45㎛"), r.get("SO3(XRF)"),
             r.get("LOI"), r.get("MgO"), r.get("석회석비율")] for r in rows] or \
           [[_d(base), "1종", 3810, 6.6, 2.58, 1.95, 3.81, 3]]
    _rows(ws, data)
    _note(ws, "※ 주신 LIMS 샘플(1종시멘트 벌크)의 실제 값입니다. '품명'은 자동으로 '1종'으로 인식됩니다.",
          len(data) + 3, len(h))

    # ── 물성 (physical) — 실제 LIMS 샘플 값 ──
    ws = wb.create_sheet("물성")
    h = ["채취일", "품명", "초결시간(분)", "종결시간(HH:MM)", "Autoclave", "압축강도1일(MPa)", "압축강도3일(MPa)",
         "압축강도7일(MPa)", "압축강도28일(MPa)", "6가크롬", "담당자"]
    _header(ws, h)
    data = [[_d(r.get("채취일")), r.get("품명"), r.get("초결시간(분)"), r.get("종결시간(HH:MM)"), r.get("Autoclave"),
             r.get("압축강도1일(MPa)"), r.get("압축강도3일(MPa)"), r.get("압축강도7일(MPa)"), r.get("압축강도28일(MPa)"),
             r.get("6가크롬"), r.get("담당자") or ""] for r in rows] or \
           [[_d(base), "1종", 262, "06:33", 0.35, 15.8, 38.0, 46.4, 57.1, 0.0, ""]]
    _rows(ws, data)
    _note(ws, "※ 주신 LIMS 샘플의 실제 값입니다. 종결시간은 'HH:MM'(예 06:33)·분 모두 인식합니다. 6가크롬 단위는 LIMS 그대로.",
          len(data) + 3, len(h))

    # ── 클링커일일 (xrd) — 예시(추정) ──
    ws = wb.create_sheet("클링커일일")
    h = ["채취일", "Alite(C3S)", "Belite(C2S)", "C3A", "C4AF", "F-CaO", "Periclase", "총크롬", "6가크롬"]
    _header(ws, h)
    ex = [[_d(base + timedelta(days=i)), 62.0, 14.5, 7.8, 10.1, 1.0, 1.2, 85, 2.1] for i in range(5)]
    _rows(ws, ex)
    _note(ws, "※ 예시(추정) — 클링커 XRD 광물상(%) + 총크롬·6가크롬(mg/kg). XRD·크롬 분석을 하는 경우만 채우세요.",
          len(ex) + 3, len(h))

    # ── 출하시멘트 — 실제 LIMS 샘플 값 ──
    ws = wb.create_sheet("출하시멘트")
    h = ["채취일", "품명", "첨가재명칭", "석회석비율", "첨가재비율", "LOI", "Al2O3", "MgO", "SO3(XRF)", "Cl(XRF)",
         "TA", "압축강도1일(MPa)", "압축강도3일(MPa)", "압축강도7일(MPa)", "압축강도28일(MPa)", "45㎛", "Blaine",
         "주도", "Flow", "종결시간(HH:MM)", "비고"]
    _header(ws, h)
    data = [[_d(r.get("채취일")), r.get("품명"), r.get("첨가재명칭"), r.get("석회석비율"), r.get("첨가재비율"),
             r.get("LOI"), r.get("Al2O3"), r.get("MgO"), r.get("SO3(XRF)"), r.get("Cl(XRF)"), r.get("TA"),
             r.get("압축강도1일(MPa)"), r.get("압축강도3일(MPa)"), r.get("압축강도7일(MPa)"), r.get("압축강도28일(MPa)"),
             r.get("45㎛"), r.get("Blaine"), r.get("주도"), r.get("Flow"), r.get("종결시간(HH:MM)"),
             r.get("비고") or ""] for r in rows]
    _rows(ws, data or [[_d(base), "1종", "플라이 애시", 3, 5, 1.95, 5.99, 3.81, 2.58, 605, 0.92, 15.8, 38.0,
                        46.4, 57.1, 6.6, 3810, 27.5, 95, "06:33", "○○호"]])
    _note(ws, "※ 주신 LIMS 샘플의 실제 값입니다. '비고'에 출항 선박명을 적으면 출하시멘트 화면에서 선박·양하지로 인식합니다.",
          len(data) + 3, len(h))

    # ── LIMS매핑 ──
    ws = wb.create_sheet("LIMS매핑")
    _header(ws, ["이 양식의 시트", "머리글(LIMS 이름)", "QMS 항목", "비고"])
    m = [
        ["원료", "채취일 / CaO·SiO2·Al2O3·Fe2O3·MgO·SO3(XRF)·K2O·Na2O / 90㎛·200㎛ / LOI", "원료(킬른 피드) 화학·분말도", "LIMS 원료 결과"],
        ["클링커", "채취일 / 오산화물 / F-CaO(습식) / 용적비중", "클링커 화학·자유석회", "LIMS 클링커 결과"],
        ["시멘트", "채취일·품명 / Blaine·45㎛·SO3(XRF)·LOI·MgO·석회석비율", "시멘트(밀 출구)", "품명=1종/3종 자동 인식"],
        ["물성", "채취일·품명 / 초결·종결 / Autoclave / 압축강도 1·3·7·28일 / 6가크롬 / 담당자", "물성·강도(로트)", "종결 HH:MM 가능"],
        ["클링커일일", "채취일 / XRD 광물상 / 총크롬·6가크롬", "클링커 일일(XRD·크롬)", "분석 시에만"],
        ["출하시멘트", "채취일·품명 / 첨가재·화학 / 강도 / 분말도 / 주도·Flow / 비고(선박명)", "출하시멘트", "비고=출항 선박"],
    ]
    _rows(ws, m)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 20

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("저장:", OUT)
    return OUT


if __name__ == "__main__":
    build()
