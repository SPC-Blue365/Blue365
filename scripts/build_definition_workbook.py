"""관리항목·기준 정의서(엑셀) 생성 — 시스템 설정(qms.standards, qms.knowledge)에서 직접 만든다.

사용:  python scripts/build_definition_workbook.py [출력경로]
기본 출력: docs/QMS_관리항목_기준정의서.xlsx
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from qms.knowledge import KB  # noqa: E402
from qms.spc import RULE_NAMES  # noqa: E402
from qms.standards import DEFAULT_SETTINGS, STAGES, TABLES, SpecRegistry  # noqa: E402
from qms.store import RAW_COLUMNS, SHEET_NAMES  # noqa: E402

FONT = "맑은 고딕"  # Arial에는 한글 글리프가 없어 한글 표준 글꼴 사용
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True)
F_SUB = Font(name=FONT, size=10, color="5F6B73")
F_KS = Font(name=FONT, size=10, color="0000FF")          # 규격값(외부 근거 하드코딩)
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="37474F")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")      # 공장이 입력할 칸
FILL_INIT = PatternFill("solid", fgColor="ECEFF1")       # 시스템 초기값(경험칙)
FILL_KEY = PatternFill("solid", fgColor="FFF3E0")
THIN = Side(style="thin", color="CFD8DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SOURCE_BY_TABLE = {"raw_meal": "XRF(시험실/온라인)·잔사 시험", "kiln": "DCS(킬른 운전)", "clinker": "XRF·f-CaO·리터중량 시험",
                   "cement": "XRF·분말도 시험 + DCS(밀 운전)", "physical": "물성 시험실(LIMS)",
                   "xrd": "클링커 일일 시료(XRD·크롬, LIMS)"}
METHOD = {"cem_blaine": "KS L 5106(블레인)", "phy_ist": "KS L 5108(비카)", "phy_fst": "KS L 5108(비카)",
          "phy_autoclave": "KS L 5107(오토클레이브)", "phy_s1": "KS L ISO 679", "phy_s3": "KS L ISO 679",
          "phy_s7": "KS L ISO 679", "phy_s28": "KS L ISO 679", "cem_so3": "KS L 5120 / XRF", "cem_loi": "KS L 5120",
          "cem_mgo": "KS L 5120 / XRF", "lab_temp": "KS L ISO 679", "lab_rh": "KS L ISO 679",
          "lab_cure_temp": "KS L ISO 679", "phy_crvi": "KS L 5221(수용성 6가크롬)", "clk_crvi": "KS L 5221 준용",
          "clk_cr": "XRF / ICP(총 크롬)", "rm_loi": "KS L 5120(강열감량)", "cem_ls": "CO₂·강열감량 역산 / 공급기 실적",
          **{k: "XRD-Rietveld 정량" for k in ("xrd_alite", "xrd_belite", "xrd_c3a", "xrd_c4af", "xrd_fcao",
                                                "xrd_periclase")}}


def header(ws, row: int, cols: list[str], widths: list[float] | None = None) -> None:
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font, cell.fill, cell.border, cell.alignment = F_HEAD, FILL_HEAD, BORDER, CENTER
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[row].height = 30


def put(ws, row: int, col: int, value, font=F_BODY, fill=None, align=WRAP, fmt: str | None = None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.border, cell.alignment = font, BORDER, align
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def build(path: Path) -> None:
    reg = SpecRegistry()
    wb = Workbook()

    # ── 1. 안내 ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "시멘트 통합 품질관리 시스템(QMS) — 관리항목·기준 정의서"
    ws["A1"].font = F_TITLE
    ws["A2"] = f"작성: 품질관리팀 · 생성일 {date.today():%Y-%m-%d} · 시스템 설정(qms/standards.py, qms/knowledge.py)에서 자동 생성"
    ws["A2"].font = F_SUB
    rows = [
        ("목적", "QMS가 감시·알림·진단에 사용하는 관리항목, KS 규격, 사내 관리기준, 알림 규칙, 원인진단 지식베이스를 한 문서로 정의한다. "
                "공장 실제 기준으로 확정한 값을 '관리항목' 시트의 노란 칸에 입력하면 시스템 설정(⚙️ 기준·알림 설정)에 반영한다."),
        ("작성 방법(입력 칸)", "'관리항목' 시트의 노란색 칸(사내 하한·상한·목표 '확정값')만 입력한다. 비워 두면 회색 칸의 시스템 초기값을 계속 사용한다."),
        ("범례 — 노란 칸", "공장이 입력할 확정값"),
        ("범례 — 회색 칸", "시스템 초기값: 업계 경험칙 기반(추정). 공장 데이터·실험으로 검증 후 확정 필요"),
        ("범례 — 파란 글씨", "KS 규격값(외부 근거). 수정 대상 아님 — 근거는 'KS규격_근거' 시트"),
        ("작성 예시", "예) 클링커 자유석회(f-CaO): 확정 하한 0.5 / 확정 상한 1.6 / 확정 목표 1.1  ← 숫자만 입력(단위 제외)"),
        ("사실/추정 구분", "KS 값 = 규격(사실, 원문 대조 권장) · 사내 초기값·지식베이스 수치 = 경험칙(추정) · 진단 판정 '데이터 지지' = 측정 데이터로 확인된 변화(사실)"),
        ("시트 구성", "요약 · 관리항목 · 알림규칙 · SPC판정규칙 · 원인진단_지식베이스 · 데이터입력_항목 · KS규격_근거 · 구축로드맵"),
    ]
    for i, (k, v) in enumerate(rows, 4):
        put(ws, i, 1, k, F_BOLD)
        c = put(ws, i, 2, v)
        if k.endswith("노란 칸"):
            c.fill = FILL_INPUT
        elif k.endswith("회색 칸"):
            c.fill = FILL_INIT
        elif k.endswith("파란 글씨"):
            c.font = F_KS
        ws.row_dimensions[i].height = 16 * (len(v) // 58 + 1) + 6

    # ── 3. 관리항목 (요약보다 먼저 만들어 행 범위를 알 수 있게) ────────────
    wm = wb.create_sheet("관리항목")
    cols = ["No", "공정", "코드", "항목", "단위", "측정 주기", "데이터 출처", "시험방법", "품종", "KS 하한", "KS 상한",
            "KS 근거", "사내 하한(초기)", "사내 상한(초기)", "목표(초기)", "초기값 근거", "사내 하한(확정)", "사내 상한(확정)",
            "목표(확정)", "SPC 판정규칙", "핵심관리", "설명"]
    widths = [5, 12, 13, 24, 7, 9, 20, 16, 7, 8, 8, 26, 10, 10, 9, 22, 10, 10, 9, 14, 8, 44]
    header(wm, 1, cols, widths)
    r = 2
    no = 0
    for stage in STAGES:
        for it in reg.by_stage(stage):
            items = list(it.limits.items()) or [("*", None)]
            for prod, lim in items:
                no += 1
                vals = [no, it.stage, it.key, it.name, it.unit, TABLES[it.table]["freq"], SOURCE_BY_TABLE[it.table],
                        METHOD.get(it.key, "-"), "전 품종" if prod == "*" else prod]
                for j, v in enumerate(vals, 1):
                    put(wm, r, j, v)
                put(wm, r, 10, getattr(lim, "ks_min", None), F_KS, align=CENTER)
                put(wm, r, 11, getattr(lim, "ks_max", None), F_KS, align=CENTER)
                put(wm, r, 12, it.ks_ref or "-")
                put(wm, r, 13, getattr(lim, "lsl", None), fill=FILL_INIT, align=CENTER)
                put(wm, r, 14, getattr(lim, "usl", None), fill=FILL_INIT, align=CENTER)
                put(wm, r, 15, getattr(lim, "target", None), fill=FILL_INIT, align=CENTER)
                put(wm, r, 16, it.basis)
                for j in (17, 18, 19):
                    put(wm, r, j, None, fill=FILL_INPUT, align=CENTER)
                put(wm, r, 20, ", ".join(it.rules) if it.rules else "-")
                put(wm, r, 21, "●" if it.key_item else "", align=CENTER)
                put(wm, r, 22, it.description or ("예측값 — 실측 전 조기 경보" if it.prediction else ""))
                if it.key_item:
                    wm.cell(row=r, column=4).fill = FILL_KEY
                r += 1
    last = r - 1
    wm.freeze_panes = "E2"
    wm.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    wm["Q1"].comment = Comment("공장 확정값 입력 칸(노란색). 비우면 초기값(회색) 사용.", "QMS")
    dv = DataValidation(type="decimal", allow_blank=True, operator="between", formula1="-100000", formula2="100000",
                        error="숫자만 입력하세요(단위 제외).", errorTitle="입력 오류")
    wm.add_data_validation(dv)
    dv.add(f"Q2:S{last}")

    # ── 2. 요약(수식) ──────────────────────────────────────────────────
    wsum = wb.create_sheet("요약", 1)
    wsum["A1"] = "관리항목 요약 — 공정별 현황 (수식으로 '관리항목' 시트에서 자동 집계)"
    wsum["A1"].font = Font(name=FONT, size=13, bold=True)
    header(wsum, 3, ["공정", "관리항목 수(행)", "KS 기준 보유", "핵심관리항목", "SPC 적용", "확정값 입력 완료", "확정 진행률"],
           [16, 14, 13, 13, 11, 15, 12])
    rng = lambda col: f"관리항목!${col}$2:${col}${last}"  # noqa: E731
    for i, stage in enumerate(STAGES, 4):
        put(wsum, i, 1, stage, F_BOLD)
        put(wsum, i, 2, f'=COUNTIF({rng("B")},A{i})', align=CENTER)
        put(wsum, i, 3, f'=COUNTIFS({rng("B")},A{i},{rng("J")},"<>")+COUNTIFS({rng("B")},A{i},{rng("J")},"",{rng("K")},"<>")',
            align=CENTER)
        put(wsum, i, 4, f'=COUNTIFS({rng("B")},A{i},{rng("U")},"●")', align=CENTER)
        put(wsum, i, 5, f'=COUNTIFS({rng("B")},A{i},{rng("T")},"<>-")', align=CENTER)
        put(wsum, i, 6, f'=SUMPRODUCT(({rng("B")}=A{i})*((LEN({rng("Q")})+LEN({rng("R")})+LEN({rng("S")}))>0))',
            align=CENTER)
        put(wsum, i, 7, f"=IFERROR(F{i}/B{i},0)", align=CENTER, fmt="0.0%")
    tr = 4 + len(STAGES)
    put(wsum, tr, 1, "합계", F_BOLD)
    for col in "BCDEF":
        put(wsum, tr, " BCDEF".index(col) + 1, f"=SUM({col}4:{col}{tr - 1})", F_BOLD, align=CENTER)
    put(wsum, tr, 7, f"=IFERROR(F{tr}/B{tr},0)", F_BOLD, align=CENTER, fmt="0.0%")
    wsum.cell(row=tr + 2, column=1, value="※ '관리항목 수(행)'은 품종별 기준을 따로 세므로 항목 수보다 많을 수 있음. "
                                         "KS 기준 보유 = KS 하한 또는 상한이 있는 행.").font = F_SUB

    # ── 4. 알림규칙 ────────────────────────────────────────────────────
    wa = wb.create_sheet("알림규칙")
    header(wa, 1, ["심각도", "판정 조건", "예시", "통보 대상(제안)", "대응 기한(제안)", "발송"], [10, 46, 46, 24, 18, 22])
    alert_rows = [
        ("🔴 위험", "제품 항목이 KS L 5201 규격을 벗어남", "시멘트 SO₃ 3.5% 초과(1종), 28일 강도 42.5 MPa 미만",
         "품질팀장·생산팀장·공장장", "즉시(1시간 내) 출하 판정", "메일+메신저 즉시"),
        ("🟠 경고", "사내 관리기준 이탈 / 시험조건(KS L ISO 679) 이탈 / 28일 예측치 사내기준 미달",
         "f-CaO 1.8% 초과, 분말도 3,250 미만, 양생수 19~21 ℃ 이탈", "품질 담당·해당 공정 담당", "당일(8시간 내) 원인 조사",
         "메일+메신저"),
        ("🟡 주의", "규격 이내이나 SPC 판정규칙 위반(8시간 평균 기준 R1·R2·R3 등)", "생료 LSF 9연속 중심선 위(평균 이동)",
         "품질 담당", "3일 내 점검", "대시보드 표시(기본 미발송)"),
        ("공통", "같은 항목·품종에서 12시간 이내(물성 3일 이내) 이어지는 위반점은 1건으로 묶음. 규격 이탈과 겹치는 SPC 패턴은 '동반 패턴'으로 기록",
         "알림 피로도 감소", "-", "-", "-"),
        ("예측", "예측값은 실측이 아니므로 최고 심각도를 '경고'로 제한", "28일 강도 예측 46 MPa → 경고(예측)", "-", "-", "-"),
    ]
    for i, row in enumerate(alert_rows, 2):
        for j, v in enumerate(row, 1):
            put(wa, i, j, v)
        wa.row_dimensions[i].height = 44

    # ── 5. SPC 판정규칙 ────────────────────────────────────────────────
    wsp = wb.create_sheet("SPC판정규칙")
    header(wsp, 1, ["코드", "규칙", "의미", "기본 적용"], [8, 34, 56, 36])
    meaning = {"R1": "돌발 이상(한 점이 관리한계 밖)", "R2": "공정 평균의 이동(shift)", "R3": "점진적 추세(drift)",
               "R5": "중간 크기의 평균 이동 조기 감지", "R6": "작은 평균 이동 감지(허위경보 증가 주의)"}
    for i, (code, name) in enumerate(RULE_NAMES.items(), 2):
        put(wsp, i, 1, code, align=CENTER)
        put(wsp, i, 2, name.format(run=DEFAULT_SETTINGS["run_length"], trend=DEFAULT_SETTINGS["trend_length"]))
        put(wsp, i, 3, meaning[code])
        put(wsp, i, 4, "활성(기본)" if code in DEFAULT_SETTINGS["enabled_rules"] else "선택(설정 화면에서 활성화)")
    n = 2 + len(RULE_NAMES)
    notes = ["관리한계: 기준기간(기본 데이터 시작 30일) 평균 ± 3σ, σ = 기준기간 표준편차(자기상관 고려), 기준기간 내 3σ 밖 점은 1회 제외",
             "2시간·1시간 데이터는 8시간(근무조) 평균에 판정규칙 적용 — 원 데이터에 런 규칙을 적용하면 허위 경보 과다",
             "규격·관리기준 이탈은 개별 측정값 기준으로 판정",
             "공정능력: Cpk = 군내변동(MR̄/1.128), Ppk = 전체 표준편차 — 관용 기준 1.33 이상 충분, 1.0 미만 개선 필요"]
    for k, t in enumerate(notes, n + 1):
        wsp.cell(row=k, column=1, value="※ " + t).font = F_SUB

    # ── 6. 원인진단 지식베이스 ─────────────────────────────────────────
    wk = wb.create_sheet("원인진단_지식베이스")
    kcols = ["현상(항목·방향)", "축", "원인 가설", "메커니즘", "자동 평가 데이터", "확인 방법(데이터·시험)", "단기 조치", "근본 대책",
             "근거 수준", "담당(제안)", "사전 가중치"]
    header(wk, 1, kcols, [24, 8, 26, 50, 30, 44, 36, 36, 11, 13, 9])
    r = 2
    seen = set()
    for ph in KB.values():
        sig = (ph.title, tuple(h.title for h in ph.hypotheses))
        if sig in seen:
            continue
        seen.add(sig)
        for h in ph.hypotheses:
            checks = ", ".join(
                {"__resid": "28일 실측−조기강도 예측 잔차", "__isolated": "이벤트 단발성"}.get(c.item, reg[c.item].name if c.item in reg else c.item)
                + {"high": "↑", "low": "↓", "var": " 변동↑", "dev": " 목표편차↑", "out": " 기준 이탈", "change": " 변경"}[c.expect]
                + ("(반대)" if c.invert else "") for c in h.checks) or "없음(현장 확인)"
            vals = [ph.title, h.axis, h.title, h.mechanism, checks, " / ".join(h.verify), " / ".join(h.short_term),
                    " / ".join(h.root_cause), h.basis, h.owner, h.prior]
            for j, v in enumerate(vals, 1):
                put(wk, r, j, v)
            r += 1
    wk.freeze_panes = "D2"
    wk.auto_filter.ref = f"A1:K{r - 1}"
    wk.cell(row=r + 1, column=1, value="※ 메커니즘의 정량 수치(예: Blaine 100 cm²/g 당 강도 변화)는 업계 경험칙(추정)이며 공장 데이터로 검증 후 수정할 것. "
                                       "사전 가중치 = 경험상 발생 빈도(1.0 보통).").font = F_SUB

    # ── 7. 데이터 입력 항목 ────────────────────────────────────────────
    wd = wb.create_sheet("데이터입력_항목")
    header(wd, 1, ["엑셀 시트", "열 코드", "설명", "단위", "측정 주기", "필수"], [12, 16, 34, 9, 10, 8])
    r = 2
    for name, cols_ in RAW_COLUMNS.items():
        for c in ["timestamp"] + cols_:
            it = reg.get(c)
            desc = it.name if it else {"timestamp": "측정일시(예: 2026-10-07 14:00)", "product": "품종(1종/3종)",
                                       "sand_lot": "표준사 Lot", "operator": "시험 조"}.get(c, c)
            req = "필수" if c in ("timestamp", "product") else "선택"
            for j, v in enumerate([SHEET_NAMES[name], c, desc, it.unit if it else "-", TABLES[name]["freq"], req], 1):
                put(wd, r, j, v)
            r += 1
    wd.cell(row=r + 1, column=1, value="※ LSF·SM·IM·C₃S·액상량 등 파생값은 시스템이 자동 계산하므로 입력하지 않음. "
                                       "같은 측정일시(+품종)는 값이 있는 칸만 덮어씀(빈 칸은 기존 값 유지 — LIMS 결과 순차 반영)."
            ).font = F_SUB

    # ── 8. KS 규격 근거 ────────────────────────────────────────────────
    wks = wb.create_sheet("KS규격_근거")
    header(wks, 1, ["구분", "1종(보통)", "3종(조강)", "근거", "확인 상태"], [22, 22, 26, 34, 34])
    ks_rows = [
        ("압축강도 1일", "-", "10.0 MPa 이상", "KS L 5201", "웹 검색 확인(2026-10)"),
        ("압축강도 3일", "12.5 MPa 이상", "20.0 MPa 이상", "KS L 5201", "웹 검색 확인(2026-10)"),
        ("압축강도 7일", "22.5 MPa 이상", "32.5 MPa 이상", "KS L 5201", "웹 검색 확인(2026-10)"),
        ("압축강도 28일", "42.5 MPa 이상", "47.5 MPa 이상", "KS L 5201", "웹 검색 확인(2026-10)"),
        ("분말도(비표면적)", "2,800 cm²/g 이상", "3,300 cm²/g 이상", "KS L 5201 (시험: KS L 5106)", "웹 검색 확인(2026-10)"),
        ("응결시간(비카)", "초결 60분 이상, 종결 10시간 이하", "1종과 동일 적용", "KS L 5201 (시험: KS L 5108)",
         "1종 확인 · 3종 원문 대조 필요"),
        ("안정도(오토클레이브 팽창도)", "0.8% 이하", "0.8% 이하", "KS L 5201 (시험: KS L 5107)", "웹 검색 확인(2026-10)"),
        ("산화마그네슘(MgO)", "5.0% 이하", "5.0% 이하", "KS L 5201", "웹 검색 확인(2026-10)"),
        ("삼산화황(SO₃)", "3.5% 이하", "4.5% 이하", "KS L 5201", "웹 검색 확인(2026-10)"),
        ("강열감량", "5.0% 이하", "5.0% 이하", "KS L 5201 (2016 개정 3.0→5.0%)", "웹 검색 확인(2026-10)"),
        ("강도 시험조건", "시험실 20±2 ℃·RH 50% 이상, 양생수 20±1 ℃, 재하 2,400±200 N/s", "동일", "KS L ISO 679 (ISO 679)",
         "ISO 679 규정값 적용 — KS 원문 대조 필요"),
        ("수용성 6가크롬(참고)", "20 mg/kg 이하", "20 mg/kg 이하", "환경부–시멘트업계 자율협약(시험: KS L 5221) — KS L 5201 아님",
         "웹 검색 확인(2026-10) · 협약 원문 확인 권장"),
        ("수용성 6가크롬(EU 참고)", "2 mg/kg 이하", "2 mg/kg 이하", "EU REACH 부속서 XVII 47항(시험: EN 196-10)",
         "시험법이 달라 국내 수치와 직접 비교 불가"),
    ]
    for i, row in enumerate(ks_rows, 2):
        for j, v in enumerate(row, 1):
            put(wks, i, j, v, F_KS if j in (2, 3) else F_BODY)
    k = len(ks_rows) + 3
    for t in ["출처: e나라표준인증 KS L 5201 https://standard.go.kr/KSCI/standardIntro/getStandardSearchView.do?ksNo=KSL5201",
              "출처: 한국시멘트협회 KS 규격 요약 http://www.cement.or.kr/tech_2014/standard.asp?sm=3_6_1",
              "※ 2종(중용열)·4종(저열)·5종(내황산염) 기준은 원문 확인 후 추가. KS는 개정될 수 있으므로 연 1회 최신판 대조 권장."]:
        wks.cell(row=k, column=1, value=t).font = F_SUB
        k += 1

    # ── 9. 구축 로드맵 ─────────────────────────────────────────────────
    wr = wb.create_sheet("구축로드맵")
    header(wr, 1, ["단계", "기간(추정)", "주요 내용", "산출물", "필요 자원(추정)", "상태"], [16, 12, 56, 36, 32, 12])
    road = [
        ("1단계 MVP", "완료", "모니터링·3단계 기준·알림·5축 원인진단·28일 예측·보고서(PPT/엑셀)·엑셀 업로드·데모 데이터",
         "Streamlit 시스템, 정의서, 구축 계획서", "품질 담당 1명 + 사내 PC 1대", "완료"),
        ("2단계 확장 기능", "완료", "원료 배합 최적화(LSF·SM·IM·C₃S, 원가 최소)·클링커 계수·XRD 광물 예측, 재령별 강도·응결 예측과 "
         "제어 솔루션, 시멘트 6가크롬 물질수지·환원제 투입량, LIMS 연동(SQL·REST·파일), AI(Claude) 솔루션 보고서",
         "배합 설계서·6가크롬 평가표(엑셀), 기능 확장 보고(PPT)", "품질 담당 1명", "개발 완료(실데이터 검증 필요)"),
        ("3단계 실데이터 시범", "약 2~3개월", "LIMS 실연결(읽기 전용 계정)·실데이터 3개월 적재 → 사내 기준·예측 모델·킬른 전환율·환원제 과잉계수 보정, "
         "알림 채널 연결, AI 보고서 보안 승인 후 시범 사용", "확정 관리기준, 시범 운영 보고", "IT·LIMS 담당 협조, 정보보안 승인", "예정"),
        ("4단계 고도화", "약 3~6개월", "조치 이력 기반 지식베이스 학습, 입도분포(PSD) 반영 강도 모델, 원료 품위 예측 연계 배합 자동 보정",
         "고도화 기능, 효과 분석", "품질·생산·설비 협업", "계획"),
        ("5단계 확산", "추후", "DB/OPC-UA 실시간 연계, 모바일 알림, 타 공정·타 공장 확산", "운영 표준·교육 자료", "IT 투자 검토", "계획"),
    ]
    for i, row in enumerate(road, 2):
        for j, v in enumerate(row, 1):
            put(wr, i, j, v)
        wr.row_dimensions[i].height = 48
    wr.cell(row=len(road) + 3, column=1, value="※ 기간·자원은 일반적인 시범 구축 경험에 따른 추정치이며, 데이터 연계 범위에 따라 달라짐.").font = F_SUB

    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 90
    wb.calculation.fullCalcOnLoad = True  # 엑셀에서 열 때 수식 재계산
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "QMS_관리항목_기준정의서.xlsx"
    build(out)
    print(out)
