"""사내망(오프라인) Windows 설치 패키지 빌드.

휴대용 Python(python.org 공식 NuGet 배포본) + 구성요소 사전 설치 + 프로그램 + 더블클릭 배치 파일을
분할 ZIP으로 만든다. 설치 PC에는 인터넷·관리자 권한·Python 설치가 필요 없다.

사용(빌드 PC: 인터넷 + uv 필요)
  python scripts/build_windows_package.py                          # dist/windows/ 에 분할 ZIP(파일당 최대 95MB)
  python scripts/build_windows_package.py --max-part-mb 0          # 한 파일로 생성
  python scripts/build_windows_package.py --relock                 # 잠금 파일 갱신(constraints-tested.txt 버전 고정)

결과(dist/windows/)
  Blue365_QMS_win64_v<날짜>_1ofN.zip   핵심: 프로그램·Python·설치 도구(이 파일만 압축 해제)
  Blue365_QMS_win64_v<날짜>_kofN.zip   구성요소: 1_SETUP.bat 이 SHA-256 확인 후 자동 결합
  SHA256SUMS.txt · components.xlsx(구성요소·라이선스·해시 명세, IT 보안 검토용)
"""

from __future__ import annotations

import argparse
import base64
import csv
import email.parser
import hashlib
import io
import json
import math
import shutil
import subprocess
import urllib.request
import zipfile
import zlib
from datetime import UTC, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PKG = ROOT / "packaging" / "windows"
REQ = PKG / "requirements-windows.txt"
LOCK = PKG / "requirements-windows.lock"
TESTED = PKG / "constraints-tested.txt"        # 테스트를 통과한 버전(잠금 파일 고정 기준)
TOP = "Blue365_QMS"
PY_VERSION = "3.13.16"
UV_PLATFORM = "x86_64-pc-windows-msvc"
NUGET_PKG = "https://api.nuget.org/v3-flatcontainer/python/{v}/python.{v}.nupkg"
NUGET_REG = "https://api.nuget.org/v3/registration5-semver1/python/{v}.json"
APP_EXCLUDE = ("tests/", "scripts/", "packaging/", ".github/", ".devcontainer/", "data/", "docs/src/",
               ".gitignore", "requirements-dev.txt")
SITE_TRIM_TOP = ("bin", "etc", "share")       # 리눅스용 실행 스크립트·Jupyter 확장(실행에 불필요)
SITE_TRIM_DIRNAME = "tests"                    # 각 라이브러리 자체 테스트(실행에 불필요)
SITE_TRIM_PATHS = ("pyarrow/include",)         # C++ 헤더
SITE_TRIM_SUFFIXES = (".lib",)                 # 링크용 import library(실행 시 사용 안 함)
PY_TRIM = ("include", "libs")                  # Python 개발용 헤더·링크 파일
CORE_GROUP_BYTES = 3_000_000                   # 이보다 작은 라이브러리는 핵심(1번) 파일에 넣음
CORE_DISTS = {"streamlit", "pandas", "numpy", "scipy", "plotly", "openpyxl", "python-pptx"}
OPTIONAL_DISTS = {"anthropic": "AI 솔루션 보고서", "sqlalchemy": "LIMS DB 조회", "requests": "LIMS REST API",
                  "pyodbc": "MS SQL Server(ODBC)", "pymssql": "MS SQL Server", "oracledb": "Oracle",
                  "psycopg2-binary": "PostgreSQL", "pymysql": "MySQL·MariaDB", "xlrd": "구형 엑셀(.xls)"}


def log(msg: str) -> None:
    print(msg, flush=True)


def http_get(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "Blue365-QMS-build"})
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.read()


def norm(name: str) -> str:
    return name.lower().replace("_", "-").replace(".", "-")


# ───────────────────────── Python(NuGet) ─────────────────────────
def fetch_python(version: str, cache: Path) -> tuple[Path, dict]:
    """python.org 공식 NuGet 패키지를 받아 NuGet 카탈로그의 SHA-512·게시자와 대조한다."""
    cache.mkdir(parents=True, exist_ok=True)
    path = cache / f"python.{version}.nupkg"
    cat = json.loads(http_get(json.loads(http_get(NUGET_REG.format(v=version)))["catalogEntry"]))
    if not path.exists():
        log(f"  Python {version} 내려받기(nuget.org)…")
        path.write_bytes(http_get(NUGET_PKG.format(v=version)))
    digest = base64.b64encode(hashlib.sha512(path.read_bytes()).digest()).decode()
    if cat.get("packageHashAlgorithm") != "SHA512" or digest != cat.get("packageHash"):
        path.unlink()
        raise SystemExit("NuGet 패키지 해시가 카탈로그와 다릅니다 — 다시 실행하세요.")
    if cat.get("authors") != "Python Software Foundation":
        raise SystemExit(f"게시자가 예상과 다릅니다: {cat.get('authors')}")
    return path, {"version": version, "source": f"https://www.nuget.org/packages/python/{version}",
                  "authors": cat["authors"], "published": cat.get("published"), "sha512_base64": digest,
                  "nupkg_bytes": path.stat().st_size}


def stage_python(nupkg: Path, py_dir: Path) -> None:
    with zipfile.ZipFile(nupkg) as z:
        for info in z.infolist():
            if info.filename.startswith("tools/") and not info.is_dir():
                target = py_dir / info.filename[len("tools/"):]
                target.parent.mkdir(parents=True, exist_ok=True)
                with z.open(info) as src, target.open("wb") as dst:
                    shutil.copyfileobj(src, dst)
    for name in PY_TRIM:
        shutil.rmtree(py_dir / name, ignore_errors=True)


# ───────────────────────── 구성요소 ─────────────────────────
def uv() -> str:
    exe = shutil.which("uv")
    if not exe:
        raise SystemExit("uv 가 필요합니다: pip install uv")
    return exe


def relock(py_minor: str, constraints: Path | None) -> None:
    cmd = [uv(), "pip", "compile", str(REQ.relative_to(ROOT)), "--python-platform", UV_PLATFORM,
           "--python-version", py_minor, "--only-binary", ":all:", "--generate-hashes", "-o", str(LOCK.relative_to(ROOT)),
           "--custom-compile-command", "python scripts/build_windows_package.py --relock"]
    if constraints:
        cmd += ["-c", str(constraints.relative_to(ROOT) if constraints.is_relative_to(ROOT) else constraints)]
    subprocess.run(cmd, cwd=ROOT, check=True)


def install_packages(site: Path, py_minor: str) -> None:
    subprocess.run([uv(), "pip", "install", "--target", str(site), "--python-platform", UV_PLATFORM,
                    "--python-version", py_minor, "--no-deps", "--require-hashes", "--only-binary", ":all:",
                    "--link-mode", "copy", "-q", "-r", str(LOCK)], cwd=ROOT, check=True)


def trim_site(site: Path) -> int:
    removed = 0

    def rm(p: Path) -> None:
        nonlocal removed
        if p.is_dir():
            removed += sum(f.stat().st_size for f in p.rglob("*") if f.is_file())
            shutil.rmtree(p)
        elif p.is_file():
            removed += p.stat().st_size
            p.unlink()

    for name in SITE_TRIM_TOP:
        rm(site / name)
    for rel in SITE_TRIM_PATHS:
        rm(site / rel)
    for d in sorted((p for p in site.rglob(SITE_TRIM_DIRNAME) if p.is_dir()), key=lambda p: len(p.parts)):
        if d.exists() and not d.parent.name.endswith(".dist-info"):
            rm(d)
    for p in list(site.rglob("*")):
        if p.is_file() and p.suffix.lower() in SITE_TRIM_SUFFIXES or p.name == "__pycache__":
            rm(p)
    return removed


def check_binaries(site: Path) -> None:
    linux = [p for p in site.rglob("*.so")]
    if linux:
        raise SystemExit(f"리눅스용 확장 모듈이 섞였습니다: {linux[:3]}")
    bad = []
    for wheel in site.glob("*.dist-info/WHEEL"):
        for line in wheel.read_text(encoding="utf-8").splitlines():
            if line.startswith("Tag:") and line.split("-")[-1].strip() not in ("any", "win_amd64"):
                bad.append((wheel.parent.name, line))
    if bad:
        raise SystemExit(f"Windows 64비트용이 아닌 구성요소: {bad[:5]}")


def dist_info(site: Path) -> list[dict]:
    """설치된 배포 묶음별 이름·버전·라이선스·요약·소유 파일."""
    out = []
    for di in sorted(site.glob("*.dist-info")):
        meta = email.parser.Parser().parsestr((di / "METADATA").read_text(encoding="utf-8", errors="replace"))
        lic = meta.get("License-Expression") or ""
        if not lic:
            classifiers = [c.split("::")[-1].strip() for c in meta.get_all("Classifier") or [] if c.startswith("License ::")]
            lines = [x.strip() for x in (meta.get("License") or "").splitlines() if x.strip()]
            if lines and len(lines[0]) <= 60:       # 라이선스 전문이 들어 있으면 제목·버전 줄만 사용
                lines = lines[:2] if len(lines) > 1 and lines[1].lower().startswith("version") else lines[:1]
            lic = ", ".join(classifiers) or " ".join(lines if lines and len(lines[0]) <= 60 else [])
        urls = [u.split(",", 1)[-1].strip() for u in meta.get_all("Project-URL") or []]
        tops = set()
        rec = di / "RECORD"
        if rec.exists():
            for row in csv.reader(io.StringIO(rec.read_text(encoding="utf-8"))):
                if row and row[0] and not row[0].startswith(".."):
                    tops.add(row[0].split("/")[0])
        tops.add(di.name)
        out.append({"name": meta["Name"], "version": meta["Version"], "license": lic or "(METADATA 참고)",
                    "summary": (meta.get("Summary") or "").strip(), "home": meta.get("Home-page") or (urls[0] if urls else ""),
                    "dist_info": di.name, "tops": sorted(tops)})
    return out


# ───────────────────────── 프로그램·도구·배치 ─────────────────────────
def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()


def app_files() -> list[str]:
    raw = subprocess.run(["git", "ls-files", "-z", "--cached", "--others", "--exclude-standard"], cwd=ROOT,
                         capture_output=True, check=True).stdout.decode("utf-8")
    files = sorted({f for f in raw.split("\0") if f and (ROOT / f).is_file() and not f.startswith(APP_EXCLUDE)})
    if any(Path(f).name == "secrets.toml" for f in files):
        raise SystemExit("secrets.toml 이 포함될 뻔했습니다 — 빌드를 중단합니다.")
    return files


def to_bat(text: str) -> bytes:
    """배치 파일: 한글 Windows 기본 코드 페이지(CP949) + CRLF."""
    return text.replace("\r\n", "\n").replace("\n", "\r\n").encode("cp949")


def to_txt(text: str) -> bytes:
    """메모장용 텍스트: UTF-8(BOM) + CRLF."""
    return text.replace("\r\n", "\n").replace("\n", "\r\n").encode("utf-8-sig")


def max_rel_path(stage_top: Path) -> int:
    """설치 폴더 기준 가장 긴 경로(설치 후 생기는 __pycache__\\*.cpython-3xx.pyc 포함)."""
    tag = f".cpython-{PY_VERSION.split('.')[0]}{PY_VERSION.split('.')[1]}.pyc"
    longest = 0
    for p in stage_top.rglob("*"):
        if p.is_file():
            rel = p.relative_to(stage_top).as_posix()
            longest = max(longest, len(rel))
            if p.suffix == ".py":
                longest = max(longest, len(p.parent.relative_to(stage_top).as_posix()) + 13 + len(p.stem) + len(tag))
    return longest


# ───────────────────────── 분할 ─────────────────────────
def zsize(files: list[Path]) -> int:
    return sum(len(zlib.compress(f.read_bytes(), 6)) + 120 + 2 * len(str(f)) for f in files)


def pack(sizes: dict[str, int], limit: int) -> list[list[str]]:
    """큰 묶음부터 가장 가벼운 파일에 넣는 방식(파일 수는 한도를 넘지 않을 때까지 늘림)."""
    items = sorted(sizes.items(), key=lambda kv: -kv[1])
    n = max(1, math.ceil(sum(sizes.values()) / limit))
    while True:
        bins: list[list[str]] = [[] for _ in range(n)]
        load = [0] * n
        for name, size in items:
            i = min(range(n), key=lambda k: load[k])
            bins[i].append(name)
            load[i] += size
        if max(load) <= limit or n >= 30:
            return bins
        n += 1


def write_zip(path: Path, entries: list[tuple[Path, str]], last: tuple[Path, str]) -> None:
    """항목을 경로순으로 쓰고, 완료 표식(last)은 맨 마지막에 쓴다 — 표식이 있으면 앞의 파일이 모두 풀린 것."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for src, arc in sorted(entries, key=lambda t: t[1]):
            z.write(src, arc)
        z.write(*last)


def write_marker(top: Path, k: int, n: int, zip_name: str, version: str) -> tuple[Path, str]:
    rel = f"tools/parts/{k}of{n}.done"
    (top / rel).parent.mkdir(parents=True, exist_ok=True)
    (top / rel).write_text(json.dumps({"part": zip_name, "version": version}, ensure_ascii=False), encoding="utf-8")
    return top / rel, f"{TOP}/{rel}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ───────────────────────── 구성요소 명세(엑셀) ─────────────────────────
def components_workbook(path: Path, dists: list[dict], dist_part: dict[str, str], dist_bytes: dict[str, int],
                        manifest: dict, part_rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    head_fill, head_font = PatternFill("solid", fgColor="1F4E79"), Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, header: list[str], rows: list[list], widths: list[int]) -> None:
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = w
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"

    def role(name: str) -> tuple[int, str]:
        n = norm(name)
        if n in CORE_DISTS:
            return 0, "핵심(필수)"
        if n in OPTIONAL_DISTS:
            return 1, f"선택 기능 — {OPTIONAL_DISTS[n]}"
        return (3, "Python 기본 포함(패키지 관리)") if n == "pip" else (2, "의존 구성요소")

    wb = Workbook()
    ws = wb.active
    ws.title = "구성요소"
    rows = sorted(dists, key=lambda d: (role(d["name"])[0], norm(d["name"])))
    sheet(ws, ["번호", "구성요소", "버전", "용도", "라이선스", "설명(배포처 요약)", "설치 용량(MB)", "포함 파일", "홈페이지"],
          [[i, d["name"], d["version"], role(d["name"])[1], d["license"], d["summary"],
            round(dist_bytes.get(d["dist_info"], 0) / 1e6, 1), f"{dist_part.get(d['dist_info'], '1')}번", d["home"]]
           for i, d in enumerate(rows, 1)], [6, 22, 12, 28, 26, 50, 12, 9, 40])

    ws2 = wb.create_sheet("분할 파일")
    sheet(ws2, ["순서", "파일명", "크기(MB)", "SHA-256", "파일 수", "내용"],
          [[r["order"], r["file"], round(r["size"] / 1e6, 1) if r["size"] else "SHA256SUMS.txt 참고",
            r.get("sha256") or "SHA256SUMS.txt 참고", r["files"] or "", r["desc"]]
           for r in part_rows], [7, 40, 12, 68, 9, 50])

    py = manifest["python"]
    ws3 = wb.create_sheet("보안·구성 정보")
    sheet(ws3, ["항목", "내용"], [
        ["패키지 버전", f"{manifest['version']} (빌드 {manifest['built_at']})"],
        ["Python", f"{py['version']} 64비트 — python.org 공식 NuGet 배포본, 게시자 {py['authors']}, {py['source']}"],
        ["Python 무결성", f"NuGet 카탈로그 SHA-512 일치 확인: {py['sha512_base64']}"],
        ["설치 방식", "휴대용(압축 해제만) — Windows에 Python을 설치하지 않음, 레지스트리 변경 없음, 관리자 권한 불필요"],
        ["외부 통신", ("기본 없음(오프라인 동작, 사용 통계 전송 끔). 예외: 사용자가 켜는 AI 보고서(api.anthropic.com), "
                     "사용자가 설정한 메일(SMTP)·웹훅·LIMS(DB·REST)")],
        ["네트워크 포트", "TCP 8501(기본). 2_RUN.bat 은 127.0.0.1(이 PC)만, 3_RUN_SHARE.bat 은 사내망 접속 허용"],
        ["접속 통제", "선택: app\\.streamlit\\secrets.toml [access] password 설정 시 접속 비밀번호 요구"],
        ["데이터 저장", "설치 폴더\\data (SQLite DB·JSON 설정), 실행 기록은 설치 폴더\\logs"],
        ["설치 폴더 밖 변경", ("작업 스케줄러 작업 Blue365_QMS_Monitor(5번 등록 시), 방화벽 규칙 Blue365_QMS_TCP8501·"
                           "Blue365_QMS_python(4번 실행 시)")],
        ["관리자 권한 필요", "방화벽 규칙 등록(4_FIREWALL_ADMIN.bat)만"],
        ["구성요소 수", f"{len(dists)}종(목록: '구성요소' 시트)"],
    ], [22, 110])
    wb.save(path)


# ───────────────────────── 빌드 ─────────────────────────
def build(args) -> None:
    py_minor = ".".join(PY_VERSION.split(".")[:2])
    out = Path(args.out).resolve()
    stage = out / "stage"
    top = stage / TOP
    if args.relock:
        log("[잠금 파일 갱신]")
        relock(py_minor, Path(args.constraints).resolve() if args.constraints else TESTED)
    if not LOCK.exists():
        raise SystemExit(f"잠금 파일이 없습니다: {LOCK} — --relock 으로 만드세요.")

    shutil.rmtree(stage, ignore_errors=True)
    for old in out.glob(f"{TOP}_win64_*.zip"):
        old.unlink()
    top.mkdir(parents=True)

    log(f"[1/6] Python {PY_VERSION}(휴대용) 준비")
    nupkg, py_info = fetch_python(PY_VERSION, out / "cache")
    stage_python(nupkg, top / "py")
    site = top / "py" / "Lib" / "site-packages"

    log("[2/6] 구성요소 설치(Windows 64비트용, 해시 검증)")
    install_packages(site, py_minor)
    removed = trim_site(site)
    check_binaries(site)
    dists = dist_info(site)
    log(f"  {len(dists)}종 설치 · 실행에 불필요한 파일 {removed / 1e6:.0f}MB 정리")

    log("[3/6] 프로그램·설치 도구·배치 파일")
    files = app_files()
    for f in files:
        (top / "app" / f).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ROOT / f, top / "app" / f)
    (top / "tools").mkdir()
    for f in (PKG / "launcher").iterdir():
        if f.suffix in (".py", ".pyw"):
            shutil.copy2(f, top / "tools" / f.name)
    shutil.copy2(LOCK, top / "tools" / LOCK.name)
    for bat in sorted((PKG / "bat").glob("*.bat")):
        (top / bat.name).write_bytes(to_bat(bat.read_text(encoding="utf-8")))
    log(f"  프로그램 파일 {len(files)}개 · 배치 파일 {len(list((PKG / 'bat').glob('*.bat')))}개")

    log("[4/6] 분할 구성")
    commit = git("rev-parse", "--short", "HEAD") + ("-dirty" if git("status", "--porcelain") else "")
    now = datetime.now(UTC)
    version = f"{now:%Y.%m.%d}-{commit}"
    tag = f"v{now:%Y.%m.%d}"
    owner = {t: d["dist_info"] for d in dists for t in d["tops"]}
    groups: dict[str, list[Path]] = {}
    for entry in site.iterdir():
        key = owner.get(entry.name, "_unowned")
        groups.setdefault(key, []).extend([entry] if entry.is_file() else [p for p in entry.rglob("*") if p.is_file()])
    dist_bytes = {k: sum(p.stat().st_size for p in v) for k, v in groups.items()}
    limit = int(args.max_part_mb * 1e6)
    big = {k: v for k, v in groups.items() if limit and k != "_unowned" and dist_bytes[k] >= CORE_GROUP_BYTES}
    bins = pack({k: zsize(v) for k, v in big.items()}, limit) if big else []
    n_parts = 1 + len(bins)

    def name(k: int) -> str:
        return f"{TOP}_win64_{tag}_{k}of{n_parts}.zip"

    lib_files = {p for k in big for p in big[k]}

    log(f"[5/6] ZIP 생성({n_parts}개)")
    parts, part_rows, dist_part, markers = [], [], {}, set()
    for i, keys in enumerate(bins, 2):
        entries = [(p, f"{TOP}/{p.relative_to(top).as_posix()}") for k in keys for p in big[k]]
        zpath = out / name(i)
        mark = write_marker(top, i, n_parts, zpath.name, version)
        markers.add(mark[0])
        write_zip(zpath, entries, mark)
        names = sorted(k.split("-")[0] for k in keys)
        parts.append({"file": zpath.name, "sha256": sha256(zpath), "size": zpath.stat().st_size, "files": len(entries) + 1,
                      "marker": mark[1][len(TOP) + 1:], "dists": names})
        part_rows.append({"order": f"{i}/{n_parts}", "file": zpath.name, "size": zpath.stat().st_size,
                          "sha256": parts[-1]["sha256"], "files": len(entries) + 1, "desc": "구성요소: " + ", ".join(names)})
        dist_part.update({k: f"{i}" for k in keys})
        log(f"  {zpath.name}: {zpath.stat().st_size / 1e6:.1f} MB · {len(entries):,}개 파일")

    manifest = {"name": "Blue365 QMS", "version": version, "tag": tag, "git_commit": commit,
                "built_at": f"{now:%Y-%m-%d %H:%M} UTC", "python": py_info,
                "packages": [{"name": d["name"], "version": d["version"]} for d in dists],
                "parts": parts, "core_file": name(1), "core_marker": f"tools/parts/1of{n_parts}.done",
                "max_rel_path": max_rel_path(top)}
    (top / "tools" / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")

    core_rows = [{"order": f"1/{n_parts}", "file": name(1), "size": 0, "files": 0,
                  "desc": "핵심: 프로그램·Python·설치 도구·안내서 — 이 파일만 압축 해제"}] + part_rows
    table = []
    for r in core_rows:
        sha = f"SHA-256 {r['sha256']}" if r.get("sha256") else "SHA-256: SHA256SUMS.txt 참고(받은 파일과 함께 제공)"
        size = f"{r['size'] / 1e6:.1f} MB" if r["size"] else ""
        table.append(f"  [{r['order']}] {r['file']}  {size}\n        {r['desc']}\n        {sha}")
    streamlit_v = next((d["version"] for d in dists if norm(d["name"]) == "streamlit"), "?")
    readme = (PKG / "README.txt").read_text(encoding="utf-8")
    for key, val in {"{VERSION}": f"{tag} ({commit})", "{PY_VERSION}": PY_VERSION, "{N_PKGS}": str(len(dists)),
                     "{STREAMLIT}": streamlit_v, "{PARTS_TABLE}": "\n".join(table), "{N}": str(n_parts)}.items():
        readme = readme.replace(key, val)
    (top / "README.txt").write_bytes(to_txt(readme))
    components_workbook(top / "tools" / "components.xlsx", dists, dist_part, dist_bytes, manifest, core_rows)

    core_mark = write_marker(top, 1, n_parts, name(1), version)
    core_entries = [(p, f"{TOP}/{p.relative_to(top).as_posix()}") for p in top.rglob("*")
                    if p.is_file() and p not in lib_files and p not in markers and p != core_mark[0]]
    core = out / name(1)
    write_zip(core, core_entries, core_mark)
    core_entries.append(core_mark)
    log(f"  {core.name}: {core.stat().st_size / 1e6:.1f} MB · {len(core_entries):,}개 파일")

    sums = [(sha256(out / name(k)), name(k)) for k in range(1, n_parts + 1)]
    (out / "SHA256SUMS.txt").write_bytes(to_txt("".join(f"{h} *{f}\n" for h, f in sums)))
    core_rows[0].update(size=core.stat().st_size, sha256=sums[0][0], files=len(core_entries))   # 배포용 명세는 1번 정보까지
    components_workbook(out / "components.xlsx", dists, dist_part, dist_bytes, manifest, core_rows)

    log("[6/6] 검증")
    validate(out, top, [out / name(k) for k in range(1, n_parts + 1)], manifest)
    total = sum((out / name(k)).stat().st_size for k in range(1, n_parts + 1))
    log(f"\n완료: {out}  (총 {total / 1e6:.1f} MB, {n_parts}개 파일, 가장 긴 상대 경로 {manifest['max_rel_path']}자)")
    for h, f in sums:
        log(f"  {h}  {f}")


def validate(out: Path, top: Path, zips: list[Path], manifest: dict) -> None:
    seen: set[str] = set()
    for z in zips:
        with zipfile.ZipFile(z) as zf:
            bad = zf.testzip()
            if bad:
                raise SystemExit(f"ZIP 손상: {z.name} / {bad}")
            names = set(zf.namelist())
            if seen & names:
                raise SystemExit(f"분할 파일 간 중복 항목: {sorted(seen & names)[:3]}")
            seen |= names
            for n in names:
                if n.endswith(".bat"):
                    raw = zf.read(n)
                    raw.decode("cp949")
                    if raw.count(b"\n") != raw.count(b"\r\n"):
                        raise SystemExit(f"배치 파일 줄바꿈 오류: {n}")
    expect = {f"{TOP}/{p.relative_to(top).as_posix()}" for p in top.rglob("*") if p.is_file()}
    if seen != expect:
        raise SystemExit(f"누락/초과 항목: {sorted(expect - seen)[:3]} / {sorted(seen - expect)[:3]}")
    if any(n.endswith("/secrets.toml") for n in seen):
        raise SystemExit("secrets.toml 포함 — 배포 금지")
    for z, marker in zip(zips, [manifest["core_marker"]] + [p["marker"] for p in manifest["parts"]]):
        with zipfile.ZipFile(z) as zf:
            if zf.infolist()[-1].filename != f"{TOP}/{marker}":
                raise SystemExit(f"완료 표식이 맨 마지막 항목이 아닙니다: {z.name}")
    log(f"  ZIP 무결성·중복·누락·비밀파일·배치 인코딩 확인 완료({len(seen):,}개 항목)")


def main() -> None:
    ap = argparse.ArgumentParser(description="사내망(오프라인) Windows 설치 패키지 빌드")
    ap.add_argument("--out", default=str(ROOT / "dist" / "windows"), help="출력 폴더(기본 dist/windows)")
    ap.add_argument("--max-part-mb", type=float, default=95, help="분할 파일 최대 크기(MB), 0이면 한 파일")
    ap.add_argument("--relock", action="store_true", help="잠금 파일(requirements-windows.lock) 갱신")
    ap.add_argument("-c", "--constraints", help="--relock 때 버전 고정 파일(기본 packaging/windows/constraints-tested.txt)")
    build(ap.parse_args())


if __name__ == "__main__":
    main()
