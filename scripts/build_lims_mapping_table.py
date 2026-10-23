"""LIMS 공정별 내보내기 파일 → QMS 항목 매칭표(엑셀) 생성.

업로드한 LIMS 샘플들의 머리글을 qms.store 의 인식 규칙(detect_lims_sheet·COL_ALIASES)으로 돌려,
'어느 파일의 어느 열이 QMS 어느 항목으로 들어가는지'를 요약·상세·참고 3개 시트로 정리한다.

사용:  python scripts/build_lims_mapping_table.py <LIMS 파일들이 있는 폴더> [출력.xlsx]
"""

from __future__ import annotations

import glob
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from qms import store

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs" / "LIMS_공정별_열매칭표.xlsx"

NAVY = "1F3A5F"
HEAD = PatternFill("solid", fgColor=NAVY)
HEADF = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
C = Alignment("center", "center", wrap_text=True)
L = Alignment("left", "center", wrap_text=True)
GREEN = PatternFill("solid", fgColor="E8F3E8")
GREY = PatternFill("solid", fgColor="F2F2F2")

# 파일 제목(정규화 포함) → 사람이 읽는 공정명. detect_lims_sheet 가 표를 정하고, 여기선 이름만 붙인다.
PROC_NAME = [("조합원료", "조합원료성적(생료·킬른 피드)"), ("수출크링커", "수출크링커성적(6·7호)"),
             ("크링커", "크링커성적(4·5호)"), ("시멘트공정", "시멘트공정(밀 출구)"),
             ("3종시멘트성적", "3종시멘트성적(출하)"), ("시멘트성적", "시멘트성적(1종·출하)")]
META = {store._norm_header(x) for x in (
    "로트번호", "검사일", "담당자", "비고", "자료생성일", "최종수정일", "전송", "ERP LOT.NO", "LIMS LOT.NO", "분석방법",
    "채취시간", "공정구분", "투입SILO", "파일구분", "샘플위치", "시멘트밀공정", "시멘트구분", "크링커공정", "클링커공정",
    "클링커구분", "스크린구분", "수출크링커구분", "입도상태", "요입위치", "요입인출SILO", "직송여부", "특이사항",
    "투입인출사이로", "구치장여부", "품위구분", "거래처", "수량", "합부판정", "색상", "채취일")}
DERIVED = {store._norm_header(x) for x in ("HM", "SM", "IM", "LSF", "C3S", "C2S", "C3A", "C4AF", "TA", "SD", "A/S",
                                          "mC3S", "Coating tendency", "Liquid Phase", "KFUI", "TargetC3S", "ActualC3S")}
TITLE_SKIP = tuple(store._norm_header(s) for s in store.LIMS_SKIP_TITLES)
QNAME = {"timestamp": "채취일(기준)", "product": "품명(품종)", "rm_cao": "생료 CaO", "rm_sio2": "생료 SiO2",
         "rm_al2o3": "생료 Al2O3", "rm_fe2o3": "생료 Fe2O3", "rm_mgo": "생료 MgO", "rm_so3": "생료 SO3",
         "rm_k2o": "생료 K2O", "rm_na2o": "생료 Na2O", "rm_r90": "생료 90㎛", "rm_r200": "생료 200㎛", "rm_loi": "생료 LOI",
         "clk_cao": "클링커 CaO", "clk_sio2": "클링커 SiO2", "clk_al2o3": "클링커 Al2O3", "clk_fe2o3": "클링커 Fe2O3",
         "clk_mgo": "클링커 MgO", "clk_so3": "클링커 SO3", "clk_k2o": "클링커 K2O", "clk_na2o": "클링커 Na2O",
         "clk_fcao": "클링커 유리석회(f-CaO)", "clk_lw": "클링커 용적비중(용중)", "cem_blaine": "시멘트 Blaine",
         "cem_r45": "시멘트 45㎛", "cem_so3": "시멘트 SO3", "cem_loi": "시멘트 LOI", "cem_mgo": "시멘트 MgO",
         "cem_ls": "시멘트 석회석비율", "cem_mill_feed": "시멘트 밀피드", "cem_mill_temp": "시멘트 밀온도",
         "phy_ist": "초결시간", "phy_fst": "종결시간", "phy_autoclave": "Autoclave", "phy_s1": "압축강도 1일",
         "phy_s3": "압축강도 3일", "phy_s7": "압축강도 7일", "phy_s28": "압축강도 28일", "phy_crvi": "6가크롬",
         "operator": "담당자"}
TBL_SHEET = {"raw_meal": "생료", "clinker": "클링커", "cement": "시멘트", "physical": "물성 + 출하시멘트"}


def _title(ws) -> str:
    return " ".join(store._norm_header(ws.cell(r, c).value) for r in range(1, 4)
                    for c in range(1, ws.max_column + 1) if store._norm_header(ws.cell(r, c).value))


def _proc_name(t: str) -> str:
    for key, nm in PROC_NAME:
        if store._norm_header(key) in t:
            return nm
    return "LIMS 파일"


def _style(ws, r, cols, fill=None, left_cols=()):
    for c in range(1, cols + 1):
        cell = ws.cell(r, c)
        cell.border = BORDER
        cell.alignment = L if c in left_cols else C
        if fill and c == cols:
            cell.fill = fill


def _headrow(ws, headers):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, HEADF, C, BORDER
    ws.freeze_panes = "A2"


def build():
    files = sorted(glob.glob(str(SRC / "*LIMS*.xlsx")))
    wb = openpyxl.Workbook()
    summ = wb.active
    summ.title = "요약"
    _headrow(summ, ["LIMS 파일(공정)", "들어갈 QMS 화면/표", "인식 열수", "참고(미반영)", "상태"])
    detail = wb.create_sheet("상세 매칭")
    _headrow(detail, ["LIMS 파일(공정)", "QMS 표", "LIMS 열 이름", "→ QMS 항목"])
    unmap = wb.create_sheet("미반영 참고")
    _headrow(unmap, ["LIMS 파일(공정)", "LIMS 열 이름", "사유"])
    unmap.append(["※ HM·SM·IM·LSF·C3S 등 지수·Bogue 광물상은 QMS가 화학성분에서 자동 계산하므로 올릴 필요가 없습니다.",
                  "", ""])
    unmap.merge_cells("A2:C2")
    unmap["A2"].font = Font(italic=True, color="7A7A7A", size=9)
    ur = 3
    sr = dr = 2

    for f in files:
        ws = openpyxl.load_workbook(f, data_only=True).active
        t = _title(ws)
        raw_title = str(ws.cell(1, 1).value or "").strip() or "LIMS 파일"
        hdr = [ws.cell(3, c).value for c in range(1, ws.max_column + 1) if ws.cell(3, c).value]
        if any(s in t for s in TITLE_SKIP):
            summ.append([raw_title, "📦 인수검사 화면(별도 관리)", f"메타+성분 {len(hdr)}열", "합부·거래처·성분 저장", "✅ 반영"])
            _style(summ, sr, 5, GREY, (1, 2))
            sr += 1
            continue
        name = _proc_name(t)
        det = store.detect_lims_sheet(openpyxl_to_df(f))
        if det is None:
            continue
        tbl = det[0]
        amap = store.COL_ALIASES[tbl]
        valid = set(store.RAW_COLUMNS[tbl]) | {"timestamp", "product"}
        mapped, derived, nostore = [], [], []
        for h in hdr:
            nh = store._norm_header(h)
            if nh in META:
                continue
            k = amap.get(nh)
            if k and k in valid:
                mapped.append((h, QNAME.get(k, k)))
            elif nh in DERIVED:
                derived.append(h)
            else:
                nostore.append(h)
        summ.append([name, f"{TBL_SHEET[tbl]} ({tbl})", len(mapped), f"자동계산 {len(derived)} · 미저장 {len(nostore)}",
                     "✅ 반영"])
        _style(summ, sr, 5, GREEN, (1, 2))
        sr += 1
        for h, q in mapped:
            detail.append([name, TBL_SHEET[tbl], h, q])
            _style(detail, dr, 4, None, (1, 3, 4))
            dr += 1
        for h in nostore:
            unmap.append([name, h, "현재 QMS 표에 저장 항목 없음(필요 시 추가 가능)"])
            _style(unmap, ur, 3, None, (1, 2))
            ur += 1

    for col, wd in zip("ABCDE", [30, 22, 12, 20, 10]):
        summ.column_dimensions[col].width = wd
    for col, wd in zip("ABCD", [30, 18, 26, 24]):
        detail.column_dimensions[col].width = wd
    for col, wd in zip("ABC", [30, 26, 46]):
        unmap.column_dimensions[col].width = wd
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("저장:", OUT)


def openpyxl_to_df(path):
    import pandas as pd
    return pd.read_excel(path, sheet_name=0, header=None)


if __name__ == "__main__":
    build()
