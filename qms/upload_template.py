"""통합 업로드 양식(엑셀) — 한 파일에 공정별 시트 + 출하시멘트 + 인수검사.

📥 데이터 관리 → '입력 템플릿' 다운로드와 docs/Blue365_QMS_업데이트_양식.xlsx(scripts/build_lims_update_workbook.py)가
같은 양식을 쓴다. 머리글은 LIMS 이름(또는 관리항목 이름)이라 qms.store·qms.shipment·qms.incoming 이 그대로 인식한다.

시트 구성(2026-10-28 사용자 수정본 기준): 사용법 · 공정구분(코드표) · ①원료 · ①원료 잔사 · ②킬른 · ③클링커(6K,7K)
          · ③클링커(4K,5K) · ④클링커일일 · ⑤시멘트 · ⑦출하시멘트(1종/3종/고로슬래그 2종) · ⑧인수검사(원료·석고·고령토) · 항목사전
          (물성 시트는 쓰지 않음 — 강도는 출하시멘트 시트. 예전 '물성' 시트가 있어도 그대로 읽는다)
각 자료 시트: 1행 제목 · 2행 안내 · 3행 머리글(주황=필수) · 4행 단위·형식(회색, 읽지 않음) · 5행부터 입력.
공정구분·품명 칸은 목록에서 고르게(데이터 유효성 검사) 되어 있다.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, ORANGE, GRAY = "1F3A5F", "D9731A", "6B7686"
HEAD = PatternFill("solid", fgColor=NAVY)
REQ = PatternFill("solid", fgColor=ORANGE)
UNIT = PatternFill("solid", fgColor="EEF2F8")
INPUT = PatternFill("solid", fgColor="FFFFFF")
WHITE_B = Font(color="FFFFFF", bold=True, size=10)
TITLE = Font(bold=True, size=14, color=NAVY)
NOTE = Font(color=GRAY, italic=True, size=9)
THIN = Side(style="thin", color="D0D7E2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
INPUT_ROWS = 1000           # 목록 선택·서식을 미리 걸어 두는 입력 행 수

# ── 코드표(공정구분 시트) — 사용자가 고친 양식(2026-10-28) 기준. 근거 열에 사실/추정 표시 ─────────────
KILNS = ["3K", "4K", "5K", "6K", "7K"]
RAW_MILLS = ["3R/M", "4R/M", "5R/M", "5-1R/M", "6R/M", "7R/M"]
CEMENT_MILLS = ["2C/M", "3C/M", "4C/M", "5C/M", "6C/M", "7C/M"]
CEMENT_PRODUCTS = ["1종", "3종", "고로슬래그 2종", "고로슬래그 3종"]
RAW_GRADES = ["1종조합원료", "1/2종조합원료"]
CLK_GRADES = ["1종크링커", "1/2종크링커", "내수", "수출클링커"]
CEMENT_NAMES = ["1종시멘트 벌크", "3종시멘트 벌크", "2종시멘트 벌크", "2종고로슬래그시멘트 벌크"]
CLK_SILOS = ["3호", "4호", "5호", "6-1호", "6-2호", "6-3호", "6-4호", "6-5호", "6-6호"]
CEM_SILOS = [f"{i}호" for i in range(1, 12)]
KAOLIN = ["기존", "신설"]
CODE_TABLE = [
    # (구분, 코드 목록, 의미, 쓰는 시트, 근거)
    ("킬른(호기)", KILNS, "킬른 호기 — 킬른 요입(킬른 피드) 시료·클링커·DCS 운전값", "원료·킬른·클링커·클링커일일",
     "사용자 지정(킬른 5기: 3K~7K)"),
    ("원료밀", RAW_MILLS, "원료밀 출구(밀출분) 시료 — 배합비(투입비율)는 이 행에만", "원료·원료 잔사",
     "사용자 지정. LIMS 표기 5_1R/M 은 5-1R/M 과 같은 것으로 읽음"),
    ("시멘트밀", CEMENT_MILLS, "시멘트밀 출구 시료", "시멘트", "LIMS 시멘트공정 원본 2C/M~7C/M 확인"),
    ("시멘트 품종", CEMENT_PRODUCTS, "KS L 5201 1종·3종, KS L 5210 고로슬래그 2·3종", "(참고) 품명에서 자동 판정",
     "품명 '…2종시멘트 벌크' = 고로슬래그 2종(첨가재 고로슬래그 확인)"),
    ("원료 종류(품명)", RAW_GRADES, "조합원료 품명(참고 정보)", "원료·원료 잔사", "LIMS 조합원료 원본 품명"),
    ("클링커 종류(품명)", CLK_GRADES, "클링커 품명·구분(참고 정보)", "클링커", "사용자 지정"),
    ("시멘트 종류(품명)", CEMENT_NAMES, "LIMS 품명 그대로 — QMS가 1종·3종·고로슬래그 2종으로 자동 구분", "시멘트·출하시멘트",
     "사용자 지정 + LIMS 실제 표기 '2종시멘트 벌크'"),
    ("클링커 투입사일로", CLK_SILOS, "클링커를 넣은 사일로(추적용)", "클링커", "사용자 지정"),
    ("시멘트 투입사일로", CEM_SILOS, "시멘트를 넣은 사일로(추적용). LIMS 표기 '구내9호' 등도 그대로 저장", "시멘트", "사용자 지정"),
    ("고령토 구분", KAOLIN, "고령토 채광 구분", "인수검사(고령토)", "사용자 지정"),
]


@dataclass
class Col:
    header: str               # 3행 머리글(LIMS 이름) — QMS가 이 이름으로 인식
    unit: str = ""            # 4행 단위·형식
    req: bool = False         # 필수(주황)
    codes: str | None = None  # 목록 선택: CODE_TABLE 구분 이름
    key: str = ""             # QMS 내부 항목(항목사전용)
    desc: str = ""            # 항목사전 설명
    width: int = 0


@dataclass
class SheetSpec:
    name: str
    title: str
    guide: str
    cols: list[Col]
    screens: str              # 이 시트가 채우는 화면


def _ox(prefix: str, so3: str = "SO3") -> list[Col]:
    return [Col("CaO", "%", key=f"{prefix}_cao"), Col("SiO2", "%", key=f"{prefix}_sio2"),
            Col("Al2O3", "%", key=f"{prefix}_al2o3"), Col("Fe2O3", "%", key=f"{prefix}_fe2o3"),
            Col("MgO", "%", key=f"{prefix}_mgo"), Col(so3, "%", key=f"{prefix}_so3"),
            Col("K2O", "%", key=f"{prefix}_k2o"), Col("Na2O", "%", key=f"{prefix}_na2o")]


DATE = Col("채취일", "YYYY-MM-DD", True, key="timestamp", desc="시료 채취 날짜(필수). 1행 = 시료 1건", width=12)
TIME = Col("채취시간", "HH:MM", key="timestamp", desc="채취 시각. 같은 날 여러 시료를 구분(없으면 00:00)", width=9)
LINE_KR = Col("공정구분", "목록 선택", True, "킬른·원료밀", "line", "3K~7K(킬른 요입) / 3R/M~7R/M(원료밀 출구)", 11)
LINE_K = Col("공정구분", "목록 선택", True, "킬른(호기)", "line", "3K~7K", 10)


def _clinker(name: str, kilns: str) -> SheetSpec:
    return SheetSpec(
        name, f"③ 클링커 — {kilns} XRF·자유석회",
        "LIMS '크링커성적/수출크링커성적'과 같은 열입니다. 공정구분은 호기(3K~7K). Bogue 광물·LSF는 QMS가 계산합니다.",
        [DATE, TIME, LINE_K,
         Col("품명", "목록 선택", codes="클링커 종류(품명)", key="grade", desc="1종크링커 / 1/2종크링커 / 내수 / 수출클링커(참고)", width=13),
         Col("클링커 투입사일로", "목록 선택", codes="클링커 투입사일로", key="silo", desc="클링커를 넣은 사일로(추적용)", width=12),
         *_ox("clk"), Col("F-CaO(XRF)", "%", key="clk_fcao", desc="자유석회. 습식만 하면 머리글을 'F-CaO(습식)'으로"),
         Col("Cl(XRF)", "ppm", key="clk_cl"), Col("F(XRF)", "ppm", key="clk_f"), Col("온도", "℃", key="clk_temp")],
        "공정 모니터링(클링커) · SPC · 원인 진단 · 강도 예측 · 원료 배합 설계(실적 검증)")


def _ship(name: str, title: str, mid: list[Col], strengths: tuple[str, ...], tail: list[Col]) -> SheetSpec:
    return SheetSpec(
        name, title,
        "LIMS '시멘트성적' 내려받기 열과 같습니다. 품명(LIMS 이름)으로 품종을 자동 구분합니다. 비고에 출항 선박명을 적으면 선박·양하지로 연결됩니다.",
        [Col("채취일", "YYYY-MM-DD", True, key="date", desc="출하 시료 채취일", width=12),
         Col("LIMS LOT.NO", "문자", key="lot", desc="같은 LOT는 갱신(없으면 채취일+품명)", width=12),
         Col("품명", "목록 선택", True, "시멘트 종류(품명)", "product", "LIMS 품명 그대로(예: 1종시멘트 벌크)", 16),
         *mid,
         Col("LOI", "%", key="loi"), Col("Al2O3", "%", key="al2o3"), Col("MgO", "%", key="mgo"),
         Col("SO3(XRF)", "%", key="so3"), Col("Cl(XRF)", "ppm", key="cl"), Col("TA", "%", key="ta"),
         *([Col("C3S", "%"), Col("C2S", "%"), Col("C3A", "%"), Col("C4AF", "%")] if name.endswith("(3종)") else []),
         *[Col(f"압축강도{a}일(MPa)", "MPa", key=f"s{a}") for a in strengths],
         Col("밀도", "g/cm³"), Col("45㎛", "% 잔사", key="r45"), Col("Blaine", "cm²/g", key="blaine"),
         Col("주도", "%", key="cons"), Col("Flow", "mm", key="flow"), *tail,
         Col("비고", "문자", key="ships", desc="출항 선박명(여러 척이면 , 로 구분)", width=16)],
        "출하시멘트(품종별 KS·사내기준 판정·선박·양하지) · 그래프 보기 · 강도 예측(재령별)")


_SET = [Col("Autoclave", "%"), Col("초결시간(분)", "min"), Col("종결시간(HH:MM)", "HH:MM", key="fst")]

SHEETS: list[SheetSpec] = [
    SheetSpec(
        "원료", "① 원료(조합원료) — 원료밀 출구(R/M)·킬른 요입(K) 시료",
        "LIMS '조합원료성적'과 같은 열입니다. 공정구분: 킬른 요입 시료는 3K~7K, 원료밀 출구 시료는 3R/M~7R/M. "
        "LSF·SM·IM은 QMS가 계산하므로 넣지 않습니다. 배합비(%)는 원료밀(R/M) 행에만 넣으세요.",
        [DATE, TIME, LINE_KR,
         Col("품명", "목록 선택", codes="원료 종류(품명)", key="grade", desc="1종조합원료 / 1/2종조합원료(참고)", width=14),
         *_ox("rm"), Col("LOI", "%", key="rm_loi", desc="강열감량. 없으면 CaO·MgO로 추정해 클링커 예측에 사용"),
         Col("90㎛", "% 잔사", key="rm_r90"), Col("200㎛", "% 잔사", key="rm_r200"),
         Col("석회석(%)", "배합 %", key="rm_mx_lime"), Col("경석(%)", "배합 %", key="rm_mx_shale"),
         Col("규석(%)", "배합 %", key="rm_mx_silica"), Col("실리카(%)", "배합 %", key="rm_mx_silica2"),
         Col("석탄회(%)", "배합 %", key="rm_mx_ash"), Col("삼고(%)", "배합 %", key="rm_mx_sago"),
         Col("철질(%)", "배합 %", key="rm_mx_iron"), Col("주물사(%)", "배합 %", key="rm_mx_sand"),
         Col("투입량 계(%)", "배합 %", key="rm_mx_total")],
        "공정 모니터링(원료) · SPC · 그래프 보기 · 원인 진단 · 원료 배합 설계(실적 검증)"),
    SheetSpec(
        "원료 잔사", "① 원료 잔사 — 원료밀 출구 90·200㎛ 체 잔사",
        "잔사만 따로 내려받는 경우 이 시트에 넣습니다. 같은 채취일·시간·공정구분이면 '원료' 시트 행과 합쳐집니다.",
        [DATE, TIME, LINE_KR,
         Col("품명", "목록 선택", codes="원료 종류(품명)", key="grade", width=14),
         Col("90㎛", "% 잔사", key="rm_r90"), Col("200㎛", "% 잔사", key="rm_r200")],
        "공정 모니터링(원료) · 원인 진단(분말도) · f-CaO 예측"),
    SheetSpec(
        "킬른", "② 킬른 운전(DCS) — 호기별 시간값",
        "DCS 시간 평균값을 호기(3K~7K)별로 넣습니다. 측정하지 않는 항목은 비워 두세요.",
        [DATE, TIME, LINE_K,
         Col("원료 투입량(t/h)", "t/h", key="kiln_feed"), Col("석탄 투입량(t/h)", "t/h", key="coal_rate"),
         Col("소성대 온도(℃)", "℃", key="kiln_bzt", desc="파이로미터. 절대값은 측정 위치·방사율에 따라 다름"),
         Col("킬른 입구 O2(%)", "%", key="kiln_o2"), Col("킬른 입구 CO(%)", "%", key="kiln_co"),
         Col("주모터 부하(%)", "%", key="kiln_torque"), Col("하소로 출구 온도(℃)", "℃", key="kiln_calc_temp"),
         Col("2차 공기 온도(℃)", "℃", key="kiln_sec_air")],
        "공정 모니터링(소성) · 원인 진단(소성 조건 근거) · f-CaO 예측 · 6가크롬"),
    _clinker("클링커(6K,7K)", "6K·7K"),
    _clinker("클링커(4K,5K)", "3K·4K·5K"),
    SheetSpec(
        "클링커일일", "④ 클링커 일일 시료 — XRD 광물상·크롬(분석하는 경우만)",
        "하루 대표 시료의 XRD·크롬 결과. 호기별 1행. 분석하지 않으면 비워 두세요.",
        [Col("채취일", "YYYY-MM-DD", True, key="timestamp", desc="시료 날짜", width=12), LINE_K,
         Col("Alite(C3S)", "%", key="xrd_alite"), Col("Belite(C2S)", "%", key="xrd_belite"),
         Col("C3A", "%", key="xrd_c3a"), Col("C4AF", "%", key="xrd_c4af"), Col("F-CaO", "%", key="xrd_fcao"),
         Col("Periclase", "%", key="xrd_periclase"), Col("총크롬", "mg/kg", key="clk_cr"),
         Col("6가크롬", "mg/kg", key="clk_crvi")],
        "강도 예측(XRD) · 6가크롬 관리 · 원인 진단"),
    SheetSpec(
        "시멘트", "⑤ 시멘트(밀 출구) — 시멘트밀별 공정 시료",
        "LIMS '시멘트공정'과 같은 열·순서입니다(45㎛ → Blaine 순서 주의). 공정구분은 시멘트밀(2C/M~7C/M), 품명은 LIMS 이름 그대로.",
        [DATE, TIME, Col("공정구분", "목록 선택", True, "시멘트밀", "line", "2C/M~7C/M(시멘트밀)", 10),
         Col("품명", "목록 선택", True, "시멘트 종류(품명)", "product", "LIMS 품명(1종시멘트 벌크 등) → 1종·3종·고로슬래그 2종 자동", 16),
         Col("시멘트 투입사일로", "목록 선택", True, "시멘트 투입사일로", "silo", "시멘트를 넣은 사일로(추적용)", 12),
         Col("45㎛", "% 잔사", key="cem_r45"), Col("Blaine", "cm²/g", key="cem_blaine"),
         Col("SO3(XRF)", "%", key="cem_so3"), Col("LOI", "%", key="cem_loi"), Col("MgO", "%", key="cem_mgo"),
         Col("CaO", "%", key="cem_cao"), Col("SiO2", "%", key="cem_sio2"), Col("Al2O3", "%", key="cem_al2o3"),
         Col("Fe2O3", "%", key="cem_fe2o3"), Col("Na2O", "%", key="cem_na2o"), Col("K2O", "%", key="cem_k2o"),
         Col("Cl(XRF)", "ppm", key="cem_cl"), Col("F-CaO(XRF)", "%", key="cem_fcao"),
         Col("석회석비율", "%", key="cem_ls", desc="석회석 혼합률"),
         Col("시멘트밀 투입량(t/h)", "t/h", key="cem_mill_feed"),
         Col("원료온도", "℃", key="cem_mill_temp", desc="LIMS '원료온도' = 시멘트밀 출구 온도")],
        "공정 모니터링(시멘트) · SPC · 강도 예측 · 원인 진단 · 6가크롬"),
    _ship("출하시멘트(1종)", "⑦-1 출하시멘트 1종 — 제품 출하 로트(LIMS 시멘트성적 형식)",
          [Col("첨가재명칭", "문자", key="additive"), Col("석회석비율", "%", key="ls"), Col("첨가재비율", "%", key="addr")],
          ("1", "3", "7", "28"),
          [Col("공기량", "%"), *_SET, Col("간이 콘크리트 Flow", "mm"), Col("간이 콘크리트 단위수량", "kg/m³"),
           Col("간이 콘크리트 강도 3일", "MPa"), Col("간이 콘크리트 강도 7일", "MPa"), Col("간이 콘크리트 강도 28일", "MPa")]),
    _ship("출하시멘트(3종)", "⑦-2 출하시멘트 3종 — 제품 출하 로트(LIMS 시멘트성적 형식)", [], ("1", "3", "7", "28"), list(_SET)),
    _ship("출하시멘트(고로슬래그시멘트2종)", "⑦-3 출하시멘트 고로슬래그 2종 — 제품 출하 로트(LIMS 시멘트성적 형식)",
          [Col("첨가재명칭", "문자", key="additive"), Col("수량", "%"), Col("첨가재비율", "%", key="addr")],
          ("3", "7", "28"), list(_SET)),
    SheetSpec(
        "인수검사(원료)", "⑧-1 원료·부원료 인수검사",
        "LIMS 'RAW MATERIAL 인수검사'와 같은 열입니다. 합부판정을 비우면 📦 인수검사 화면의 사내 기준으로 자동 판정합니다.",
        [Col("로트번호", "문자", key="로트번호"), Col("검사일", "YYYY-MM-DD", key="검사일", width=12),
         Col("채취일", "YYYY-MM-DD", True, key="채취일", width=12), Col("거래처", "문자", key="거래처", width=16),
         Col("품명", "문자", True, key="품명", desc="원료명(규석·슬래그 등)", width=18), Col("수분", "%"), Col("LOI", "%"),
         Col("SiO2", "%"), Col("Al2O3", "%"), Col("Fe2O3", "%"), Col("CaO", "%"), Col("MgO", "%"), Col("SO3", "%"),
         Col("Na2O", "%"), Col("K2O", "%"), Col("Cl(XRF)", "ppm"), Col("Cl(습식)", "ppm"), Col("F(XRF)", "ppm"),
         Col("총크롬", "ppm"), Col("Cr", "ppm"), Col("염기도", ""), Col("철물", "%"),
         Col("합부판정", "합격/불합격", key="합부판정"), Col("비고", "문자", key="비고")],
        "📦 인수검사(사내 기준 자동 합부판정)"),
    SheetSpec(
        "인수검사(석고)", "⑧-2 석고 인수검사",
        "LIMS '석고류 인수검사'와 같은 열입니다.",
        [Col("로트번호", "문자"), Col("검사일", "YYYY-MM-DD", width=12), Col("채취일", "YYYY-MM-DD", True, width=12),
         Col("거래처", "문자", width=16), Col("품명", "문자", True, width=18), Col("결정수", "%"), Col("수분", "%"),
         Col("LOI", "%"), Col("SiO2", "%"), Col("Al2O3", "%"), Col("Fe2O3", "%"), Col("CaO", "%"), Col("MgO", "%"),
         Col("SO3", "%"), Col("P2O5(XRF)", "%"), Col("Cl(습식)", "ppm"), Col("순도", "%"), Col("F(습식)", "ppm"),
         Col("합부판정", "합격/불합격"), Col("비고", "문자")],
        "📦 인수검사"),
    SheetSpec(
        "인수검사(고령토)", "⑧-3 고령토 인수검사",
        "LIMS '삼척고령토'와 같은 열입니다. 공정구분은 고령토 채광 구분(기존/신설).",
        [Col("로트번호", "문자"), Col("검사일", "YYYY-MM-DD", width=12), Col("채취일", "YYYY-MM-DD", True, width=12),
         Col("품명", "문자", True, width=18), Col("채취시간", "HH:MM"), Col("공정구분", "목록 선택", codes="고령토 구분"),
         Col("품위구분", "문자"), Col("LOI", "%"), Col("수분", "%"),
         Col("SiO2", "%"), Col("Al2O3", "%"), Col("Fe2O3", "%"), Col("CaO", "%"), Col("MgO", "%"), Col("Na2O", "%"),
         Col("K2O", "%"), Col("합부판정", "합격/불합격"), Col("비고", "문자")],
        "📦 인수검사"),
]


def _width(text: str) -> int:
    """머리글 폭(한글은 2칸) + 여백, 9~20."""
    return max(9, min(20, sum(2 if ord(ch) > 0x1100 else 1 for ch in text) + 3))


def _code_ranges() -> dict[str, str]:
    """CODE_TABLE 구분 → '공정구분' 시트의 목록 범위(데이터 유효성 검사용). '킬른·원료밀'은 두 목록을 합친 열."""
    out, col = {}, 2
    for name, codes, *_ in CODE_TABLE:
        L = get_column_letter(col)
        out[name] = f"'공정구분'!${L}$4:${L}${3 + len(codes)}"
        col += 1
    L = get_column_letter(col)
    out["킬른·원료밀"] = f"'공정구분'!${L}$4:${L}${3 + len(KILNS) + len(RAW_MILLS)}"
    return out


def _guide(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "사용법"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 118
    ws["B2"] = "Blue365 QMS — 통합 업로드 양식 (공정별 시트 + 출하시멘트 + 인수검사)"
    ws["B2"].font = TITLE
    lines = [
        "",
        "■ 쓰는 순서",
        "  1) LIMS에서 공정별 결과를 내려받아, 이 파일의 맞는 시트 5행부터 붙여넣습니다(1~4행 제목·안내·머리글·단위는 그대로).",
        "  2) '공정구분'(호기·밀)·'품명'·'투입사일로' 칸은 목록에서 고릅니다. 예: 킬른 3K~7K / 원료밀 6R/M / 시멘트밀 3C/M.",
        "  3) 저장 후 [📥 데이터 관리]에 올리거나, 매일 자동 반영 파일(data\\Blue365_QMS.xlsx)로 저장합니다(06:00 자동 반영).",
        "",
        "■ 꼭 지킬 것",
        "  · 주황색 머리글 = 필수 칸입니다. '채취일'이 빈 행은 읽지 않습니다.",
        "  · 공정구분(호기)을 꼭 넣으세요. 같은 시각이라도 6K와 7K는 따로 저장·판정합니다(넣지 않으면 '구분 없음'으로 한데 묶입니다).",
        "  · 같은 채취일·시간·공정구분(+품명)이 이미 있으면 값이 있는 칸만 덮어씁니다. 늦게 나온 28일 강도를 나중에 채워도 됩니다.",
        "  · 0 또는 빈 칸 = 미시험으로 처리합니다(LIMS가 미시험을 0으로 내보내는 것과 같게). 배합비·석회석비율의 0은 실제 0%로 봅니다.",
        "  · 머리글 이름은 바꾸지 마세요(LIMS 이름 그대로여야 자동 연결). 열 순서는 LIMS와 같게 맞춰 두었습니다 — 통째로 붙여넣을 때 머리글과 값의 순서가 같은지 확인하세요(예: 시멘트 45㎛ → Blaine).",
        "  · 한 공정을 여러 시트로 나눠도 됩니다(예: 클링커(6K,7K)·클링커(4K,5K), 원료·원료 잔사). 시트 이름이 공정 이름으로 시작하면 모두 읽어 합칩니다.",
        "  · 계산값(LSF·SM·IM·C3S 등 Bogue)은 넣지 않습니다. QMS가 산화물로 계산합니다.",
        "",
        "■ 시트 구성과 연결 화면",
    ]
    lines += [f"  · {s.name:<10} {s.title.split(' — ')[0]}  →  {s.screens}" for s in SHEETS]
    lines += [
        "  · 공정구분   코드표(호기·밀·품종 목록). 공장 설비가 바뀌면 이 시트의 목록만 고치면 됩니다.",
        "  · 항목사전   모든 열의 뜻·단위·필수 여부·QMS 내부 항목.",
        "",
        "■ 데이터가 없는 공정은?  해당 시트를 비워 두면 됩니다(머리글만 있는 시트는 건너뜀).",
        "■ LIMS 원본 파일을 그대로 써도 됩니다.  예전처럼 LIMS에서 받은 파일을 그대로 폴더에 넣어도 자동으로 인식합니다.",
        "■ 물성 시트는 없습니다.  강도·응결·안정도는 출하시멘트 시트(품종별)에서 읽습니다. 예전 양식의 '물성' 시트를 추가해도 그대로 읽습니다.",
        "",
        "※ 근거: 시트 구성·코드 목록은 사용자가 고친 양식(2026-10-28)을 그대로 따랐고, 열 이름은 LIMS 원본 내려받기 이름입니다.",
    ]
    for i, line in enumerate(lines, 4):
        c = ws.cell(i, 2, line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("■"):
            c.font = Font(bold=True, size=11, color=NAVY)


def _codes(wb: Workbook) -> None:
    ws = wb.create_sheet("공정구분")
    ws["A1"] = "공정구분·품명 코드표 — 각 시트의 목록 선택에 쓰입니다(값을 고치면 목록도 바뀜)"
    ws["A1"].font = TITLE
    ws["A2"] = "아래 3행이 구분 이름, 4행부터 코드입니다. 칸을 추가하면 목록 범위도 함께 늘려야 하니, 설비 변경 시 관리자에게 알려 주세요."
    ws["A2"].font = NOTE
    col = 2
    for name, codes, *_ in CODE_TABLE:
        c = ws.cell(3, col, name)
        c.fill, c.font, c.alignment, c.border = HEAD, WHITE_B, CENTER, BOX
        for r, v in enumerate(codes, 4):
            cc = ws.cell(r, col, v)
            cc.alignment, cc.border = CENTER, BOX
        ws.column_dimensions[get_column_letter(col)].width = 16
        col += 1
    c = ws.cell(3, col, "킬른·원료밀(원료 시트용)")
    c.fill, c.font, c.alignment, c.border = HEAD, WHITE_B, CENTER, BOX
    for r, v in enumerate(KILNS + RAW_MILLS, 4):
        ws.cell(r, col, v).alignment = CENTER
    ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["A"].width = 3
    # 설명 표
    r0 = 4 + max(len(c[1]) for c in CODE_TABLE) + len(KILNS) + 2
    for j, h in enumerate(["구분", "코드", "의미", "쓰는 시트", "근거(사실/추정)"], 2):
        c = ws.cell(r0, j, h)
        c.fill, c.font, c.alignment, c.border = HEAD, WHITE_B, CENTER, BOX
    for i, (name, codes, mean, sheets, basis) in enumerate(CODE_TABLE, r0 + 1):
        for j, v in enumerate([name, ", ".join(codes), mean, sheets, basis], 2):
            c = ws.cell(i, j, v)
            c.alignment, c.border = LEFT, BOX
    for i in range(r0 + 1, r0 + 1 + len(CODE_TABLE)):
        ws.row_dimensions[i].height = 44


def _data_sheet(wb: Workbook, spec: SheetSpec, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet(spec.name)
    n = len(spec.cols)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n, 2))
    ws.cell(1, 1, spec.title).font = TITLE
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n, 2))
    g = ws.cell(2, 1, spec.guide + "  (주황 머리글 = 필수 · 4행 회색 = 단위, 읽지 않음 · 5행부터 입력)")
    g.font, g.alignment = NOTE, LEFT
    ws.row_dimensions[2].height = 30
    for j, col in enumerate(spec.cols, 1):
        h = ws.cell(3, j, col.header)
        h.fill, h.font, h.alignment, h.border = (REQ if col.req else HEAD), WHITE_B, CENTER, BOX
        u = ws.cell(4, j, col.unit)
        u.fill, u.font, u.alignment, u.border = UNIT, NOTE, CENTER, BOX
        ws.column_dimensions[get_column_letter(j)].width = col.width or _width(col.header)
        L = get_column_letter(j)
        if col.codes and col.codes in ranges:
            dv = DataValidation(type="list", formula1=f"={ranges[col.codes]}", allow_blank=True, showErrorMessage=True,
                                errorTitle="목록에서 선택", error=f"'{col.header}'은(는) 공정구분 시트의 목록 값만 넣을 수 있습니다.")
            dv.add(f"{L}5:{L}{4 + INPUT_ROWS}")
            ws.add_data_validation(dv)
        if col.unit == "YYYY-MM-DD":
            for r in range(5, 5 + INPUT_ROWS):
                ws.cell(r, j).number_format = "yyyy-mm-dd"
        elif col.unit == "HH:MM":
            for r in range(5, 5 + INPUT_ROWS):
                ws.cell(r, j).number_format = "@"          # 문자 '06:00' 그대로(엑셀 시각 변환 방지)
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = ws.cell(5, 4 if n > 4 else 2)


def _dictionary(wb: Workbook) -> None:
    ws = wb.create_sheet("항목사전")
    ws["A1"] = "항목사전 — 모든 열의 뜻·단위·필수 여부"
    ws["A1"].font = TITLE
    hdr = ["시트", "열 이름(머리글)", "필수", "단위·형식", "QMS 항목", "설명"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(3, j, h)
        c.fill, c.font, c.alignment, c.border = HEAD, WHITE_B, CENTER, BOX
    from . import shipment as ship
    from .standards import SpecRegistry
    reg = SpecRegistry()
    r = 4
    for spec in SHEETS:
        for col in spec.cols:
            it = (ship.BY_KEY.get(col.key) if spec.name.startswith("출하시멘트") else reg.get(col.key)) if col.key else None
            vals = [spec.name, col.header, "필수" if col.req else "", col.unit, (it.name if it else col.key) or "",
                    col.desc or (getattr(it, "description", "") if it else "") or ("" if col.key else "기타 항목(그래프 보기에서 선택 가능)")]
            for j, v in enumerate(vals, 1):
                c = ws.cell(r, j, v)
                c.alignment, c.border = LEFT, BOX
            r += 1
    for L, w in zip("ABCDEF", (14, 22, 6, 14, 26, 70)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = "A4"


def _print_setup(ws, rows: int | None = None) -> None:
    """인쇄: 가로 방향 · 한 쪽 너비 맞춤(빈 양식을 인쇄해 기록지로 써도 열이 잘리지 않게)."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if rows:
        ws.print_area = f"A1:{get_column_letter(ws.max_column)}{rows}"


def build_upload_template() -> bytes:
    """통합 업로드 양식(xlsx 바이트)."""
    wb = Workbook()
    _guide(wb)
    _codes(wb)
    ranges = _code_ranges()
    for spec in SHEETS:
        _data_sheet(wb, spec, ranges)
    _dictionary(wb)
    for ws in wb.worksheets:
        _print_setup(ws, 30 if ws.title in {s.name for s in SHEETS} else None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
