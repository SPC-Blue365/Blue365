"""시멘트 6가크롬(Cr⁶⁺) 물질수지 · 예측 · 환원제 투입량 · 저감 솔루션.

계산 구조
---------
1) 클링커 총 크롬(mg/kg-clk) = Σ 투입량_i(kg/t-clk, 건조) × 총Cr_i(mg/kg) × 잔류율_i(%)/100 / 1000
       + 내화물 마모분 = 마모량(kg/t-clk) × Cr₂O₃(%) × 6.842      (Cr₂O₃ → Cr 0.6842, 1 kg/t = 1,000 mg/kg)
   크롬은 비휘발성이라 투입분 대부분이 클링커에 남는다(석탄 Cr는 회분에 농축되어 95% 이상 이행 — 문헌).
2) 클링커 Cr⁶⁺(수용성) = 총 크롬 × 킬른 전환율(%)
   전환율: 문헌상 클링커 Cr의 약 8~20%가 6가로 전환(Costeri 2016; Lizarraga 2003, Hills & Johansen 2007 인용).
   산소 과잉·알칼리 과잉에서 증가(Hills; IJERPH 2022). 공장마다 크게 달라 ③ 실측 보정이 필수이다.
3) 시멘트 Cr⁶⁺(환원제 투입 전) = 클링커 비율 × 클링커 Cr⁶⁺ + Σ 혼합재 비율 × 혼합재 Cr⁶⁺
       + 분쇄매체 마모(g/t) × 매체 Cr(%)/100 × 밀 내 산화율(%)/100            (g/t = mg/kg)
4) 환원제   Cr⁶⁺ + 3Fe²⁺ → Cr³⁺ + 3Fe³⁺ ,   2Cr⁶⁺ + 3Sn²⁺ → 2Cr³⁺ + 3Sn⁴⁺
   화학양론 소요량(g/g-Cr⁶⁺) = 몰비 × 환원제 몰질량 / Cr 원자량(51.996)
       FeSO₄·7H₂O 3×278.01/52.00 = 16.04,  FeSO₄·H₂O 3×169.92/52.00 = 9.80,  SnSO₄ 1.5×214.77/52.00 = 6.20
   유효 환원능력(mg-Cr⁶⁺/kg) = 투입량(mg/kg) × 순도 × 잔존율 / (화학양론 × 현장 과잉계수)
   잔존율 = (1 − 월 열화율)^저장개월 × 열화계수(고온 투입 시)
   현장 과잉계수(경험칙): FeSO₄ 약 10배 이상, SnSO₄ 약 2~4배 — 공장 실측으로 보정.
5) 판정  국내: 환경부–시멘트업계 자율협약 20 mg/kg 이하(수용성 Cr⁶⁺, KS L 5221) — 법정 KS 기준 아님
         EU : REACH 부속서 XVII 47항 2 mg/kg(수용성 Cr⁶⁺, EN 196-10)
   ※ 두 기준은 시험법(추출 조건)이 달라 수치를 직접 비교할 수 없다.

※ 원료·연료 크롬 함량, 전환율, 과잉계수, 열화율, 단가의 기본값은 예시(추정)이다. 공장 분석값으로 교체할 것.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass

import numpy as np
import pandas as pd

M_CR = 51.996
CR2O3_TO_CR = 2 * M_CR / (2 * M_CR + 3 * 15.999)          # 0.6842
LIMIT_KR = 20.0
LIMIT_EU = 2.0
LIMIT_KR_TEXT = "국내 자율기준 20 mg/kg(환경부–업계 자율협약, KS L 5221)"
LIMIT_EU_TEXT = "EU 2 mg/kg(REACH 부속서 XVII 47항, EN 196-10)"

REDUCERS: dict[str, dict] = {
    "FeSO4·7H2O": {"label": "황산제1철 7수화물(FeSO₄·7H₂O)", "molar_mass": 278.01, "mol_per_cr": 3.0, "purity": 90.0,
                   "excess": 10.0, "loss_month": 15.0, "heat_sensitive": True, "price": 200.0},
    "FeSO4·H2O": {"label": "황산제1철 1수화물(FeSO₄·H₂O)", "molar_mass": 169.92, "mol_per_cr": 3.0, "purity": 88.0,
                  "excess": 10.0, "loss_month": 10.0, "heat_sensitive": True, "price": 300.0},
    "SnSO4": {"label": "황산제1주석(SnSO₄)", "molar_mass": 214.77, "mol_per_cr": 1.5, "purity": 97.0,
              "excess": 3.0, "loss_month": 3.0, "heat_sensitive": False, "price": 30000.0},
}


def stoich(key: str) -> float:
    """화학양론 소요량(g 환원제 / g Cr⁶⁺)."""
    r = REDUCERS[key]
    return r["mol_per_cr"] * r["molar_mass"] / M_CR


KILN_COLUMNS = ["use", "name", "group", "amount", "cr", "retention", "note"]
MILL_COLUMNS = ["use", "name", "pct", "cr", "crvi", "note"]
GROUPS = ["원료", "부원료", "연료", "기타"]


def default_kiln_inputs(mix_table: pd.DataFrame | None = None, mats: pd.DataFrame | None = None) -> pd.DataFrame:
    """킬른 투입물(클링커 1 t 기준). mix_table(배합 결과)·mats(원료 Cr)가 있으면 원료 행을 그 값으로 채운다."""
    rows = []
    if mix_table is not None and mats is not None and len(mix_table):
        cr_by = dict(zip(mats["name"].astype(str), pd.to_numeric(mats["Cr"], errors="coerce").fillna(0.0)))
        for _, r in mix_table.iterrows():
            grp = "원료" if r["구분"] in ("석회석", "실리카원", "알루미나원") else "부원료"
            rows.append((True, str(r["원료"]), grp, float(r["건조 투입량(kg/t-clk)"]), float(cr_by.get(str(r["원료"]), 0.0)),
                         100.0, "배합 계산 결과 연동"))
    else:
        rows += [
            (True, "석회석", "원료", 1260.0, 14.0, 100.0, "예시값 — 입고 분석으로 교체"),
            (True, "규석", "원료", 60.0, 20.0, 100.0, "예시값"),
            (True, "점토(셰일)", "원료", 124.0, 90.0, 100.0, "예시값(셰일 20~120 mg/kg 문헌 범위)"),
            (True, "철광석", "부원료", 20.0, 300.0, 100.0, "예시값(추정) — 철질원은 Cr 편차가 큼"),
            (True, "석탄재(플라이애시)", "부원료", 58.0, 190.0, 100.0, "예시값(문헌 약 190 mg/kg)"),
        ]
    rows += [
        (True, "유연탄", "연료", 115.0, 20.0, 98.0, "예시값(문헌 10~40 mg/kg) — Cr는 회분에 농축되어 대부분 클링커로 이행"),
        (True, "대체연료(폐합성수지)", "연료", 25.0, 40.0, 98.0, "예시값(추정) — 성적서·분석으로 확인"),
        (False, "하수슬러지(건조)", "연료", 10.0, 150.0, 98.0, "예시값(추정)"),
        (False, "제강슬래그(철질원 대체)", "부원료", 30.0, 5000.0, 100.0, "예시값(추정) — Cr₂O₃ 1~3% 수준이면 수천 mg/kg"),
    ]
    df = pd.DataFrame(rows, columns=KILN_COLUMNS)
    return df


def default_mill_inputs(gypsum_pct: float = 5.0, ls_pct: float = 4.0) -> pd.DataFrame:
    return pd.DataFrame([
        (True, "석고(천연·탈황)", gypsum_pct, 5.0, 0.1, "예시값(추정) — 탈황석고는 Lot별 확인"),
        (True, "석회석 미분말", ls_pct, 14.0, 0.0, "예시값 — 석회석 Cr는 대부분 3가"),
        (False, "기타 혼합재(슬래그 등)", 0.0, 500.0, 0.2, "예시값(추정)"),
    ], columns=MILL_COLUMNS)


@dataclass
class CrParams:
    conv_kiln: float = 12.0          # 킬른 전환율(%) — 수용성 Cr⁶⁺ / 클링커 총 Cr. 문헌 8~20%, 실측 보정 필수
    refr_wear: float = 0.20          # 내화물 마모(kg/t-clk). 현대식 킬른 0.2 미만(문헌)
    refr_cr2o3: float = 10.0         # 마그네시아-크롬 벽돌 Cr₂O₃(%) — 제품 사양서 확인(예시)
    refr_ret: float = 100.0          # 마모분의 클링커 이행률(%)
    media_wear: float = 30.0         # 분쇄매체 마모(g/t-시멘트) — 예시
    media_cr: float = 12.0           # 매체 Cr(%) — 고크롬 주철볼 12~30%, 단조강 1% 내외(예시)
    media_conv: float = 5.0          # 밀 내 Cr⁶⁺ 산화율(%) — 추정
    reducer: str = "FeSO4·7H2O"
    purity: float = 90.0             # 환원제 순도(%)
    excess: float = 10.0             # 현장 과잉계수(화학양론 대비 배수)
    loss_month: float = 15.0         # 저장 중 월 열화율(%)
    months: float = 1.0              # 출하~사용까지 저장 기간(개월)
    mill_temp: float = 105.0         # 투입 지점 온도(℃) — 밀 투입 시 밀 출구 온도
    after_mill: bool = False         # True: 밀 출구 이후(저온부) 투입
    dose: float = 0.9                # 현재 투입량(kg/t-시멘트) — 예시
    price: float = 200.0             # 원/kg
    target: float = 10.0             # 사내 목표(mg/kg)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict | None) -> "CrParams":
        if not d:
            return cls()
        keys = cls.__dataclass_fields__
        return cls(**{k: v for k, v in d.items() if k in keys})

    def with_reducer(self, key: str) -> "CrParams":
        r = REDUCERS[key]
        out = CrParams(**self.to_dict())
        out.reducer, out.purity, out.excess, out.loss_month, out.price = key, r["purity"], r["excess"], r["loss_month"], r["price"]
        return out


def heat_factor(temp: float, sensitive: bool, after_mill: bool) -> float:
    """고온 투입 시 환원제 열화 계수(추정). FeSO₄: 60 ℃ 이하 1.0 → 120 ℃ 0.6 선형. SnSO₄: 열에 안정(1.0)."""
    if after_mill or not sensitive:
        return 1.0
    return float(np.clip(1.0 - (temp - 60.0) / 60.0 * 0.4, 0.6, 1.0))


def reducer_retention(p: CrParams) -> float:
    r = REDUCERS.get(p.reducer, REDUCERS["FeSO4·7H2O"])
    keep = (1 - np.clip(p.loss_month, 0, 100) / 100) ** max(p.months, 0)
    return float(keep * heat_factor(p.mill_temp, r["heat_sensitive"], p.after_mill))


def capacity_per_kg(p: CrParams) -> float:
    """환원제 1 kg/t(= 1,000 mg/kg) 투입 시 제거되는 Cr⁶⁺(mg/kg)."""
    return 1000.0 * (p.purity / 100) * reducer_retention(p) / (stoich(p.reducer) * max(p.excess, 1e-9))


def required_dose(crvi0: float, target: float, p: CrParams) -> float:
    """목표 이하로 낮추는 데 필요한 투입량(kg/t-시멘트)."""
    cap = capacity_per_kg(p)
    return max(crvi0 - target, 0.0) / cap if cap > 0 else float("inf")


def after_reducer(crvi0: float, dose_kg_t: float, p: CrParams) -> float:
    return max(crvi0 - dose_kg_t * capacity_per_kg(p), 0.0)


# ── 물질수지 ────────────────────────────────────────────────────────────
@dataclass
class CrBalance:
    kiln: pd.DataFrame              # 투입원별 클링커 총Cr 기여
    mill: pd.DataFrame              # 혼합재·분쇄매체 기여
    total_clk: float                # 클링커 총 Cr(mg/kg)
    crvi_clk: float                 # 클링커 Cr⁶⁺(mg/kg)
    clinker_fraction: float
    crvi_cement0: float             # 시멘트 Cr⁶⁺(환원제 전)
    crvi_cement: float              # 시멘트 Cr⁶⁺(현재 환원제 투입 후)
    params: CrParams
    contrib: pd.DataFrame           # 시멘트 Cr⁶⁺(환원제 전) 기여도(파레토)

    def summary(self) -> dict:
        p = self.params
        return {
            "클링커 총 Cr(mg/kg)": self.total_clk, "킬른 전환율(%)": p.conv_kiln, "클링커 Cr⁶⁺(mg/kg)": self.crvi_clk,
            "클링커 비율(%)": self.clinker_fraction * 100, "시멘트 Cr⁶⁺ 환원 전(mg/kg)": self.crvi_cement0,
            "환원제": REDUCERS[p.reducer]["label"], "현재 투입량(kg/t)": p.dose,
            "시멘트 Cr⁶⁺ 환원 후(mg/kg)": self.crvi_cement,
            "필요 투입량 — 사내 목표(kg/t)": required_dose(self.crvi_cement0, p.target, p),
            "필요 투입량 — 국내 20(kg/t)": required_dose(self.crvi_cement0, LIMIT_KR, p),
            "필요 투입량 — EU 2(kg/t)": required_dose(self.crvi_cement0, LIMIT_EU, p),
        }


def _clean_kiln(df: pd.DataFrame) -> pd.DataFrame:
    d = df.copy()
    for c in KILN_COLUMNS:
        if c not in d:
            d[c] = {"use": True, "name": "", "group": "기타", "note": "", "retention": 100.0}.get(c, 0.0)
    for c in ("amount", "cr", "retention"):
        d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0.0)
    d["use"] = d["use"].fillna(False).astype(bool)
    return d


def _clean_mill(df: pd.DataFrame) -> pd.DataFrame:
    d = df.copy()
    for c in MILL_COLUMNS:
        if c not in d:
            d[c] = {"use": True, "name": "", "note": ""}.get(c, 0.0)
    for c in ("pct", "cr", "crvi"):
        d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0.0)
    d["use"] = d["use"].fillna(False).astype(bool)
    return d


def cr_balance(kiln_inputs: pd.DataFrame, mill_inputs: pd.DataFrame, p: CrParams) -> CrBalance:
    k = _clean_kiln(kiln_inputs)
    k = k[k["use"]].copy()
    k["contribution"] = k["amount"] * k["cr"] * k["retention"] / 100 / 1000      # mg/kg-clk
    refr = {"use": True, "name": "내화물 마모(크롬계 벽돌)", "group": "내화물", "amount": p.refr_wear,
            "cr": p.refr_cr2o3 * 10000 * CR2O3_TO_CR, "retention": p.refr_ret,
            "note": f"Cr₂O₃ {p.refr_cr2o3:g}% × 0.684", "contribution": p.refr_wear * p.refr_cr2o3 * 10000 * CR2O3_TO_CR
            * p.refr_ret / 100 / 1000}
    k = pd.concat([k, pd.DataFrame([refr])], ignore_index=True)
    total = float(k["contribution"].sum())
    k["share"] = k["contribution"] / total * 100 if total > 0 else 0.0
    crvi_clk = total * p.conv_kiln / 100

    m = _clean_mill(mill_inputs)
    m = m[m["use"]].copy()
    cf = max(1 - float(m["pct"].sum()) / 100, 0.0)
    m["contribution"] = m["pct"] / 100 * m["crvi"]
    media = p.media_wear * p.media_cr / 100 * p.media_conv / 100
    crvi0 = cf * crvi_clk + float(m["contribution"].sum()) + media
    crvi_after = after_reducer(crvi0, p.dose, p)

    rows = [{"구분": "킬른", "투입원": r["name"], "그룹": r["group"],
             "기여(mg/kg)": cf * r["contribution"] * p.conv_kiln / 100} for _, r in k.iterrows()]
    rows += [{"구분": "시멘트밀", "투입원": r["name"], "그룹": "혼합재", "기여(mg/kg)": r["contribution"]}
             for _, r in m.iterrows()]
    rows.append({"구분": "시멘트밀", "투입원": "분쇄매체 마모", "그룹": "분쇄매체", "기여(mg/kg)": media})
    contrib = pd.DataFrame(rows).sort_values("기여(mg/kg)", ascending=False).reset_index(drop=True)
    contrib["비율(%)"] = contrib["기여(mg/kg)"] / crvi0 * 100 if crvi0 > 0 else 0.0
    contrib["누적(%)"] = contrib["비율(%)"].cumsum()
    return CrBalance(k, m, total, crvi_clk, cf, crvi0, crvi_after, p, contrib)


def tornado(kiln_inputs: pd.DataFrame, mill_inputs: pd.DataFrame, p: CrParams, rel: float = 0.3,
            top_sources: int = 5) -> pd.DataFrame:
    """주요 입력 ±rel 변화 시 시멘트 Cr⁶⁺(현재 환원제 투입 후)의 변화(민감도)."""
    base = cr_balance(kiln_inputs, mill_inputs, p)
    rows = []

    def run(ki=None, mi=None, pp=None) -> float:
        return cr_balance(ki if ki is not None else kiln_inputs, mi if mi is not None else mill_inputs,
                          pp if pp is not None else p).crvi_cement

    k = _clean_kiln(kiln_inputs)
    top = base.kiln[base.kiln["group"] != "내화물"].nlargest(top_sources, "contribution")["name"].tolist()
    for name in top:
        vals = []
        for f in (1 - rel, 1 + rel):
            kk = k.copy()
            kk.loc[kk["name"] == name, "cr"] *= f
            vals.append(run(ki=kk))
        rows.append({"입력": f"{name} 총Cr ±{rel * 100:.0f}%", "낮음": vals[0], "높음": vals[1]})
    for attr, label in (("conv_kiln", "킬른 전환율"), ("refr_wear", "내화물 마모량"), ("excess", "환원제 과잉계수"),
                        ("months", "저장 기간"), ("media_wear", "분쇄매체 마모")):
        vals = []
        for f in (1 - rel, 1 + rel):
            pp = CrParams(**p.to_dict())
            setattr(pp, attr, getattr(p, attr) * f)
            vals.append(run(pp=pp))
        rows.append({"입력": f"{label} ±{rel * 100:.0f}%", "낮음": vals[0], "높음": vals[1]})
    out = pd.DataFrame(rows)
    out["기준"] = base.crvi_cement
    out["영향폭"] = (out["높음"] - out["낮음"]).abs()
    return out.sort_values("영향폭", ascending=True).reset_index(drop=True)


# ── 실측 기반 보정 ──────────────────────────────────────────────────────
def pairs_from_store(store) -> pd.DataFrame:
    """클링커 일일 시료의 총 Cr·수용성 Cr⁶⁺ + 같은 날 킬른 O₂·클링커 Na₂Oeq·황산화도."""
    xrd = store.tables.get("xrd", pd.DataFrame())
    if len(xrd) == 0 or not {"clk_cr", "clk_crvi"} <= set(xrd.columns):
        return pd.DataFrame(columns=["date", "total_cr", "crvi", "o2", "na2oeq", "sd"])
    d = xrd[["timestamp", "clk_cr", "clk_crvi"]].dropna().copy()
    d["date"] = d["timestamp"].dt.normalize()
    kiln = store.tables.get("kiln", pd.DataFrame())
    if len(kiln) and "kiln_o2" in kiln:
        o2 = kiln.set_index("timestamp")["kiln_o2"].resample("1D").mean()
        d["o2"] = d["date"].map(o2)
    clk = store.tables.get("clinker", pd.DataFrame())
    if len(clk) and "clk_na2oeq" in clk:
        cd = clk.set_index("timestamp")[["clk_na2oeq", "clk_so3", "clk_na2o", "clk_k2o"]].resample("1D").mean()
        d["na2oeq"] = d["date"].map(cd["clk_na2oeq"])
        d["sd"] = d["date"].map(sulfatization_degree(cd["clk_so3"], cd["clk_na2o"], cd["clk_k2o"]))
    out = d.rename(columns={"clk_cr": "total_cr", "clk_crvi": "crvi"})
    for c in ("o2", "na2oeq", "sd"):
        if c not in out:
            out[c] = np.nan
    return out[["date", "total_cr", "crvi", "o2", "na2oeq", "sd"]].reset_index(drop=True)


def sulfatization_degree(so3, na2o, k2o):
    """황산화도 SD(%) = 100·SO₃ / (1.292·Na₂O + 0.850·K₂O). 알칼리 대비 황 균형 지표."""
    den = 1.292 * na2o + 0.850 * k2o
    return 100 * so3 / den


@dataclass
class ConvCalibration:
    n: int
    median: float                   # 전환율(%)
    q1: float
    q3: float
    coef: dict | None               # 회귀(전환율% ~ O₂ + Na₂Oeq), 없으면 None
    r2: float | None
    rmse: float | None

    def predict(self, o2: float | None = None, na2oeq: float | None = None) -> float:
        if self.coef and o2 is not None and na2oeq is not None and np.isfinite(o2) and np.isfinite(na2oeq):
            c = self.coef
            return float(max(c["intercept"] + c["o2"] * o2 + c["na2oeq"] * na2oeq, 0.1))
        return self.median


def calibrate_conversion(pairs: pd.DataFrame, min_reg: int = 8) -> ConvCalibration | None:
    d = pairs.copy()
    for c in ("total_cr", "crvi", "o2", "na2oeq"):
        if c in d:
            d[c] = pd.to_numeric(d[c], errors="coerce")
    d = d.dropna(subset=["total_cr", "crvi"])
    d = d[d["total_cr"] > 0]
    if len(d) == 0:
        return None
    ratio = d["crvi"] / d["total_cr"] * 100
    coef = r2 = rmse = None
    reg = d.assign(ratio=ratio).dropna(subset=["o2", "na2oeq"]) if {"o2", "na2oeq"} <= set(d.columns) else pd.DataFrame()
    if len(reg) >= min_reg and reg["o2"].std() > 0 and reg["na2oeq"].std() > 0:
        X = np.column_stack([np.ones(len(reg)), reg["o2"], reg["na2oeq"]])
        y = reg["ratio"].to_numpy(float)
        b, *_ = np.linalg.lstsq(X, y, rcond=None)
        e = y - X @ b
        ss = float(np.sum((y - y.mean()) ** 2)) or 1e-9
        coef = {"intercept": float(b[0]), "o2": float(b[1]), "na2oeq": float(b[2])}
        r2 = 1 - float(np.sum(e ** 2)) / ss
        rmse = float(np.sqrt(np.mean(e ** 2)))
    return ConvCalibration(len(d), float(ratio.median()), float(ratio.quantile(0.25)), float(ratio.quantile(0.75)),
                           coef, r2, rmse)


def reducer_effect_series(store, media_crvi: float = 0.2, gypsum_pct: float = 5.0) -> pd.DataFrame:
    """로트별 '예측 환원 전 Cr⁶⁺(클링커 실측 기반) − 제품 실측' = 환원제 실효 제거량(mg/kg)."""
    phy = store.tables.get("physical", pd.DataFrame())
    xrd = store.tables.get("xrd", pd.DataFrame())
    if len(phy) == 0 or "phy_crvi" not in phy or len(xrd) == 0 or "clk_crvi" not in xrd:
        return pd.DataFrame()
    daily = xrd.set_index("timestamp")["clk_crvi"].resample("1D").mean()
    lots = phy.dropna(subset=["phy_crvi"]).copy()
    if len(lots) == 0:
        return pd.DataFrame()
    day = lots["timestamp"].dt.normalize()
    full = pd.date_range(min(daily.index.min(), day.min()) - pd.Timedelta(days=3), max(daily.index.max(), day.max()))
    used = daily.reindex(full).shift(1).rolling(2, min_periods=1).mean()
    ls = pd.Series(np.nan, index=lots.index)
    cem = store.tables.get("cement", pd.DataFrame())
    if len(cem) and "cem_ls" in cem:
        ld = cem.groupby([cem["timestamp"].dt.normalize(), "product"])["cem_ls"].mean()
        ls = pd.Series([ld.get((dd, pr), np.nan) for dd, pr in zip(day, lots["product"])], index=lots.index)
    cf = 1 - ls.fillna(4.0) / 100 - gypsum_pct / 100
    before = cf.to_numpy() * used.reindex(day).to_numpy() + media_crvi
    out = pd.DataFrame({"timestamp": lots["timestamp"].to_numpy(), "product": lots["product"].to_numpy(),
                        "예측 환원 전": before, "제품 실측": lots["phy_crvi"].to_numpy()})
    out["실효 제거량"] = out["예측 환원 전"] - out["제품 실측"]
    return out.dropna(subset=["예측 환원 전"])


# ── 저감 솔루션 지식베이스 ─────────────────────────────────────────────
@dataclass(frozen=True)
class CrAction:
    group: str
    title: str
    mechanism: str
    effect: str
    side_effects: str
    how: str
    verify: str
    basis: str
    horizon: str          # 단기 | 근본


CR_ACTIONS: tuple[CrAction, ...] = (
    CrAction("환원제", "환원제 투입량 보정(예측 Cr⁶⁺ 연동)",
             "예측 환원 전 Cr⁶⁺와 목표값 차이만큼 투입량을 정해 과·부족 투입을 막는다.",
             "목표 대비 과잉 투입 원가 절감·초과 위험 감소(공장별 상이, 추정)", "과량 투입 시 원가↑·FeSO₄는 물 요구량·응결 영향",
             "⑵ 탭 필요 투입량 기준으로 정량공급기 설정, 일 1회 이상 보정", "로트별 KS L 5221 시험, 공급기 교정(실측 투입량)",
             "화학양론+경험칙", "단기"),
    CrAction("환원제", "투입 위치를 저온부로 이동(밀 출구 이후)",
             "FeSO₄는 고온(약 60~80 ℃ 이상)·습기에서 산화·탈수되어 환원 능력이 떨어진다.",
             "열화 감소분만큼 투입량 절감(추정)", "설비 개조 필요(투입 장치·혼합 균일성 확보)",
             "분리기 이후 또는 시멘트 이송 라인에서 투입, 혼합 균일성 확인", "투입 위치별 제품 Cr⁶⁺ 비교 시험", "문헌+경험칙", "근본"),
    CrAction("환원제", "SnSO₄(주석계) 등 안정형 환원제 검토",
             "SnSO₄는 FeSO₄보다 소요량이 적고 열·저장 안정성이 높다.",
             "투입량 약 1/10 수준 가능(제품·조건별 상이, 추정)", "단가 높음 — 총원가 비교 필요",
             "실험실·실기 비교시험 후 전환 검토", "저장 기간별 Cr⁶⁺ 재시험(1·2·3개월)", "문헌+제조사 자료", "근본"),
    CrAction("환원제", "저장·유통 기간 관리(선입선출)",
             "환원제 효과는 시멘트 저장 기간이 길수록 줄어든다(EU는 유효기간 표기 의무).",
             "장기 저장 로트 초과 위험 감소", "물류 운영 제약", "출하 사일로·포장 제품 선입선출, 저장 기간별 관리기준 설정",
             "저장 기간별 Cr⁶⁺ 재시험", "규격+문헌", "단기"),
    CrAction("원료·연료", "원·부원료·연료 Cr 입고 기준 설정",
             "크롬은 대부분 비휘발성이라 투입 총 Cr가 클링커 Cr로 이어진다(물질수지).",
             "고Cr 원료 1종 대체 시 클링커 총 Cr를 수십 mg/kg 낮출 수 있음(⑴ 기여도 참조, 추정)",
             "원료 조달 제약·원가", "Lot별 성적서·자체 분석(XRF/ICP), Cr 상한(mg/kg) 계약 조건화", "월간 Cr 물질수지 점검",
             "물질수지", "근본"),
    CrAction("원료·연료", "고Cr 부원료 사용 비율 조정(철질원·석탄재·슬러지)",
             "기여도 상위 투입원을 줄이거나 저Cr 대체재로 바꾼다.", "기여도 비례 감소(⑴ 파레토)",
             "배합(IM·SM) 변화 → 🧪 배합 화면에서 재계산 필요", "🧪 원료 배합 화면에서 대체 원료로 재최적화", "클링커 총 Cr·Cr⁶⁺",
             "물질수지", "단기"),
    CrAction("내화물", "크롬프리 내화물(마그네시아-스피넬) 전환",
             "마그네시아-크롬 벽돌 마모분이 클링커 Cr의 주요 원인이 될 수 있다.", "내화물 기여분 제거(⑴ 기여도 참조)",
             "단가↑, 단열성·내스폴링성 차이 — 소성대 적용 시 코팅 관리 필요", "정기 보수 시 구간별 전환 계획 수립",
             "내화물 원단위·클링커 총 Cr 추이", "문헌", "근본"),
    CrAction("킬른 운전", "과잉 산소 억제(O₂ 적정화)",
             "Cr³⁺ → Cr⁶⁺ 산화는 산화 분위기에서 촉진된다.", "전환율 감소(공장별 상이 — ⑶ 보정 회귀로 확인)",
             "CO 상승·소성 품질 저하 위험 — O₂ 하한 준수", "O₂ 목표를 품질 허용 범위 하단으로 운전", "O₂·CO·클링커 Cr⁶⁺/총Cr",
             "문헌", "단기"),
    CrAction("킬른 운전", "알칼리·황 균형 관리(황산화도)",
             "알칼리가 과잉이면 알칼리 크롬산염 형성으로 Cr⁶⁺가 늘어나는 경향이 있다.",
             "공장별 상이(추정)", "SO₃ 증가 시 프리히터 부착·코팅 문제", "원료·연료 알칼리 투입 관리, 바이패스 운전 검토",
             "클링커 Na₂Oeq·황산화도·Cr⁶⁺", "문헌", "근본"),
    CrAction("킬른 운전", "냉각대 환원 분위기 형성(2차 연료 투입) 검토",
             "냉각 구간에 연료를 일부 투입해 국부 환원 분위기를 만들면 수용성 Cr⁶⁺가 줄었다는 실기 연구가 있다.",
             "연구 사례 기준 — 공장 적용 효과는 시험 필요(추정)", "클링커 품질·연소 안정성 영향 검토 필요",
             "시험 운전 계획 수립(소성 품질 동시 확인)", "클링커 Cr⁶⁺, f-CaO, 강도", "문헌", "근본"),
    CrAction("분쇄", "분쇄매체 재질 관리",
             "고크롬 매체 마모분은 밀 내에서 일부 산화될 수 있다.", "기여도 작음(⑴ 참조, 추정)", "마모 저항성 변화",
             "매체 재질·마모량(g/t) 관리", "매체 원단위", "문헌+추정", "근본"),
    CrAction("품질관리", "시험 빈도·판정 체계 강화",
             "6가크롬은 원료 Lot 변경 시 급변할 수 있어 조기 감지가 중요하다.", "초과 출하 위험 감소",
             "시험 부하 증가", "클링커 Cr⁶⁺ 일 1회 + 제품 로트별 KS L 5221, 18 mg/kg 경고 → 20 mg/kg 출하 보류",
             "관리도(R1)·알림", "규격+경험칙", "단기"),
)


def cr_solution_markdown(bal: CrBalance) -> str:
    p = bal.params
    lines = ["### 규칙 기반 6가크롬 저감 솔루션", ""]
    lines.append(f"- 클링커 총 Cr **{bal.total_clk:.1f} mg/kg** × 전환율 {p.conv_kiln:.1f}% → 클링커 Cr⁶⁺ **{bal.crvi_clk:.1f}**, "
                 f"시멘트(환원 전) **{bal.crvi_cement0:.1f}**, 현재 환원제 {p.dose:.2f} kg/t 투입 후 **{bal.crvi_cement:.1f} mg/kg**")
    judge = ("🔴 국내 자율기준 20 mg/kg 초과 위험" if bal.crvi_cement > LIMIT_KR else
             "🟠 사내 목표 초과" if bal.crvi_cement > p.target else "🟢 사내 목표 이내")
    lines.append(f"- 판정: {judge} (EU 2 mg/kg 기준은 시험법이 달라 참고용)")
    need = required_dose(bal.crvi_cement0, p.target, p)
    lines.append(f"- 사내 목표 {p.target:g} mg/kg 달성 필요 투입량: **{need:.2f} kg/t** ({REDUCERS[p.reducer]['label']}, "
                 f"순도 {p.purity:g}%, 과잉계수 {p.excess:g}배, 잔존율 {reducer_retention(p) * 100:.0f}%)")
    top = bal.contrib.head(3)
    lines.append("- 기여도 상위: " + ", ".join(f"{r['투입원']} {r['비율(%)']:.0f}%" for _, r in top.iterrows()))
    lines.append("")
    for horizon in ("단기", "근본"):
        lines.append(f"#### {'단기 조치' if horizon == '단기' else '근본 대책'}")
        for a in CR_ACTIONS:
            if a.horizon == horizon:
                lines.append(f"- **[{a.group}] {a.title}** ({a.basis}): {a.mechanism} 효과: {a.effect}. 부작용: {a.side_effects}. "
                             f"실행: {a.how}. 확인: {a.verify}.")
        lines.append("")
    lines.append("※ 크롬 함량·전환율·과잉계수·열화율 기본값은 예시(추정)입니다. 실측 보정 후 사용하세요.")
    return "\n".join(lines)


def cr_workbook(bal: CrBalance, tor: pd.DataFrame | None, kiln_inputs: pd.DataFrame, mill_inputs: pd.DataFrame) -> bytes:
    """6가크롬 평가표(엑셀): 요약·투입원 기여·파레토·환원제 필요량·민감도·입력값."""
    import io

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = bal.params
    summary = pd.DataFrame([(k, v) for k, v in bal.summary().items()], columns=["항목", "값"])
    summary.loc[len(summary)] = ["판정 기준", f"{LIMIT_KR_TEXT} / {LIMIT_EU_TEXT} — 시험법 상이, 직접 비교 불가"]
    summary.loc[len(summary)] = ["주의", "크롬 함량·전환율·과잉계수·열화율 기본값은 예시(추정) — 실측 보정 후 사용"]
    rows = []
    for key in REDUCERS:
        pp = p.with_reducer(key) if key != p.reducer else p
        for lab, tgt in ((f"사내 목표 {p.target:g}", p.target), ("국내 자율기준 20", LIMIT_KR), ("EU 2(참고)", LIMIT_EU)):
            d = required_dose(bal.crvi_cement0, tgt, pp)
            rows.append({"환원제": REDUCERS[key]["label"], "목표": lab, "투입량(kg/t)": d, "원가(원/t)": d * pp.price,
                         "화학양론(g/g)": stoich(key), "과잉계수": pp.excess, "잔존율(%)": reducer_retention(pp) * 100})
    kiln = bal.kiln[["name", "group", "amount", "cr", "retention", "contribution", "share"]].rename(columns={
        "name": "투입원", "group": "구분", "amount": "투입량(kg/t-clk)", "cr": "총Cr(mg/kg)", "retention": "잔류율(%)",
        "contribution": "클링커 기여(mg/kg)", "share": "비율(%)"})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="요약", index=False)
        kiln.round(3).to_excel(xw, sheet_name="클링커 총Cr 기여", index=False)
        bal.contrib.round(3).to_excel(xw, sheet_name="시멘트 Cr6 파레토", index=False)
        pd.DataFrame(rows).round(4).to_excel(xw, sheet_name="환원제 필요량", index=False)
        if tor is not None and len(tor):
            tor.round(3).to_excel(xw, sheet_name="민감도", index=False)
        _clean_kiln(kiln_inputs).to_excel(xw, sheet_name="입력_킬른", index=False)
        _clean_mill(mill_inputs).to_excel(xw, sheet_name="입력_시멘트밀", index=False)
        pd.DataFrame([(k, v) for k, v in p.to_dict().items()], columns=["매개변수", "값"]).to_excel(
            xw, sheet_name="입력_조건", index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A5F")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="맑은 고딕", size=10)
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width * 1.5), 70)
            ws.freeze_panes = "A2"
    return buf.getvalue()
