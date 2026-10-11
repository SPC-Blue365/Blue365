"""LIMS(실험실정보관리시스템) 연동 — LIMS 시험 결과를 QMS 모니터링 데이터로 자동 반영한다.

연동 방식(공장 LIMS 제품·보안 정책에 맞게 선택)
  sql   LIMS DB 직접 조회(읽기 전용 계정 권장) — SQLAlchemy URL (MSSQL·Oracle·PostgreSQL·MySQL·SQLite 등)
  rest  LIMS REST API(JSON) — Bearer 토큰 인증
  file  LIMS 내보내기 파일(CSV·XLSX, 긴 형식) 폴더
  demo  시연용 LIMS DB(data/lims_demo.db, SQLite) — 데모 데이터에서 생성

표준 형식(긴 형식) — 한 행 = 시험 결과 1건
  sampled_at    시료 채취 일시            sample_point  채취 지점(KILN_FEED, CLINKER, CLINKER_DAILY, CEMENT_MILL, CEMENT_LOT)
  test_code     시험 항목 코드(CAO, BLAINE, CS28 …)    value  결과값('<0.1' 같은 부등호 값은 숫자만 쓰고 건수를 보고)
  product       품종 코드(OPC, HES …, 시멘트·물성만)     status 결과 상태(승인된 결과만 반영)
  unit          단위                          updated_at    결과 확정·수정 일시(증분 동기화 기준)
  LIMS 열 이름이 다르면 column_map 으로 바꾸거나, SQL 조회문에서 AS 별칭으로 맞춘다.
가로형  한 행 = 시료 1건, 열 = 시험항목(엑셀 성적서형)도 받는다 — test_code·value 열이 없으면 항목 열을 자동으로 긴 형식으로 바꾼다
        (열 이름 = 시험코드). 시료번호·비고처럼 시험이 아닌 열은 column_map 에서 "-" 로 지정해 제외한다.
채취 지점 열이 없으면 파일 이름의 지점 코드(예: CLINKER_1010.xlsx) → default_sample_point 순서로 채운다.

매핑   (sample_point, test_code) → (QMS 테이블, 열, 환산계수, 오프셋)  — 화면에서 수정, data/lims.json 에 저장
동기화 워터마크(updated_at 최댓값) 이후 결과만 조회 → 승인 필터 → 매핑 → 넓은 형식 → 열 단위 병합 저장(store.save_raw)
       → 늦게 나온 28일 강도도 같은 로트 행에 채워지고, 먼저 들어온 값은 지워지지 않는다.
비밀값 DB 비밀번호·API 토큰은 환경변수 QMS_LIMS_PASSWORD / QMS_LIMS_TOKEN 또는 .streamlit/secrets.toml [lims]
       (password, token) 에 둔다. 설정 파일(lims.json)·저장소에는 저장하지 않는다.
"""

from __future__ import annotations

import io
import json
import os
import re
import sqlite3
from collections import Counter
from contextlib import closing
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import quote_plus

import numpy as np
import pandas as pd

from .standards import DATA_DIR, TABLES
from .store import DB_PATH, RAW_COLUMNS, TEXT_COLUMNS, load_raw, read_csv_bytes, save_raw

LIMS_CONFIG_PATH = DATA_DIR / "lims.json"
LIMS_LOG_PATH = DATA_DIR / "lims_log.jsonl"
LIMS_DEMO_DB = DATA_DIR / "lims_demo.db"
SECRETS_PATH = Path(__file__).resolve().parent.parent / ".streamlit" / "secrets.toml"
STD_COLUMNS = ["sampled_at", "sample_point", "test_code", "value", "product", "status", "unit", "updated_at"]
FILE_COL = "__file__"                  # 내보내기 파일 이름(채취 지점 추론용 내부 열)
SKIP_MARKS = {"-", "제외", "skip"}     # column_map 에서 이 값으로 지정한 열은 버린다(시료번호·비고 등)
MODES = {"demo": "데모 LIMS(시연용 SQLite)", "sql": "DB 직접 조회(SQL)", "rest": "REST API", "file": "내보내기 파일 폴더"}

SAMPLE_POINTS = {"raw_meal": "KILN_FEED", "clinker": "CLINKER", "xrd": "CLINKER_DAILY", "cement": "CEMENT_MILL",
                 "physical": "CEMENT_LOT"}
# QMS 열 → LIMS 시험코드(기본값 예시 — 공장 LIMS 코드표로 교체)
TEST_CODES = {
    "rm_cao": "CAO", "rm_sio2": "SIO2", "rm_al2o3": "AL2O3", "rm_fe2o3": "FE2O3", "rm_mgo": "MGO", "rm_so3": "SO3",
    "rm_k2o": "K2O", "rm_na2o": "NA2O", "rm_r90": "R90", "rm_r200": "R200", "rm_loi": "LOI",
    "clk_cao": "CAO", "clk_sio2": "SIO2", "clk_al2o3": "AL2O3", "clk_fe2o3": "FE2O3", "clk_mgo": "MGO",
    "clk_so3": "SO3", "clk_k2o": "K2O", "clk_na2o": "NA2O", "clk_fcao": "FCAO", "clk_lw": "LW",
    "xrd_alite": "ALITE", "xrd_belite": "BELITE", "xrd_c3a": "C3A", "xrd_c4af": "C4AF", "xrd_fcao": "FCAO_XRD",
    "xrd_periclase": "PERICLASE", "clk_cr": "CR_TOTAL", "clk_crvi": "CR6",
    "cem_blaine": "BLAINE", "cem_r45": "R45", "cem_so3": "SO3", "cem_loi": "LOI", "cem_mgo": "MGO", "cem_ls": "LS_ADD",
    "phy_ist": "IST", "phy_fst": "FST", "phy_autoclave": "AUTOCLAVE", "phy_s1": "CS1", "phy_s3": "CS3", "phy_s7": "CS7",
    "phy_s28": "CS28", "phy_crvi": "CR6", "lab_temp": "LAB_TEMP", "lab_rh": "LAB_RH", "lab_cure_temp": "CURE_TEMP",
    "sand_lot": "SAND_LOT", "operator": "OPERATOR",
}
# DCS 운전값(밀 투입량·온도 등)은 LIMS 대상이 아니므로 매핑하지 않는다.
DCS_ONLY = {"cem_mill_feed", "cem_mill_temp"}

DEFAULT_QUERY = """SELECT s.sampled_at   AS sampled_at,
       s.sample_point AS sample_point,
       r.test_code    AS test_code,
       r.value        AS value,
       s.product      AS product,
       r.status       AS status,
       r.unit         AS unit,
       r.approved_at  AS updated_at
FROM lims_result r
JOIN lims_sample s ON s.sample_id = r.sample_id
WHERE r.approved_at > :since
ORDER BY r.approved_at"""


def default_mapping() -> list[dict]:
    rows = []
    for table, cols in RAW_COLUMNS.items():
        if table not in SAMPLE_POINTS:
            continue
        for c in cols:
            if c == "product" or c in DCS_ONLY or c not in TEST_CODES:
                continue
            rows.append({"sample_point": SAMPLE_POINTS[table], "test_code": TEST_CODES[c], "table": table, "column": c,
                         "factor": 1.0, "offset": 0.0})
    return rows


@dataclass
class LimsConfig:
    enabled: bool = False
    mode: str = "demo"
    sql_url: str = ""                  # 예: mssql+pyodbc://qms_reader:{password}@LIMS-DB/LIMS?driver=ODBC+Driver+18+for+SQL+Server
    sql_query: str = DEFAULT_QUERY
    rest_url: str = ""
    rest_records_path: str = "data"    # 응답 JSON에서 결과 목록 위치(점 표기, 예: result.items)
    rest_since_param: str = "updated_since"
    file_dir: str = ""
    default_sample_point: str = ""     # 채취 지점 열이 없는 결과에 쓸 지점 코드(파일 이름에 지점 코드가 있으면 그것 우선)
    column_map: dict = field(default_factory=dict)     # LIMS 열 이름 → 표준 열 이름("-" 이면 제외)
    product_map: dict = field(default_factory=lambda: {"OPC": "1종", "HES": "3종", "1종": "1종", "3종": "3종"})
    status_ok: list = field(default_factory=lambda: ["APPROVED", "승인", "A", "FINAL"])
    lookback_days: int = 35
    watermark: str | None = None
    mapping: list = field(default_factory=default_mapping)

    def to_dict(self) -> dict:
        return asdict(self)


def load_lims_config(path: Path = LIMS_CONFIG_PATH) -> LimsConfig:
    data: dict = {}
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            data = {}
    return LimsConfig(**{k: v for k, v in data.items() if k in LimsConfig.__dataclass_fields__})


def save_lims_config(cfg: LimsConfig, path: Path = LIMS_CONFIG_PATH) -> None:
    d = cfg.to_dict()
    if "{password}" not in d.get("sql_url", "") and re.search(r"://[^:/@]+:[^@{]+@", d.get("sql_url", "")):
        raise ValueError("DB 비밀번호를 URL에 직접 넣지 마세요. {password} 자리표시자를 쓰고 QMS_LIMS_PASSWORD 로 설정하세요.")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")


def _secret(name: str) -> str:
    env = {"password": "QMS_LIMS_PASSWORD", "token": "QMS_LIMS_TOKEN"}[name]
    if os.environ.get(env):
        return os.environ[env]
    if SECRETS_PATH.exists():
        try:
            import tomllib
            return str(tomllib.loads(SECRETS_PATH.read_text(encoding="utf-8")).get("lims", {}).get(name, ""))
        except Exception:  # noqa: BLE001 - 비밀 설정 파싱 실패 시 빈 값
            return ""
    return ""


def secret_status() -> dict[str, bool]:
    return {"password": bool(_secret("password")), "token": bool(_secret("token"))}


# ── 조회 ────────────────────────────────────────────────────────────────
def _since(cfg: LimsConfig, now: datetime | None = None) -> datetime:
    if cfg.watermark:
        try:
            return pd.Timestamp(cfg.watermark).to_pydatetime()
        except (ValueError, TypeError):
            pass
    return (now or datetime.now()) - timedelta(days=int(cfg.lookback_days))


def _engine_url(cfg: LimsConfig) -> str:
    if cfg.mode == "demo":
        return f"sqlite:///{LIMS_DEMO_DB}"
    url = cfg.sql_url
    if "{password}" in url:
        url = url.replace("{password}", quote_plus(_secret("password")))
    return url


def fetch_sql(cfg: LimsConfig, since: datetime) -> pd.DataFrame:
    url = _engine_url(cfg)
    if not url:
        raise ValueError("SQL 연결 주소(sql_url)가 비어 있습니다.")
    if url.startswith("sqlite:///"):
        path = Path(url[len("sqlite:///"):])
        if not path.exists():
            raise FileNotFoundError(f"LIMS DB 파일이 없습니다: {path}")
        with closing(sqlite3.connect(path)) as con:
            q = cfg.sql_query.replace(":since", "?")
            return pd.read_sql(q, con, params=(since.strftime("%Y-%m-%d %H:%M:%S"),))
    try:
        from sqlalchemy import create_engine, text
    except ImportError as exc:  # pragma: no cover - 선택 패키지
        raise RuntimeError("SQL 연동에는 SQLAlchemy와 DB 드라이버가 필요합니다(pip install sqlalchemy pyodbc 등).") from exc
    engine = create_engine(url, pool_pre_ping=True)
    try:
        with engine.connect() as conn:
            return pd.read_sql(text(cfg.sql_query), conn, params={"since": since})
    finally:
        engine.dispose()


def _dig(obj, path: str):
    for part in [p for p in path.split(".") if p]:
        obj = obj.get(part, []) if isinstance(obj, dict) else []
    return obj


def fetch_rest(cfg: LimsConfig, since: datetime, session=None) -> pd.DataFrame:
    if not cfg.rest_url:
        raise ValueError("REST API 주소(rest_url)가 비어 있습니다.")
    import requests

    sess = session or requests
    headers = {"Accept": "application/json"}
    token = _secret("token")
    if token:
        headers["Authorization"] = f"Bearer {token}"
    resp = sess.get(cfg.rest_url, params={cfg.rest_since_param: since.isoformat(timespec="seconds")}, headers=headers,
                    timeout=30)
    resp.raise_for_status()
    records = _dig(resp.json(), cfg.rest_records_path) if cfg.rest_records_path else resp.json()
    return pd.DataFrame(records if isinstance(records, list) else [])


def fetch_files(cfg: LimsConfig) -> pd.DataFrame:
    folder = Path(cfg.file_dir)
    if not cfg.file_dir or not folder.is_dir():
        raise ValueError(f"내보내기 폴더가 없습니다: {cfg.file_dir}")
    frames = []
    for f in sorted(folder.iterdir()):
        if f.name.startswith("~$") or not f.is_file():         # 엑셀이 열어 둔 파일의 잠금 파일
            continue
        if f.suffix.lower() == ".csv":
            df = read_csv_bytes(f.read_bytes())
        elif f.suffix.lower() in (".xlsx", ".xls"):
            df = pd.read_excel(f)
        else:
            continue
        df[FILE_COL] = f.stem                                    # 채취 지점 추론용(파일 이름)
        frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=STD_COLUMNS)


def fetch(cfg: LimsConfig, since: datetime, session=None) -> pd.DataFrame:
    if cfg.mode in ("sql", "demo"):
        return fetch_sql(cfg, since)
    if cfg.mode == "rest":
        return fetch_rest(cfg, since, session)
    if cfg.mode == "file":
        return fetch_files(cfg)
    raise ValueError(f"알 수 없는 연동 방식: {cfg.mode}")


# ── 변환 ────────────────────────────────────────────────────────────────
_NUM = re.compile(r"[-+]?\d+(?:[.,]\d+)?(?:[eE][-+]?\d+)?")


def parse_value(v) -> tuple[float, bool]:
    """결과값 → (숫자, 부등호 여부). '<0.1' → (0.1, True), '12,345' → 12345, 숫자 아님 → (nan, False)."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return float("nan"), False
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v), False
    s = str(v).strip()
    censored = s.startswith(("<", ">", "≤", "≥"))
    s2 = s.replace(",", "") if re.fullmatch(r"[<>≤≥]?\s*\d{1,3}(,\d{3})+(\.\d+)?", s) else s.replace(",", ".")
    m = _NUM.search(s2)
    return (float(m.group()) if m else float("nan")), censored


@dataclass
class SyncReport:
    mode: str
    since: str
    n_fetched: int = 0
    n_status_skipped: int = 0
    n_invalid: int = 0
    n_censored: int = 0
    n_unmapped: int = 0
    unmapped: list = field(default_factory=list)       # [(sample_point, test_code, 건수)]
    n_mapped: int = 0
    rows_by_table: dict = field(default_factory=dict)
    cells_new: int = 0
    cells_changed: int = 0
    watermark_after: str | None = None
    saved: bool = False
    errors: list = field(default_factory=list)
    messages: list = field(default_factory=list)
    wide: dict = field(default_factory=dict, repr=False)

    def summary_text(self) -> str:
        if self.errors:
            return "실패: " + "; ".join(self.errors)
        t = ", ".join(f"{TABLES[k]['label']} {v}행" for k, v in self.rows_by_table.items()) or "반영할 결과 없음"
        return (f"조회 {self.n_fetched}건 → 반영 {self.n_mapped}건({t}) · 신규 칸 {self.cells_new} · 변경 칸 {self.cells_changed} · "
                f"미승인 제외 {self.n_status_skipped} · 미매핑 {self.n_unmapped} · 숫자 아님 {self.n_invalid}")


def _infer_sample_point(d: pd.DataFrame, cfg: LimsConfig) -> pd.Series:
    """채취 지점 열이 없을 때: 파일 이름에 들어 있는 지점 코드(매핑표 기준, 긴 코드 우선) → 기본 채취 지점."""
    codes = sorted({str(m.get("sample_point", "")).strip().upper() for m in cfg.mapping or []} - {""}, key=len, reverse=True)
    sp = pd.Series(None, index=d.index, dtype=object)
    if FILE_COL in d.columns:
        sp = d[FILE_COL].map(lambda stem: next((c for c in codes if c in str(stem).upper()), None))
    default = (cfg.default_sample_point or "").strip().upper()
    return sp.fillna(default) if default else sp


def _to_long(d: pd.DataFrame, rep: SyncReport) -> pd.DataFrame:
    """가로형(한 행 = 시료 1건, 열 = 시험항목) → 긴 형식(한 행 = 시험 결과 1건). 열 이름이 시험코드가 된다."""
    ids = [c for c in ("sampled_at", "sample_point", "product", "status", "unit", "updated_at", FILE_COL) if c in d.columns]
    items = [c for c in d.columns if c not in ids]
    if not items:
        raise ValueError("시험 결과 열이 없습니다(가로형이면 시험항목 열이 있어야 합니다).")
    long = d.melt(id_vars=ids, value_vars=items, var_name="test_code", value_name="value")
    long = long[long["value"].notna() & (long["value"].astype(str).str.strip() != "")]
    rep.messages.append(f"가로형 결과(시료당 1행)로 인식해 항목 열 {len(items)}개를 시험 결과 {len(long)}건으로 바꿨습니다.")
    return long


def normalize(df: pd.DataFrame, cfg: LimsConfig, rep: SyncReport) -> pd.DataFrame:
    d = df.rename(columns={k: v for k, v in (cfg.column_map or {}).items() if k in df.columns}).copy()
    d = d.loc[:, [str(c).strip() not in SKIP_MARKS for c in d.columns]]      # column_map 에서 "-" 로 지정한 열 제외
    d.columns = [str(c).strip().lower() for c in d.columns]
    if "sample_point" not in d.columns:
        d["sample_point"] = _infer_sample_point(d, cfg)
        if d["sample_point"].isna().all():
            raise ValueError("채취 지점(sample_point) 열이 없습니다 — 열을 추가하거나, 파일 이름에 지점 코드(예: CLINKER_1010.xlsx)를 "
                             "넣거나, ① 연결 설정의 '기본 채취 지점'을 지정하세요.")
    if "test_code" not in d.columns and "value" not in d.columns and "sampled_at" in d.columns:
        d = _to_long(d, rep)
    missing = [c for c in ("sampled_at", "sample_point", "test_code", "value") if c not in d.columns]
    if missing:
        raise ValueError(f"LIMS 결과에 필수 열이 없습니다: {', '.join(missing)} (column_map 또는 SQL 별칭으로 맞추세요)")
    for c in ("product", "status", "unit", "updated_at"):
        if c not in d:
            d[c] = None
    d["sampled_at"] = pd.to_datetime(d["sampled_at"], errors="coerce", format="mixed")
    d["updated_at"] = pd.to_datetime(d["updated_at"], errors="coerce", format="mixed")
    d = d[d["sampled_at"].notna() & d["sample_point"].notna()].copy()
    d["sample_point"] = d["sample_point"].astype(str).str.strip().str.upper()
    d["test_code"] = d["test_code"].astype(str).str.strip().str.upper()
    if cfg.status_ok:
        ok = {str(s).strip().upper() for s in cfg.status_ok}
        st = d["status"].astype(str).str.strip().str.upper()
        keep = d["status"].isna() | st.isin(ok)
        rep.n_status_skipped = int((~keep).sum())
        d = d[keep].copy()
    pm = {str(k).strip().upper(): v for k, v in (cfg.product_map or {}).items()}
    d["product"] = d["product"].map(lambda p: pm.get(str(p).strip().upper(), p) if pd.notna(p) and str(p).strip() else None)
    return d


def apply_mapping(d: pd.DataFrame, cfg: LimsConfig, rep: SyncReport) -> dict[str, pd.DataFrame]:
    mp = pd.DataFrame(cfg.mapping or [])
    if len(mp) == 0:
        rep.n_unmapped = len(d)
        return {}
    mp = mp.copy()
    mp["sample_point"] = mp["sample_point"].astype(str).str.strip().str.upper()
    mp["test_code"] = mp["test_code"].astype(str).str.strip().str.upper()
    mp["factor"] = pd.to_numeric(mp.get("factor", 1.0), errors="coerce").fillna(1.0)
    mp["offset"] = pd.to_numeric(mp.get("offset", 0.0), errors="coerce").fillna(0.0)
    m = d.merge(mp[["sample_point", "test_code", "table", "column", "factor", "offset"]],
                on=["sample_point", "test_code"], how="left")
    unm = m[m["table"].isna()]
    rep.n_unmapped = len(unm)
    rep.unmapped = [(sp, tc, int(n)) for (sp, tc), n in Counter(zip(unm["sample_point"], unm["test_code"])).most_common(20)]
    m = m[m["table"].notna()].copy()
    m = m[m["table"].isin(RAW_COLUMNS) & m.apply(lambda r: r["column"] in RAW_COLUMNS.get(r["table"], []), axis=1)] \
        if len(m) else m
    text = m["column"].isin(TEXT_COLUMNS)
    parsed = m["value"].map(parse_value)
    m["num"] = [p[0] for p in parsed]
    m["censored"] = [p[1] for p in parsed]
    m.loc[~text, "num"] = m.loc[~text, "num"] * m.loc[~text, "factor"] + m.loc[~text, "offset"]
    bad = (~text) & m["num"].isna()
    rep.n_invalid = int(bad.sum())
    rep.n_censored = int(m.loc[~text, "censored"].sum())
    m = m[~bad].copy()
    text = m["column"].isin(TEXT_COLUMNS)
    m["val"] = m["num"].astype(object)
    m.loc[text, "val"] = m.loc[text, "value"].astype(str)
    m = m.sort_values("updated_at", na_position="first", kind="stable")
    out: dict[str, pd.DataFrame] = {}
    for table, g in m.groupby("table"):
        by_product = TABLES[table]["by_product"]
        if by_product:
            miss = g["product"].isna() | ~g["product"].isin(["1종", "3종"] + list(set((cfg.product_map or {}).values())))
            if miss.any():
                rep.messages.append(f"[{TABLES[table]['label']}] 품종을 알 수 없는 결과 {int(miss.sum())}건 제외(product_map 확인)")
            g = g[~miss]
        if len(g) == 0:
            continue
        keys = ["sampled_at"] + (["product"] if by_product else [])
        # 같은 시료·항목이 여러 번 나오면(재시험·수정) 가장 늦게 확정된 값을 쓴다
        wide = g.groupby(keys + ["column"], sort=False)["val"].last().unstack("column").reset_index()
        wide = wide.rename(columns={"sampled_at": "timestamp"})
        wide.columns.name = None
        for c in wide.columns:
            if c not in TEXT_COLUMNS and c not in ("timestamp", "product"):
                wide[c] = pd.to_numeric(wide[c], errors="coerce")
        out[table] = wide
        rep.rows_by_table[table] = len(wide)
    rep.n_mapped = int(len(m))
    return out


def _cell_diff(before: dict[str, pd.DataFrame], wide: dict[str, pd.DataFrame]) -> tuple[int, int]:
    new = changed = 0
    for table, w in wide.items():
        b = before.get(table)
        keys = ["timestamp"] + (["product"] if TABLES[table]["by_product"] else [])
        cols = [c for c in w.columns if c not in keys]
        if b is None or len(b) == 0:
            new += int(w[cols].notna().sum().sum())
            continue
        b = b.copy()
        b["timestamp"] = pd.to_datetime(b["timestamp"])
        j = w.merge(b[keys + [c for c in cols if c in b.columns]], on=keys, how="left", suffixes=("", "__old"))
        for c in cols:
            old = j[f"{c}__old"] if f"{c}__old" in j else pd.Series(np.nan, index=j.index)
            has_new = j[c].notna()
            new += int((has_new & old.isna()).sum())
            if c in TEXT_COLUMNS:
                changed += int((has_new & old.notna() & (j[c].astype(str) != old.astype(str))).sum())
            else:
                diff = (pd.to_numeric(j[c], errors="coerce") - pd.to_numeric(old, errors="coerce")).abs() > 1e-9
                changed += int((has_new & old.notna() & diff).sum())
    return new, changed


def sync(cfg: LimsConfig, db_path: Path = DB_PATH, dry_run: bool = False, now: datetime | None = None,
         session=None, log_path: Path = LIMS_LOG_PATH, persist_watermark: bool = True) -> SyncReport:
    """LIMS → QMS 증분 동기화. dry_run=True 이면 저장·워터마크 갱신 없이 결과만 보고."""
    since = _since(cfg, now)
    rep = SyncReport(cfg.mode, since.strftime("%Y-%m-%d %H:%M:%S"))
    try:
        if cfg.mode == "demo" and not LIMS_DEMO_DB.exists():
            create_demo_lims(load_raw(db_path), LIMS_DEMO_DB)
            rep.messages.append("데모 LIMS DB를 새로 만들었습니다.")
        raw = fetch(cfg, since, session)
        rep.n_fetched = len(raw)
        if len(raw):
            d = normalize(raw, cfg, rep)
            if cfg.mode == "file" and d["updated_at"].notna().any():
                d = d[d["updated_at"].isna() | (d["updated_at"] > pd.Timestamp(since))]
            wide = apply_mapping(d, cfg, rep)
            rep.wide = wide
            if wide:
                rep.cells_new, rep.cells_changed = _cell_diff(load_raw(db_path), wide)
                if not dry_run:
                    save_raw(wide, path=db_path, replace=False)
                    rep.saved = True
            wm = d["updated_at"].max() if "updated_at" in d and d["updated_at"].notna().any() else None
            if wm is not None and pd.notna(wm):
                rep.watermark_after = pd.Timestamp(wm).strftime("%Y-%m-%d %H:%M:%S")
                if not dry_run and persist_watermark:
                    cfg.watermark = rep.watermark_after
                    save_lims_config(cfg)
    except Exception as exc:  # noqa: BLE001 - 연동 실패 원인을 화면·로그에 표시
        rep.errors.append(f"{type(exc).__name__}: {exc}")
    _log_sync(rep, dry_run, log_path)
    return rep


def _log_sync(rep: SyncReport, dry_run: bool, path: Path) -> None:
    entry = {k: v for k, v in asdict(rep).items() if k != "wide"}
    entry["ts"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    entry["dry_run"] = dry_run
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False, default=str) + "\n")
    except OSError:
        pass


def read_sync_log(path: Path = LIMS_LOG_PATH, limit: int = 100) -> list[dict]:
    if not path.exists():
        return []
    out = []
    for line in path.read_text(encoding="utf-8").splitlines()[-limit:]:
        try:
            out.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return out


def test_connection(cfg: LimsConfig, session=None) -> tuple[bool, str, pd.DataFrame]:
    """최근 lookback_days 범위를 조회해 연결·열 구성을 확인(저장 안 함)."""
    try:
        if cfg.mode == "demo" and not LIMS_DEMO_DB.exists():
            create_demo_lims(load_raw(), LIMS_DEMO_DB)
        df = fetch(cfg, datetime.now() - timedelta(days=int(cfg.lookback_days)), session)
    except Exception as exc:  # noqa: BLE001
        return False, f"연결 실패 — {type(exc).__name__}: {exc}", pd.DataFrame()
    if len(df) == 0:
        return True, f"연결 성공 — 최근 {cfg.lookback_days}일 결과 없음(기간·조회문 확인)", df
    rep = SyncReport(cfg.mode, "")
    try:
        d = normalize(df, cfg, rep)
    except ValueError as exc:
        return False, f"연결은 되었으나 형식을 맞춰야 합니다 — {exc}", df.drop(columns=[FILE_COL], errors="ignore").head(20)
    fmt = "가로형(시료당 1행)" if any("가로형" in m for m in rep.messages) else "긴 형식(결과당 1행)"
    return (True, f"연결 성공 — {len(df)}행 조회(최근 {cfg.lookback_days}일) · {fmt} · 시험 결과 {len(d)}건",
            df.drop(columns=[FILE_COL], errors="ignore").head(20))


# ── 매핑표 가져오기 ────────────────────────────────────────────────────────
_MAP_HEADERS = {
    "sample_point": {"sample_point", "채취지점", "회사채취지점코드", "채취지점코드", "lims채취지점코드"},
    "test_code": {"test_code", "시험코드", "회사시험코드", "lims시험코드"},
    "target": {"대상", "qms항목(테이블.열)", "table.column"},
    "table": {"table", "테이블"},
    "column": {"column", "qms항목코드", "qms열"},
    "factor": {"factor", "계수", "환산계수"},
    "offset": {"offset", "오프셋"},
}


def _norm_header(h) -> str:
    return re.sub(r"\s+", "", re.sub(r"^\[[^\]]*\]", "", str(h)).strip().lower())


def _cell(r: pd.Series, cols: dict, key: str) -> str:
    v = r[cols[key]] if key in cols else None
    return "" if v is None or (isinstance(v, float) and np.isnan(v)) or str(v).strip().lower() == "nan" else str(v).strip()


def parse_mapping_table(data: bytes, filename: str) -> tuple[list[dict], list[str]]:
    """매핑표(엑셀·CSV) → 매핑 목록과 경고. 'LIMS 연동 준비서' 엑셀의 시험항목 매핑표 시트, 화면에서 내려받은 매핑표 CSV를
    그대로 쓸 수 있다. 회사 채취지점·시험코드가 비어 있는 행은 건너뛴다."""
    if filename.lower().endswith(".csv"):
        df = read_csv_bytes(data, dtype=str)
    else:
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str)
        df = sheets[next((n for n in sheets if "매핑" in n), next(iter(sheets)))]
    cols: dict[str, str] = {}
    for c in df.columns:
        key = next((k for k, names in _MAP_HEADERS.items() if _norm_header(c) in names), None)
        if key and key not in cols:
            cols[key] = c
    if "sample_point" not in cols or "test_code" not in cols or not ({"target", "column"} & cols.keys()):
        raise ValueError("매핑표에 '채취 지점'·'시험코드'·'QMS 항목' 열이 필요합니다"
                         "(LIMS 연동 준비서의 매핑표 시트 또는 화면에서 내려받은 매핑표 CSV 형식).")
    owner = {c: t for t, cs in RAW_COLUMNS.items() for c in cs if c != "product"}
    rows: dict[tuple[str, str], dict] = {}
    warns: list[str] = []
    for i, r in df.iterrows():
        sp, tc = _cell(r, cols, "sample_point").upper(), _cell(r, cols, "test_code").upper()
        if not sp or not tc:
            continue
        target = _cell(r, cols, "target")
        if "." in target:
            table, column = target.split(".", 1)
        else:
            column = _cell(r, cols, "column")
            table = _cell(r, cols, "table") or owner.get(column, "")
        if column not in RAW_COLUMNS.get(table, []) or column == "product":
            warns.append(f"{i + 2}행: 알 수 없는 QMS 항목 '{table}.{column}' — 건너뜀")
            continue
        try:
            factor = float(_cell(r, cols, "factor") or 1.0)
            offset = float(_cell(r, cols, "offset") or 0.0)
        except ValueError:
            warns.append(f"{i + 2}행: 계수·오프셋이 숫자가 아님 — 건너뜀")
            continue
        if (sp, tc) in rows:
            warns.append(f"{i + 2}행: {sp}/{tc} 가 중복되어 아래 행 값으로 바꿈")
        rows[(sp, tc)] = {"sample_point": sp, "test_code": tc, "table": table, "column": column,
                          "factor": factor, "offset": offset}
    return list(rows.values()), warns


# ── LIMS 연동 준비서(엑셀) ──────────────────────────────────────────────────
_TEXT_NAMES = {"sand_lot": ("표준사 Lot", ""), "operator": ("시험 조", "")}
PREP_STEPS = [
    ["1", "연동 방식 결정", "LIMS 업체·IT와 DB 조회 / REST API / 내보내기 파일 중 가능한 방식을 정한다(② 시트)",
     "품질관리 + IT", "방식 결정", "1주"],
    ["2", "IT 요청", "읽기 전용 계정·결과 뷰(④ 시트 형식)·QMS PC → DB 서버 방화벽 허용을 요청한다(③ 시트)",
     "IT · LIMS 업체", "접속 정보", "1~3주"],
    ["3", "시험항목 매핑표 작성", "⑤ 시트에 회사 LIMS의 채취지점·시험코드를 적고, QMS [🔗 LIMS 연동] ② 탭 '매핑표 올리기'로 등록",
     "품질관리", "매핑표", "1~2일"],
    ["4", "연결 설정·연결 테스트", "① 탭에 연결 주소·조회문(또는 폴더) 입력, 비밀번호는 secrets.toml [lims] → '연결 테스트'",
     "품질관리(+IT)", "'연결 성공'", "1일"],
    ["5", "미리보기 → 동기화", "③ 탭 '미리보기'로 미매핑·미승인 건수를 확인한 뒤 '동기화 실행'", "품질관리", "반영 건수", "1일"],
    ["6", "값 대조 검증", "LIMS 화면 값과 QMS 값을 공정별 10건 이상 대조(⑦ 시트)", "품질관리", "체크리스트", "1~2일"],
    ["7", "자동 동기화", "5_MONITOR_TASK.bat → 1번(30분마다 동기화·이상 감지·알림)", "품질관리", "작업 등록", "10분"],
    ["8", "운영 전환", "1~2주 병행 운영(수기 입력과 비교) 후 수기 입력 중단", "품질관리", "운영 전환", "1~2주"],
]
PREP_MODES = [
    ["DB 직접 조회(SQL)", "LIMS DB 읽기 전용 계정, 결과 뷰(View), QMS PC → DB 서버 포트(예: 1433·1521) 허용",
     "30분 주기 자동 반영, 늦게 확정·수정된 결과도 자동 반영, 사람 손이 가지 않음",
     "IT 보안 승인 필요. DB 드라이버(MS SQL·Oracle·PostgreSQL·MySQL)는 설치 패키지에 포함", "IT가 DB 조회를 허용할 때(가장 권장)"],
    ["REST API", "LIMS 업체가 제공하는 API 주소·인증 토큰, 결과 JSON 형식 설명",
     "DB 구조와 무관, 업체가 지원하면 깔끔", "LIMS 제품이 API를 제공해야 함(업체 개발 비용 가능 — 추정)", "최신 LIMS·업체가 API를 제공할 때"],
    ["내보내기 파일", "LIMS의 엑셀·CSV 내보내기(자동 또는 수동) → 공유 폴더",
     "IT 개발 없이 바로 시작, 성적서형(시료당 1행) 엑셀도 그대로 사용",
     "사람이 내보내면 지연·누락 가능. 매번 폴더 전체를 다시 읽으므로 오래된 파일은 정리",
     "DB 승인 전 임시 운영, 망분리로 DB 접근이 막혔을 때"],
]
PREP_IT = [
    ["공통", "LIMS 제품명·버전", "", "(업체명) LIMS v3.2", ""],
    ["공통", "담당자(IT·LIMS 업체)", "", "IT 홍길동 / 업체 김OO", ""],
    ["DB", "DB 종류·버전", "", "MS SQL Server 2019 / Oracle 19c", "드라이버는 QMS 설치 패키지에 포함"],
    ["DB", "서버 주소(호스트명 또는 IP)", "", "LIMS-DB01 / 10.10.20.5", ""],
    ["DB", "포트", "", "1433(MS SQL) / 1521(Oracle)", ""],
    ["DB", "DB 이름(Oracle은 서비스명)", "", "LIMSDB / ORCL", ""],
    ["DB", "결과 뷰(View) 이름", "", "V_QMS_RESULT", "④ 시트 형식으로 생성 요청"],
    ["DB", "읽기 전용 계정", "", "qms_reader", "뷰 SELECT 권한만"],
    ["DB", "인증 방식", "", "SQL 로그인 / Windows 인증", "비밀번호는 QMS PC의 secrets.toml 에만 보관"],
    ["네트워크", "QMS 공유 PC의 IP(출발지)", "", "10.10.30.21", "9_CHECK.bat 결과에 표시"],
    ["네트워크", "방화벽 허용", "", "10.10.30.21 → 10.10.20.5 : 1433/TCP", ""],
    ["코드", "결과 확정 상태 코드", "", "APPROVED / 승인 / F", "이 상태의 결과만 반영(⑥ 시트)"],
    ["코드", "품종 코드", "", "OPC = 1종, HES = 3종", "⑥ 시트"],
    ["코드", "채취 지점·시험항목 코드표", "", "CLINKER / FCAO …", "⑤ 시트에 기입"],
    ["REST(대안)", "API 주소", "", "https://lims.company.local/api/results", ""],
    ["REST(대안)", "인증 토큰", "", "Bearer 토큰", "secrets.toml [lims] token"],
    ["REST(대안)", "증분 조회 매개변수·응답 예시", "", "updated_since=2026-10-01T00:00:00", "결과 목록 위치(예: data)"],
    ["파일(대안)", "내보내기 폴더(공유 폴더)", "", r"\\FILESRV\LIMS_EXPORT", "QMS PC에서 읽기 가능해야 함"],
    ["파일(대안)", "내보내기 주기·형식", "", "1시간마다, 엑셀(시료당 1행)", "파일 이름에 채취지점 코드 포함 권장"],
]
PREP_VIEW = [
    ["sampled_at", "시료 채취 일시", "날짜시간(datetime)", "필수", "2026-10-07 10:00"],
    ["sample_point", "채취 지점 코드", "문자", "필수(파일 방식은 파일 이름으로 대체 가능)", "CLINKER"],
    ["test_code", "시험항목 코드", "문자", "필수", "FCAO"],
    ["value", "결과값", "숫자 또는 문자('<0.1' 허용)", "필수", "1.31"],
    ["product", "품종 코드", "문자", "시멘트·물성 결과는 필수", "OPC"],
    ["status", "결과 상태", "문자", "권장", "APPROVED"],
    ["unit", "단위", "문자", "선택", "%"],
    ["updated_at", "결과 확정·수정 일시(증분 조회 기준)", "날짜시간(datetime)", "권장(DB·API)", "2026-10-07 11:05"],
]
PREP_SQL = [
    ("MS SQL Server 뷰 예시(테이블·열 이름은 LIMS 제품마다 다름 — 업체 확인)", """CREATE VIEW dbo.V_QMS_RESULT AS
SELECT s.SAMPLED_DATE      AS sampled_at,
       s.SAMPLE_POINT_CODE AS sample_point,
       r.TEST_CODE         AS test_code,
       r.RESULT_VALUE      AS value,
       s.PRODUCT_CODE      AS product,
       r.STATUS            AS status,
       r.UNIT              AS unit,
       r.APPROVED_DATE     AS updated_at
FROM LIMS_RESULT r
JOIN LIMS_SAMPLE s ON s.SAMPLE_ID = r.SAMPLE_ID;
GRANT SELECT ON dbo.V_QMS_RESULT TO qms_reader;"""),
    ("Oracle 뷰 예시(날짜가 문자로 저장돼 있으면 TO_DATE 로 변환)", """CREATE OR REPLACE VIEW V_QMS_RESULT AS
SELECT TO_DATE(s.SAMPLED_DT, 'YYYYMMDDHH24MISS') AS sampled_at,
       s.SAMPLE_POINT_CODE AS sample_point,
       r.TEST_CODE AS test_code, r.RESULT_VALUE AS value, s.PRODUCT_CODE AS product,
       r.STATUS AS status, r.UNIT AS unit, r.APPROVED_DATE AS updated_at
FROM LIMS_RESULT r JOIN LIMS_SAMPLE s ON s.SAMPLE_ID = r.SAMPLE_ID;
GRANT SELECT ON V_QMS_RESULT TO qms_reader;"""),
    ("QMS 조회문(① 탭 '조회문'에 입력, :since = 마지막 동기화 시점)", """SELECT sampled_at, sample_point, test_code, value, product, status, unit, updated_at
FROM V_QMS_RESULT
WHERE updated_at > :since
ORDER BY updated_at"""),
]
PREP_URLS = [
    ["MS SQL(Windows 기본 드라이버, 설치 불필요)", "mssql+pyodbc://qms_reader:{password}@LIMS-DB01/LIMSDB?driver=SQL+Server"],
    ["MS SQL(ODBC Driver 18 설치 시)",
     "mssql+pyodbc://qms_reader:{password}@LIMS-DB01/LIMSDB?driver=ODBC+Driver+18+for+SQL+Server&TrustServerCertificate=yes"],
    ["MS SQL(ODBC 없이 — pymssql)", "mssql+pymssql://qms_reader:{password}@LIMS-DB01:1433/LIMSDB"],
    ["MS SQL(Windows 인증 — 비밀번호 없음)", "mssql+pyodbc://LIMS-DB01/LIMSDB?driver=SQL+Server&trusted_connection=yes"],
    ["Oracle(Oracle Client 설치 불필요)", "oracle+oracledb://qms_reader:{password}@10.10.20.5:1521/?service_name=ORCL"],
    ["PostgreSQL", "postgresql+psycopg2://qms_reader:{password}@LIMS-DB01:5432/limsdb"],
    ["MySQL·MariaDB", "mysql+pymysql://qms_reader:{password}@LIMS-DB01:3306/limsdb"],
]
PREP_CHECKS = [
    ["1", "연결", "① 탭 '연결 테스트'", "'연결 성공'과 결과 형식 표시"],
    ["2", "미매핑 시험코드", "③ 탭 '미리보기'의 미매핑 목록", "모니터링에 필요한 항목은 미매핑 0건"],
    ["3", "미승인 제외", "미리보기 요약의 '미승인 제외' 건수", "LIMS의 승인 대기 건수와 일치"],
    ["4", "값 대조", "LIMS 화면 값 ↔ QMS [📈 공정 모니터링] 값, 공정별 10건 이상", "100% 일치(단위 환산 포함)"],
    ["5", "품종 구분", "시멘트·물성 결과의 품종", "1종·3종이 올바르게 나뉨"],
    ["6", "늦게 확정된 결과", "28일 강도 확정 후 동기화", "같은 로트 행에 채워지고 기존 값 유지"],
    ["7", "수정된 결과", "LIMS에서 값 수정·재승인 후 동기화", "새 값으로 갱신('변경 칸' 1 이상)"],
    ["8", "자동 동기화", "5_MONITOR_TASK.bat → 4번(최근 실행 기록)", "30분마다 '조회 … 반영' 기록"],
    ["9", "알림", "기준을 벗어난 결과가 반영될 때", "[🚨 알림 센터]·메일 발송 확인"],
    ["10", "병행 운영", "1~2주 수기 입력과 비교", "차이 0건 확인 후 수기 입력 중단"],
]


def prep_workbook(registry=None) -> bytes:
    """LIMS 연동 준비서(엑셀) — 절차·방식 비교·IT 요청서·표준 뷰 정의·시험항목 매핑표·코드 변환·검증 체크리스트.

    매핑표 시트는 회사 코드를 적은 뒤 그대로 [🔗 LIMS 연동] ② 탭 '매핑표 올리기'에 올릴 수 있다.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    head_fill, head_font = PatternFill("solid", fgColor="1F4E79"), Font(bold=True, color="FFFFFF")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, header, rows, widths, inputs=(), note=None):
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = w
        for row in ws.iter_rows(min_row=1, max_row=len(rows) + 1):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row > 1 and cell.column in inputs:
                    cell.fill = input_fill
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.page_setup.orientation = "landscape"                 # 인쇄: 가로, 한 페이지 너비에 맞춤
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
        if note:
            ws.cell(len(rows) + 3, 1, note).font = Font(italic=True, color="595959")

    wb = Workbook()
    ws = wb.active
    ws.title = "① 연동 절차"
    sheet(ws, ["단계", "할 일", "내용", "담당", "확인·산출물", "소요(추정)"], PREP_STEPS, [6, 20, 70, 16, 14, 11],
          note="소요 기간은 일반적인 경우의 추정치입니다. 노란 칸은 입력하는 칸입니다(③·⑤·⑥·⑦ 시트).")
    sheet(wb.create_sheet("② 연동 방식"), ["방식", "필요한 것", "장점", "주의", "이런 경우 권장"], PREP_MODES,
          [18, 40, 40, 40, 28])
    sheet(wb.create_sheet("③ IT 요청서"), ["구분", "항목", "내용(입력)", "예시", "비고"], PREP_IT, [12, 28, 34, 38, 34],
          inputs=(3,))
    ws4 = wb.create_sheet("④ 표준 뷰 정의")
    sheet(ws4, ["열 이름", "의미", "형식", "필수 여부", "예시"], PREP_VIEW, [16, 34, 26, 34, 20])
    r = len(PREP_VIEW) + 3
    for title, sql in PREP_SQL:
        ws4.cell(r, 1, title).font = Font(bold=True, color="1F4E79")
        ws4.cell(r + 1, 1, sql).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
        ws4.row_dimensions[r + 1].height = 15 * (sql.count("\n") + 1)
        ws4.cell(r + 1, 1).font = Font(name="Consolas", size=10)
        r += 3
    ws4.cell(r, 1, "연결 주소(SQLAlchemy) 예시 — {password} 자리는 그대로 두고 비밀번호는 secrets.toml [lims] password 에")\
        .font = Font(bold=True, color="1F4E79")
    for i, (label, url) in enumerate(PREP_URLS, r + 1):            # 설명(A~B) | 연결 주소(C~E)
        ws4.cell(i, 1, label).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        ws4.cell(i, 3, url).font = Font(name="Consolas", size=10)
        ws4.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)
        ws4.row_dimensions[i].height = 30
    last = r + len(PREP_URLS) + 2
    ws4.cell(last, 1, "날짜는 문자(YYYYMMDDHHMISS)가 아니라 날짜형으로 주세요. 값을 수정·재승인하면 updated_at 이 갱신돼야 다시 반영됩니다. "
             "28일 강도처럼 늦게 확정되는 결과도 같은 시료(sampled_at)로 주세요.").font = Font(italic=True, color="595959")
    ws4.cell(last, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.merge_cells(start_row=last, start_column=1, end_row=last, end_column=5)
    ws4.row_dimensions[last].height = 32

    labels = {}
    for c in [m["column"] for m in default_mapping()]:
        if registry is not None and c in registry:
            labels[c] = (registry[c].name, registry[c].unit)
        else:
            labels[c] = _TEXT_NAMES.get(c, (c, ""))
    rows = []
    for m in default_mapping():
        name, unit = labels[m["column"]]
        note = "문자값" if m["column"] in TEXT_COLUMNS else ("늦게 확정돼도 같은 시료 일시로" if m["column"] == "phy_s28" else "")
        rows.append([TABLES[m["table"]]["label"], m["column"], name, unit, m["sample_point"], m["test_code"], "", "",
                     m["factor"], m["offset"], note])
    sheet(wb.create_sheet("⑤ 시험항목 매핑표"),
          ["공정 단계", "QMS 항목 코드", "항목명", "단위", "기본 채취 지점(예시)", "기본 시험코드(예시)",
           "[입력] 회사 채취지점 코드", "[입력] 회사 시험코드", "계수", "오프셋", "비고"],
          rows, [16, 16, 30, 9, 16, 16, 20, 18, 7, 7, 26], inputs=(7, 8, 9, 10),
          note="회사 코드를 적은 행만 등록됩니다. 단위가 다르면 계수·오프셋으로 환산(QMS 값 = LIMS 값 × 계수 + 오프셋, 예: g/kg → % 는 계수 0.1).")
    sheet(wb.create_sheet("⑥ 코드 변환"), ["구분", "LIMS 코드(입력)", "QMS 값", "비고"],
          [["품종", "OPC", "1종", "보통 포틀랜드 시멘트(예시)"], ["품종", "HES", "3종", "조강 포틀랜드 시멘트(예시)"],
           ["품종", "", "", ""], ["결과 상태", "APPROVED", "반영", "승인·확정된 결과만 반영"],
           ["결과 상태", "승인", "반영", ""], ["결과 상태", "PENDING", "제외", "승인 대기(예시)"], ["결과 상태", "", "", ""]],
          [12, 22, 12, 40], inputs=(2, 3),
          note="입력한 코드는 QMS [🔗 LIMS 연동] ① 탭의 '품종 코드 → QMS 품종', '반영할 결과 상태'에 넣습니다.")
    sheet(wb.create_sheet("⑦ 검증 체크리스트"), ["순서", "확인 항목", "방법", "합격 기준", "결과(입력)", "확인자·일자(입력)"],
          [row + ["", ""] for row in PREP_CHECKS], [6, 18, 44, 34, 16, 18], inputs=(5, 6))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 데모 LIMS ───────────────────────────────────────────────────────────
TEST_TAT_HOURS = {"KILN_FEED": 1.0, "CLINKER": 1.5, "CLINKER_DAILY": 6.0, "CEMENT_MILL": 1.5}
AGE_DAYS = {"CS1": 1, "CS3": 3, "CS7": 7, "CS28": 28, "AUTOCLAVE": 1, "CR6": 1}


def create_demo_lims(raw: dict[str, pd.DataFrame], path: Path = LIMS_DEMO_DB, days: int = 10,
                     now: datetime | None = None) -> dict[str, int]:
    """QMS 데이터(최근 days 일)를 LIMS 형식(시료·결과 테이블)으로 만든다. 승인 대기·미등록 시험코드도 일부 포함."""
    rows_s, rows_r = [], []
    sid = 0
    product_code = {"1종": "OPC", "3종": "HES"}
    rng = np.random.default_rng(11)
    ends = [pd.to_datetime(df["timestamp"]).max() for df in raw.values() if len(df)]
    if not ends:
        return {"samples": 0, "results": 0}
    end = max(ends)
    now = now or end.to_pydatetime() + timedelta(hours=1)
    start = end - pd.Timedelta(days=days)
    for table, sp in SAMPLE_POINTS.items():
        df = raw.get(table)
        if df is None or len(df) == 0:
            continue
        df = df[pd.to_datetime(df["timestamp"]) >= start]
        for _, r in df.iterrows():
            sid += 1
            ts = pd.Timestamp(r["timestamp"])
            prod = product_code.get(r.get("product")) if "product" in df.columns else None
            rows_s.append((sid, sp, ts.strftime("%Y-%m-%d %H:%M:%S"), prod))
            for col in RAW_COLUMNS[table]:
                if col == "product" or col in DCS_ONLY or col not in TEST_CODES or col not in r or pd.isna(r[col]):
                    continue
                code = TEST_CODES[col]
                tat = timedelta(days=AGE_DAYS.get(code, 0), hours=TEST_TAT_HOURS.get(sp, 2.0))
                approved = ts + tat
                if approved > pd.Timestamp(now):
                    continue
                status = "PENDING" if rng.random() < 0.02 else "APPROVED"
                val = r[col]
                value = str(val) if col in TEXT_COLUMNS else (f"{float(val):.3f}")
                rows_r.append((sid, code, value, "", status, approved.strftime("%Y-%m-%d %H:%M:%S")))
            if sp == "CEMENT_LOT":       # QMS에 없는 시험(미매핑 예시)
                rows_r.append((sid, "WATER_DEMAND", f"{rng.normal(26.5, 0.4):.1f}", "%", "APPROVED",
                               (ts + timedelta(hours=3)).strftime("%Y-%m-%d %H:%M:%S")))
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()
    with closing(sqlite3.connect(path)) as con:
        con.execute("CREATE TABLE lims_sample (sample_id INTEGER PRIMARY KEY, sample_point TEXT, sampled_at TEXT, product TEXT)")
        con.execute("CREATE TABLE lims_result (sample_id INTEGER, test_code TEXT, value TEXT, unit TEXT, status TEXT, "
                    "approved_at TEXT)")
        con.executemany("INSERT INTO lims_sample VALUES (?,?,?,?)", rows_s)
        con.executemany("INSERT INTO lims_result VALUES (?,?,?,?,?,?)", rows_r)
        con.commit()
    return {"samples": len(rows_s), "results": len(rows_r)}
