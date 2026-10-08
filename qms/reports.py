"""보고서 생성: 엑셀(데이터·통계) · PPT(경영진 보고 / 이상 분석 보고).

- PPT 차트는 PowerPoint 기본(네이티브) 차트로 생성해 보고서 작성자가 직접 수정할 수 있다.
- 글꼴은 '맑은 고딕'(한글 East Asian 글꼴까지 지정)을 사용한다.
"""

from __future__ import annotations

import io
from datetime import datetime

import numpy as np
import pandas as pd

from .alerts import SEVERITY_ORDER
from .diagnosis import VERDICT_LEGEND, Diagnosis
from .spc import capability, capability_grade
from .standards import TABLES, SpecRegistry
from .store import SHEET_NAMES, DataStore

FONT = "맑은 고딕"

# ── 공정능력 표 ─────────────────────────────────────────────────────────
def capability_table(store: DataStore, registry: SpecRegistry, start=None, end=None,
                     key_only: bool = False) -> pd.DataFrame:
    rows = []
    items = registry.key_items() if key_only else registry.all()
    for item in items:
        if item.prediction or item.ks_is_method:
            continue
        for product in registry.products_for(item.key):
            lim = item.limits_for(product)
            if lim.lsl is None and lim.usl is None:
                continue
            s = store.series(item.key, registry, product, start, end)
            if len(s) < 5:
                continue
            cap = capability(s, lim.lsl, lim.usl)
            ks = ("≥ " + f"{lim.ks_min:g}") if lim.ks_min is not None else (("≤ " + f"{lim.ks_max:g}") if lim.ks_max is not None else "-")
            rows.append({"공정": item.stage, "항목": item.name, "품종": product or "-", "단위": item.unit, "n": cap["n"],
                         "평균": round(cap["mean"], item.decimals + 1), "표준편차": round(cap["std"], item.decimals + 2),
                         "사내 하한": lim.lsl, "사내 상한": lim.usl, "KS": ks,
                         "Cpk": round(cap["Cpk"], 2) if np.isfinite(cap["Cpk"]) else np.nan,
                         "Ppk": round(cap["Ppk"], 2) if np.isfinite(cap["Ppk"]) else np.nan,
                         "기준이탈(%)": round(cap["out_pct"], 1) if np.isfinite(cap["out_pct"]) else np.nan,
                         "등급": capability_grade(cap["Cpk"]), "item_key": item.key})
    return pd.DataFrame(rows)


# ── 엑셀 ────────────────────────────────────────────────────────────────
def build_excel_report(store: DataStore, registry: SpecRegistry, events_df: pd.DataFrame,
                       diagnoses: list[Diagnosis], start, end, include_raw: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="37474F")
    head_font = Font(bold=True, color="FFFFFF", name=FONT)
    body_font = Font(name=FONT, size=10)
    thin = Side(style="thin", color="CFD8DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sev_fill = {"위험": "FFCDD2", "경고": "FFE0B2", "주의": "FFF9C4"}
    grade_fill = {"매우 우수": "C8E6C9", "충분": "DCEDC8", "보통(개선 권장)": "FFF9C4", "부족(개선 필요)": "FFE0B2",
                  "매우 부족(즉시 개선)": "FFCDD2"}

    def write_table(ws, df: pd.DataFrame, r0: int = 1, widths: dict | None = None):
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(row=r0, column=j, value=col)
            c.fill, c.font, c.border = head_fill, head_font, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, (_, row) in enumerate(df.iterrows(), r0 + 1):
            for j, col in enumerate(df.columns, 1):
                v = row[col]
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                elif isinstance(v, (float, np.floating)) and not np.isfinite(v):
                    v = None
                elif isinstance(v, (np.integer,)):
                    v = int(v)
                c = ws.cell(row=i, column=j, value=v)
                c.font, c.border = body_font, border
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 30)
                if isinstance(v, datetime):
                    c.number_format = "yyyy-mm-dd hh:mm"
        for j, col in enumerate(df.columns, 1):
            w = (widths or {}).get(col)
            if w is None:
                sample = [len(str(col))] + [len(str(x)) for x in df[col].head(200)]
                w = min(max(sample) + 2, 60)
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)

    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "품질관리 현황 보고 (Blue365 QMS)"
    ws["A1"].font = Font(bold=True, size=16, name=FONT)
    ws["A2"] = f"대상 기간: {pd.Timestamp(start):%Y-%m-%d} ~ {pd.Timestamp(end):%Y-%m-%d}   |   생성: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(size=10, color="5F6B73", name=FONT)
    kpi = kpi_summary(store, registry, events_df, start, end)
    kdf = pd.DataFrame([{"구분": k, "값": v[0], "설명": v[1]} for k, v in kpi.items()])
    write_table(ws, kdf, r0=4, widths={"구분": 28, "값": 16, "설명": 70})
    r = 6 + len(kdf)
    if len(events_df):
        pivot = (events_df.pivot_table(index="stage", columns="severity", values="event_id", aggfunc="count", fill_value=0)
                 .reindex(columns=["위험", "경고", "주의"], fill_value=0).reset_index().rename(columns={"stage": "공정"}))
        ws.cell(row=r, column=1, value="공정별 이상 이벤트 건수").font = Font(bold=True, size=12, name=FONT)
        write_table(ws, pivot, r0=r + 1, widths={"공정": 28, "위험": 16, "경고": 70, "주의": 12})
    ws.freeze_panes = None

    cap = capability_table(store, registry, start, end).drop(columns=["item_key"], errors="ignore")
    ws2 = wb.create_sheet("공정능력")
    write_table(ws2, cap)
    if len(cap):
        gcol = list(cap.columns).index("등급") + 1
        for i in range(2, len(cap) + 2):
            g = ws2.cell(row=i, column=gcol).value
            if g in grade_fill:
                ws2.cell(row=i, column=gcol).fill = PatternFill("solid", fgColor=grade_fill[g])

    ws3 = wb.create_sheet("알림목록")
    cols = {"severity": "심각도", "stage": "공정", "product": "품종", "item_name": "항목", "rule_desc": "판정",
            "start": "시작", "end": "종료", "n_points": "점수", "worst_value": "최악값", "limit_value": "기준값",
            "unit": "단위", "patterns": "동반 패턴", "status": "처리상태", "assignee": "담당", "note": "메모",
            "message": "내용"}
    edf = events_df[[c for c in cols if c in events_df]].rename(columns=cols) if len(events_df) else pd.DataFrame(columns=list(cols.values()))
    write_table(ws3, edf, widths={"내용": 80, "메모": 30})
    if len(edf):
        for i in range(2, len(edf) + 2):
            v = ws3.cell(row=i, column=1).value
            if v in sev_fill:
                ws3.cell(row=i, column=1).fill = PatternFill("solid", fgColor=sev_fill[v])

    ws4 = wb.create_sheet("원인진단")
    drows = []
    for dx in diagnoses:
        top = dx.results[:3]
        drows.append({
            "심각도": dx.event.severity, "품종": dx.event.product or "-", "현상": dx.phenomenon.title,
            "기간": f"{dx.event.start:%m/%d %H:%M} ~ {dx.event.end:%m/%d %H:%M}",
            "현상 정량화(①)": dx.quant["text"], "종합 판단": dx.summary,
            "원인 1순위(②)": f"[{top[0].hyp.axis}] {top[0].hyp.title} — {top[0].verdict}" if top else "",
            "원인 2순위": f"[{top[1].hyp.axis}] {top[1].hyp.title} — {top[1].verdict}" if len(top) > 1 else "",
            "원인 3순위": f"[{top[2].hyp.axis}] {top[2].hyp.title} — {top[2].verdict}" if len(top) > 2 else "",
            "확인 방법(③)": " / ".join(top[0].hyp.verify) if top else "",
            "단기 조치(④)": "\n".join(a for a, _ in dx.short_term[:4]),
            "근본 대책(④)": "\n".join(a for a, _ in dx.root_cause[:4]),
            "KS 평가(⑤)": "\n".join(f"{r['항목']}({r['품종']}): {r['KS 판정']}" for r in dx.ks_eval),
        })
    write_table(ws4, pd.DataFrame(drows) if drows else pd.DataFrame(columns=["심각도"]),
                widths={"현상 정량화(①)": 60, "종합 판단": 70, "단기 조치(④)": 45, "근본 대책(④)": 45, "KS 평가(⑤)": 45,
                        "확인 방법(③)": 50, "원인 1순위(②)": 40, "원인 2순위": 40, "원인 3순위": 40})
    ws4.cell(row=len(drows) + 3, column=1, value="※ " + VERDICT_LEGEND).font = Font(size=9, color="5F6B73", name=FONT)

    if include_raw:
        f = store.filtered(start, end)
        for name, df in f.tables.items():
            if len(df) == 0:
                continue
            wsr = wb.create_sheet(f"데이터_{SHEET_NAMES.get(name, name)}")
            out = df.copy()
            for c in out.select_dtypes("number").columns:
                out[c] = out[c].round(4)
            write_table(wsr, out)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def kpi_summary(store: DataStore, registry: SpecRegistry, events_df: pd.DataFrame, start, end) -> dict:
    """요약 KPI: {이름: (값, 설명)}."""
    ev = events_df if len(events_df) else pd.DataFrame(columns=["severity", "rule"])
    out = {
        "KS 규격 이탈(위험)": (int((ev["severity"] == "위험").sum()), "제품 KS L 5201 규격 이탈 이벤트 — 출하 판정 필요"),
        "사내 관리기준 이탈(경고)": (int((ev["severity"] == "경고").sum()), "사내기준·시험조건 이탈 이벤트"),
        "SPC 이상 패턴(주의)": (int((ev["severity"] == "주의").sum()), "규격 이내이나 공정 불안정 징후"),
    }
    for product in registry.products_for("phy_s28"):
        s = store.series("phy_s28", registry, product, start, end)
        lim = registry["phy_s28"].limits_for(product)
        if len(s):
            out[f"28일 강도 평균({product})"] = (f"{s.mean():.1f} MPa",
                                            f"최소 {s.min():.1f} / KS {lim.ks_min:g} 대비 여유 {s.min() - lim.ks_min:+.1f} MPa, "
                                            f"사내 {lim.lsl:g} 미달 {int((s < lim.lsl).sum())}건 (n={len(s)})")
        p = store.series("pred_s28", registry, product, start, end)
        if len(p):
            out[f"28일 강도 예측({product}, 미도래)"] = (f"{p.mean():.1f} MPa",
                                                 f"최소 {p.min():.1f}, 사내 {lim.lsl:g} 미달 예측 {int((p < lim.lsl).sum())}로트")
    cap = capability_table(store, registry, start, end, key_only=True)
    if len(cap):
        weak = cap[cap["Cpk"] < 1.0]
        out["공정능력 부족 항목(Cpk<1.0)"] = (len(weak), ", ".join(f"{r['항목']}({r['품종']})" for _, r in weak.iterrows())[:200] or "없음")
    return out


# ── PPT 공통 ────────────────────────────────────────────────────────────
class _Deck:
    INK = "1F2A30"
    MUTED = "5F6B73"
    CARD = "F1F3F4"
    DARK = "263238"
    ACCENT = "C2410C"   # 킬른 화염색 — 핵심 메시지 강조(본문 대비 5:1 이상)
    # 상태색(의미가 있을 때만 사용, 항상 라벨과 함께) — dataviz 기준 팔레트
    GOOD = "0CA30C"
    WARN = "EC835A"      # serious
    DANGER = "D03B3B"    # critical
    CAUTION = "FAB219"   # warning
    SERIES = ["2A78D6", "6DA7EC"]  # 실측(진한 파랑) / 예측(같은 색상 계열의 밝은 단계, 점선)
    SEV = {"위험": "D03B3B", "경고": "EC835A", "주의": "FAB219", "정상": "0CA30C"}

    def __init__(self, title: str):
        from pptx import Presentation
        from pptx.util import Inches
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.prs.core_properties.title = title
        self.prs.core_properties.author = "Blue365 QMS"
        self._blank = self.prs.slide_layouts[6]

    # 기본 요소 -----------------------------------------------------------
    @staticmethod
    def _rgb(hexstr: str):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(hexstr)

    def _font(self, font, size: float, bold: bool = False, color: str | None = None, italic: bool = False):
        from pptx.oxml.ns import qn
        from pptx.util import Pt
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.italic = italic
        if color:
            font.color.rgb = self._rgb(color)
        rpr = font._rPr
        latin = rpr.find(qn("a:latin"))
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                (latin if latin is not None else rpr).addnext(el) if latin is not None else rpr.append(el)
                latin = el
            el.set("typeface", FONT)

    def slide(self, dark: bool = False):
        s = self.prs.slides.add_slide(self._blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = self._rgb(self.DARK if dark else "FFFFFF")
        return s

    def text(self, slide, x, y, w, h, paragraphs, size=14, color=None, bold=False, align="left",
             anchor="top", name: str | None = None, line_spacing: float | None = None, space_after: float = 4):
        """paragraphs: str | list[str | (str, dict)] — dict: size/bold/color/bullet/italic."""
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]
        for i, p in enumerate(paragraphs):
            opts = {}
            if isinstance(p, tuple):
                p, opts = p
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[opts.get("align", align)]
            para.space_after = Pt(opts.get("space_after", space_after))
            if line_spacing:
                para.line_spacing = line_spacing
            if opts.get("bullet"):
                self._bullet(para)
            run = para.add_run()
            run.text = str(p)
            self._font(run.font, opts.get("size", size), opts.get("bold", bold), opts.get("color", color or self.INK),
                       opts.get("italic", False))
        return tb

    @staticmethod
    def _bullet(para):
        from pptx.oxml.ns import qn
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", "228600")
        pPr.set("indent", "-228600")
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            el = pPr.find(qn(tag))
            if el is not None:
                pPr.remove(el)
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)

    def box(self, slide, x, y, w, h, fill: str, rounded: bool = True, name: str | None = None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                     Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self._rgb(fill)
        shp.line.fill.background()
        shp.shadow.inherit = False
        if rounded:
            shp.adjustments[0] = 0.06
        if name:
            shp.name = name
        return shp

    def title(self, slide, title: str, message: str | None = None):
        self.text(slide, 0.6, 0.38, 12.1, 0.75, title, size=30, bold=True, name="Title")
        if message:
            self.text(slide, 0.6, 1.08, 12.1, 0.5, message, size=15, color=self.ACCENT, name="Key message")

    def footer(self, slide, note: str):
        self.text(slide, 0.6, 7.02, 12.1, 0.32, note, size=10, color=self.MUTED, name="Source note")

    def table(self, slide, df: pd.DataFrame, x, y, w, col_widths: list[float] | None = None, font_size=11,
              row_h=0.36, header_fill: str | None = None, cell_fills: dict | None = None, name: str | None = None):
        """cell_fills: {(row_idx, col_name): (fill_hex, font_hex)} — row_idx 는 df 기준 0부터."""
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches
        rows, cols = len(df) + 1, len(df.columns)
        shp = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(row_h * rows))
        if name:
            shp.name = name
        tbl = shp.table
        widths = col_widths or [w / cols] * cols
        for j, cw in enumerate(widths):
            tbl.columns[j].width = Inches(cw)
        for i in range(rows):
            tbl.rows[i].height = Inches(row_h)
        hf = header_fill or self.DARK
        for j, col in enumerate(df.columns):
            self._cell(tbl.cell(0, j), str(col), font_size, True, "FFFFFF", hf)
        for i, (_, r) in enumerate(df.iterrows(), 1):
            for j, col in enumerate(df.columns):
                fill, fc = (cell_fills or {}).get((i - 1, col), ("FFFFFF" if i % 2 else "F7F8F9", self.INK))
                v = r[col]
                txt = "-" if v is None or (isinstance(v, float) and not np.isfinite(v)) else str(v)
                self._cell(tbl.cell(i, j), txt, font_size, False, fc, fill)
        for i in range(rows):
            for j in range(cols):
                tbl.cell(i, j).vertical_anchor = MSO_ANCHOR.MIDDLE
        return shp

    def _cell(self, cell, text, size, bold, color, fill):
        from pptx.util import Inches
        cell.fill.solid()
        cell.fill.fore_color.rgb = self._rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        tf = cell.text_frame
        tf.word_wrap = True
        para = tf.paragraphs[0]
        run = para.add_run()
        run.text = text
        self._font(run.font, size, bold, color)

    def line_chart(self, slide, categories: list[str], series: list[dict], x, y, w, h, title: str = "",
                   y_title: str = "", label_skip: int | None = None, name: str | None = None,
                   y_min: float | None = None, y_max: float | None = None):
        """series: [{name, values, color, dash(bool), width, marker(bool)}]"""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt

        cd = CategoryChartData()
        cd.categories = categories
        for s in series:
            cd.add_series(s["name"], [None if (v is None or not np.isfinite(v)) else float(v) for v in s["values"]])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        if name:
            gf.name = name
        ch = gf.chart
        self._font(ch.font, 10, color=self.MUTED)
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            self._font(ch.legend.font, 10, color=self.MUTED)
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            self._font(ch.chart_title.text_frame.paragraphs[0].runs[0].font, 13, True, self.INK)
        else:
            ch.has_title = False
        for s, ser in zip(series, ch.series):
            ser.smooth = False
            ln = ser.format.line
            ln.color.rgb = self._rgb(s.get("color", self.SERIES[0]))
            ln.width = Pt(s.get("width", 2.0))
            if s.get("dash"):
                ln.dash_style = MSO_LINE_DASH_STYLE.DASH
            if s.get("marker", False):
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 5
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = self._rgb(s.get("color", self.SERIES[0]))
                ser.marker.format.line.fill.background()
            else:
                ser.marker.style = XL_MARKER_STYLE.NONE
        va = ch.value_axis
        if y_min is None or y_max is None:
            vals = np.array([v for ser_ in series for v in ser_["values"] if v is not None and np.isfinite(v)], float)
            if len(vals):
                lo, hi = float(vals.min()), float(vals.max())
                pad = (hi - lo) * 0.12 or abs(hi) * 0.05 or 1.0
                step = _nice_step(hi - lo + 2 * pad)
                if y_min is None:
                    y_min = np.floor((lo - pad) / step) * step
                    if lo >= 0 > y_min:
                        y_min = 0.0
                if y_max is None:
                    y_max = np.ceil((hi + pad) / step) * step
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = self._rgb("E3E7EA")
        va.format.line.fill.background()
        if y_min is not None:
            va.minimum_scale = y_min
        if y_max is not None:
            va.maximum_scale = y_max
        self._font(va.tick_labels.font, 10, color=self.MUTED)
        if y_title:
            va.has_title = True
            va.axis_title.text_frame.text = y_title
            self._font(va.axis_title.text_frame.paragraphs[0].runs[0].font, 10, color=self.MUTED)
        ca = ch.category_axis
        ca.format.line.color.rgb = self._rgb("B0BEC5")
        self._font(ca.tick_labels.font, 9, color=self.MUTED)
        skip = label_skip or max(1, len(categories) // 10)
        if skip > 1:
            cat_ax = ca._element
            nml = cat_ax.find(qn("c:noMultiLvlLbl"))
            el = cat_ax.makeelement(qn("c:tickLblSkip"), {"val": str(skip)})
            el2 = cat_ax.makeelement(qn("c:tickMarkSkip"), {"val": str(skip)})
            if nml is not None:
                nml.addprevious(el)
                nml.addprevious(el2)
            else:
                cat_ax.append(el)
                cat_ax.append(el2)
        return gf

    def kpi_card(self, slide, x, y, w, h, value: str, label: str, sub: str = "", status: str | None = None,
                 value_size: float = 34):
        """KPI 카드. 값은 본문 잉크색, 상태는 색 점 + 라벨 칩으로 표시(색만으로 의미 전달 금지)."""
        self.box(slide, x, y, w, h, self.CARD, name=f"KPI {label}")
        self.text(slide, x + 0.2, y + 0.15, w - 1.4 if status else w - 0.4, 0.4, label, size=13, color=self.MUTED)
        if status:
            self.status_chip(slide, x + w - 1.25, y + 0.17, status)
        self.text(slide, x + 0.2, y + 0.55, w - 0.4, 0.8, value, size=value_size, bold=True, color=self.INK)
        if sub:
            self.text(slide, x + 0.2, y + h - 0.6, w - 0.4, 0.5, sub, size=11, color=self.MUTED)

    def status_chip(self, slide, x, y, status: str):
        """상태 칩: ● 색 점 + 라벨. status ∈ 위험/경고/주의/정상 또는 '양호'."""
        from pptx.util import Pt
        key = {"양호": "정상"}.get(status, status)
        col = self.SEV.get(key, self.MUTED)
        self.box(slide, x, y, 1.05, 0.34, "FFFFFF", name=f"Status {status}")
        tb = self.text(slide, x, y, 1.05, 0.34, "", size=11, align="center", anchor="middle")
        para = tb.text_frame.paragraphs[0]
        para.runs[0].text = "● "
        self._font(para.runs[0].font, 11, True, col)
        r2 = para.add_run()
        r2.text = status
        self._font(r2.font, 11, True, self.INK)
        para.space_after = Pt(0)

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def _nice_step(span: float) -> float:
    """눈금 간격(1·2·5 계열)."""
    if span <= 0 or not np.isfinite(span):
        return 1.0
    raw = span / 6
    mag = 10 ** np.floor(np.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * mag:
            return float(m * mag)
    return float(10 * mag)


def _impact(dx: Diagnosis) -> tuple:
    """보고서 상세 분석 대상 우선순위: 심각도 → 제품 단계 → 지속 점수."""
    e = dx.event
    stage_w = 2 if e.stage in ("제품 물성", "시멘트(분쇄)") else 1
    return (SEVERITY_ORDER[e.severity], stage_w * e.n_points)


def _prod(v) -> str:
    return f"[{v}] " if isinstance(v, str) and v else ""


def _daily(series: pd.Series) -> pd.Series:
    return series.resample("1D").mean()


def _chart_series_for_item(store: DataStore, registry: SpecRegistry, key: str, product, start, end):
    """차트용 시계열: 고빈도 데이터는 8시간 평균, 일 데이터는 그대로."""
    item = registry[key]
    s = store.series(key, registry, product, start, end)
    rule = TABLES[item.table].get("spc_resample")
    if rule and len(s):
        span_days = (s.index.max() - s.index.min()).days if len(s) > 1 else 0
        s = s.resample("1D" if span_days > 14 else rule).mean().dropna()
    return s


def _fmt_cat(ts: pd.Timestamp, daily: bool) -> str:
    return f"{ts:%m/%d}" if daily else f"{ts:%m/%d %H}시"


# ── 경영진 보고 PPT ─────────────────────────────────────────────────────
def build_management_ppt(store: DataStore, registry: SpecRegistry, events_df: pd.DataFrame,
                         diagnoses: list[Diagnosis], start, end, org: str = "품질관리팀") -> bytes:
    start, end = pd.Timestamp(start), pd.Timestamp(end)
    d = _Deck("품질관리 현황 보고")
    period = f"{start:%Y.%m.%d} ~ {end:%Y.%m.%d}"
    src = f"출처: Blue365 QMS 자동 집계({datetime.now():%Y-%m-%d %H:%M}) · KS L 5201 포틀랜드 시멘트 · 사내 관리기준"

    # 1. 표지
    s = d.slide(dark=True)
    d.text(s, 0.8, 2.3, 11.5, 1.0, "품질관리 현황 보고", size=44, bold=True, color="FFFFFF")
    d.text(s, 0.8, 3.35, 11.5, 0.6, f"대상 기간 {period}", size=20, color="CFD8DC")
    d.text(s, 0.8, 5.9, 11.5, 0.5, f"{org} · Blue365 통합 품질관리 시스템 자동 생성", size=14, color="B0BEC5")

    # 2. 핵심 요약
    ev = events_df if len(events_df) else pd.DataFrame(columns=["severity", "item_name", "product", "message"])
    n_r, n_w, n_c = (int((ev["severity"] == k).sum()) for k in ("위험", "경고", "주의"))
    s = d.slide()
    main = ("KS 규격 이탈이 발생해 출하 판정이 필요합니다" if n_r else
            "KS 규격 이탈은 없으나 사내 관리기준 이탈 대응이 필요합니다" if n_w else "전 항목 관리 상태가 안정적입니다")
    d.title(s, "핵심 요약", main)
    s28 = store.series("phy_s28", registry, "1종", start, end)
    lim28 = registry["phy_s28"].limits_for("1종")
    cap = capability_table(store, registry, start, end, key_only=True)
    weak = cap[cap["Cpk"] < 1.0] if len(cap) else cap
    cw, gap, x0 = 2.85, 0.27, 0.6
    d.kpi_card(s, x0, 1.75, cw, 2.0, f"{n_r}건", "KS 규격 이탈", "출하 판정·고객 영향 검토 대상",
               "위험" if n_r else "양호")
    d.kpi_card(s, x0 + (cw + gap), 1.75, cw, 2.0, f"{n_w}건", "사내기준 이탈", "원인 조사·조치 대상",
               "경고" if n_w else "양호")
    if len(s28):
        ok28 = s28.min() >= lim28.lsl
        d.kpi_card(s, x0 + 2 * (cw + gap), 1.75, cw, 2.0, f"{s28.mean():.1f} MPa", "28일 강도",
                   f"1종 · 최소 {s28.min():.1f} · KS {lim28.ks_min:g} 대비 +{s28.min() - lim28.ks_min:.1f}",
                   "양호" if ok28 else "경고", value_size=30)
    else:
        d.kpi_card(s, x0 + 2 * (cw + gap), 1.75, cw, 2.0, "-", "28일 강도(1종)", "결과 미도래")
    d.kpi_card(s, x0 + 3 * (cw + gap), 1.75, cw, 2.0, f"{len(weak)}개", "Cpk 1.0 미만",
               ", ".join(weak["항목"].head(2)) if len(weak) else "핵심 항목 모두 1.0 이상", "주의" if len(weak) else "양호")
    d.text(s, 0.6, 4.05, 12.1, 0.45, "주요 사항", size=18, bold=True)
    bullets = []
    top_ev = ev.copy()
    if len(top_ev):
        top_ev["_o"] = top_ev["severity"].map(SEVERITY_ORDER)
        top_ev = top_ev[top_ev["severity"] != "주의"].sort_values(["_o", "end"], ascending=[False, False]).head(3)
    for _, r in top_ev.iterrows():
        dx = next((x for x in diagnoses if x.event.event_id == r["event_id"]), None)
        cause = ""
        if dx and dx.results:
            top = next((x for x in dx.results if x.supported), dx.results[0])
            cause = f" → {top.hyp.title}({top.verdict})"
        when = f"{r['start']:%m/%d}" + ("" if r["start"].date() == r["end"].date() else f"~{r['end']:%m/%d}")
        bullets.append((f"[{r['severity']}] {_prod(r.get('product'))}{r['item_name']} {r['rule_desc']} ({when}){cause}",
                        {"bullet": True, "size": 14}))
    if n_c:
        bullets.append((f"SPC 이상 패턴(주의) {n_c}건 — 규격 이내이나 추세·평균 이동 징후, 공정별 점검 진행", {"bullet": True, "size": 14}))
    if not bullets:
        bullets = [("대상 기간 중 관리기준 이탈 없음", {"bullet": True})]
    d.text(s, 0.6, 4.5, 12.1, 2.4, bullets, size=14)
    d.footer(s, src)

    # 3. 28일 강도 추이
    s = d.slide()
    phy = store.tables.get("physical", pd.DataFrame())
    p1 = phy[(phy["product"] == "1종") & (phy["timestamp"] >= start) & (phy["timestamp"] <= end)] if len(phy) else phy
    if len(p1):
        cats = [_fmt_cat(t, True) for t in p1["timestamp"]]
        act = p1["phy_s28"].to_numpy(float)
        pred = p1["pred_s28"].to_numpy(float) if "pred_s28" in p1 else np.full(len(p1), np.nan)
        n_pred_low = int(np.nansum(pred < lim28.lsl))
        msg = (f"28일 강도 평균 {np.nanmean(act):.1f} MPa, KS {lim28.ks_min:g} MPa 대비 여유 확보"
               if np.isfinite(np.nanmean(act)) else "28일 강도 결과 미도래")
        if n_pred_low:
            msg = f"미도래 로트 중 {n_pred_low}로트가 사내기준({lim28.lsl:g} MPa) 미달로 예측됩니다"
        d.title(s, "28일 압축강도 추이 (1종)", msg)
        d.line_chart(s, cats, [
            {"name": "28일 실측", "values": act, "color": d.SERIES[0], "marker": True},
            {"name": "28일 예측(미도래)", "values": pred, "color": d.SERIES[1], "marker": True, "dash": True},
            {"name": f"사내 하한 {lim28.lsl:g}", "values": [lim28.lsl] * len(cats), "color": d.WARN, "width": 1.25, "dash": True},
            {"name": f"KS 하한 {lim28.ks_min:g}", "values": [lim28.ks_min] * len(cats), "color": d.DANGER, "width": 1.25},
        ], 0.6, 1.7, 8.6, 5.1, y_title="MPa", name="Strength chart",
            y_min=float(np.floor(min(lim28.ks_min, np.nanmin(np.r_[act, pred])) - 2)))
        d.box(s, 9.5, 1.75, 3.25, 5.0, d.CARD, name="Strength notes")
        notes = [("판독 포인트", {"bold": True, "size": 15})]
        if np.isfinite(np.nanmean(act)):
            notes.append((f"실측 평균 {np.nanmean(act):.1f} / 최소 {np.nanmin(act):.1f} MPa", {"bullet": True}))
            notes.append((f"사내 하한 미달 {int(np.nansum(act < lim28.lsl))}로트", {"bullet": True}))
        if np.isfinite(np.nanmean(pred)):
            notes.append((f"예측 대상 {int(np.isfinite(pred).sum())}로트, 미달 예측 {n_pred_low}로트", {"bullet": True}))
        notes.append(("예측은 조기강도·분말도·C₃S 회귀식(추정치)이며 실측으로 확정", {"bullet": True, "size": 12,
                                                                     "color": d.MUTED}))
        d.text(s, 9.7, 1.95, 2.9, 4.6, notes, size=13)
    else:
        d.title(s, "28일 압축강도 추이 (1종)", "대상 기간 데이터 없음")
    d.footer(s, src + " · 예측값: 다중회귀(LOO RMSE 기반 ±1.96σ)")

    # 4. 공정 핵심 지표
    s = d.slide()
    d.title(s, "공정 핵심 지표 추이", "소성(f-CaO)과 분쇄(분말도)가 제품 강도의 선행 지표입니다")
    for idx, (key, product) in enumerate((("clk_fcao", None), ("cem_blaine", "1종"))):
        item = registry[key]
        lim = item.limits_for(product)
        ser = _chart_series_for_item(store, registry, key, product, start, end)
        if len(ser) == 0:
            continue
        daily = (ser.index.to_series().diff().median() or pd.Timedelta("1D")) >= pd.Timedelta("1D")
        cats = [_fmt_cat(t, daily) for t in ser.index]
        lines = [{"name": item.name + ("(일평균)" if daily else "(8h 평균)"), "values": ser.to_numpy(float),
                  "color": d.SERIES[0], "width": 1.75}]
        if lim.usl is not None:
            lines.append({"name": f"사내 상한 {lim.usl:g}", "values": [lim.usl] * len(cats), "color": d.WARN, "dash": True, "width": 1.25})
        if lim.lsl is not None:
            lines.append({"name": f"사내 하한 {lim.lsl:g}", "values": [lim.lsl] * len(cats), "color": d.WARN, "dash": True, "width": 1.25})
        if lim.ks_min is not None:
            lines.append({"name": f"KS 하한 {lim.ks_min:g}", "values": [lim.ks_min] * len(cats), "color": d.DANGER, "width": 1.25})
        x = 0.6 + idx * 6.2
        title = f"{item.name}" + (f" [{product}]" if product else "") + f" ({item.unit})"
        d.line_chart(s, cats, lines, x, 1.75, 5.9, 4.4, title=title, name=f"Chart {key}")
        n_out = int(((ser > lim.usl).sum() if lim.usl is not None else 0) + ((ser < lim.lsl).sum() if lim.lsl is not None else 0))
        d.text(s, x, 6.25, 5.9, 0.6, f"평균 {ser.mean():,.{item.decimals}f} {item.unit} · 집계값 기준 사내기준 이탈 {n_out}회",
               size=12, color=d.MUTED)
    d.footer(s, src)

    # 5. 주요 이상 현황 표
    s = d.slide()
    d.title(s, "주요 이상 발생 및 조치 현황", f"위험 {n_r}건 · 경고 {n_w}건 · 주의 {n_c}건 (기간 내)")
    if len(ev):
        t = ev.copy()
        t["_o"] = t["severity"].map(SEVERITY_ORDER)
        t = t.sort_values(["_o", "end"], ascending=[False, False]).head(9)
        rows, fills = [], {}
        for i, (_, r) in enumerate(t.iterrows()):
            dx = next((x for x in diagnoses if x.event.event_id == r["event_id"]), None)
            cause = "-"
            if dx and dx.results:
                top = next((x for x in dx.results if x.supported), dx.results[0])
                cause = f"[{top.hyp.axis}] {top.hyp.title}"
            rows.append({"기간": f"{r['start']:%m/%d}~{r['end']:%m/%d}", "공정": r["stage"],
                         "항목": _prod(r.get("product")) + r["item_name"],
                         "판정": r["rule_desc"], "심각도": r["severity"], "1순위 원인(자동 진단)": cause,
                         "상태": r.get("status", "신규")})
            fills[(i, "심각도")] = ({"위험": "FFCDD2", "경고": "FFE0B2", "주의": "FFF9C4"}.get(r["severity"], "FFFFFF"), d.INK)
        d.table(s, pd.DataFrame(rows), 0.6, 1.75, 12.1, [1.25, 1.35, 2.6, 2.3, 0.85, 2.95, 0.8], font_size=11,
                row_h=0.46, cell_fills=fills, name="Event table")
    else:
        d.text(s, 0.6, 2.0, 12, 1, "기간 내 이상 이벤트가 없습니다.", size=16)
    d.footer(s, src + " · 원인은 데이터 기반 자동 진단(1차 판단)이며 현장 확인으로 확정")

    # 6. 상세 분석 (위험/경고 상위 2건)
    ranked = sorted([x for x in diagnoses if x.event.severity in ("위험", "경고")], key=_impact, reverse=True)
    for dx in ranked[:2]:
        _issue_overview_slide(d, dx, store, registry, src)

    # 7. 공정능력
    s = d.slide()
    d.title(s, "핵심 항목 공정능력", "Cpk 1.33 이상 충분 · 1.0 미만은 개선 대상 (관용 기준)")
    if len(cap):
        c = cap.copy().sort_values("Cpk").head(12)
        rows, fills = [], {}
        gcol = {"매우 우수": "C8E6C9", "충분": "DCEDC8", "보통(개선 권장)": "FFF9C4", "부족(개선 필요)": "FFE0B2",
                "매우 부족(즉시 개선)": "FFCDD2"}
        for i, (_, r) in enumerate(c.iterrows()):
            lim_txt = " ~ ".join(f"{v:g}" for v in (r["사내 하한"], r["사내 상한"]) if pd.notna(v))
            lim_txt = (("≥ " if pd.notna(r["사내 하한"]) and pd.isna(r["사내 상한"]) else
                        "≤ " if pd.isna(r["사내 하한"]) and pd.notna(r["사내 상한"]) else "") + lim_txt)
            rows.append({"공정": r["공정"], "항목": r["항목"], "품종": r["품종"], "평균": f"{r['평균']:g} {r['단위']}",
                         "사내 기준": lim_txt, "KS": r["KS"], "Cpk": f"{r['Cpk']:.2f}" if pd.notna(r["Cpk"]) else "-",
                         "등급": r["등급"]})
            fills[(i, "등급")] = (gcol.get(r["등급"], "FFFFFF"), d.INK)
        d.table(s, pd.DataFrame(rows), 0.6, 1.75, 12.1, [1.5, 2.5, 0.8, 1.6, 1.7, 1.1, 0.9, 2.0], font_size=11,
                row_h=0.4, cell_fills=fills, name="Capability table")
    d.footer(s, src + " · Cpk: 군내변동(MR/1.128) 기준")

    # 8. 후속 조치
    s = d.slide()
    d.title(s, "후속 조치 계획", "자동 진단에서 도출된 근본 대책을 담당 부서별로 확정해 주십시오")
    acts, seen = [], set()
    for dx in ranked[:5]:
        for a, tag in dx.root_cause[:3]:
            if a in seen:
                continue
            seen.add(a)
            owner = next((r.hyp.owner for r in dx.results if f"[{r.hyp.axis}] {r.hyp.title}" == tag), "품질")
            acts.append({"근본 대책": a, "관련 원인": tag, "담당(제안)": owner, "기한(제안)": "1개월"})
    if acts:
        d.table(s, pd.DataFrame(acts[:9]), 0.6, 1.75, 12.1, [5.0, 4.2, 1.6, 1.3], font_size=11, row_h=0.46,
                name="Action table")
    else:
        d.text(s, 0.6, 2.0, 12, 1, "추가 조치 사항 없음", size=16)
    d.footer(s, "담당·기한은 시스템 제안값 — 회의에서 확정 필요")
    return d.save()


def _top_cause_line(dx: Diagnosis) -> str:
    if not dx.results:
        return "원인 가설 없음"
    top = next((r for r in dx.results if r.supported), None)
    if top is None:
        r = dx.results[0]
        return f"데이터로 지지되는 원인 없음 — 우선 확인: [{r.hyp.axis}] {r.hyp.title} (추정)"
    score = f", 근거 점수 {top.evidence_score}" if top.evidence_score is not None else ""
    return f"1순위 원인: [{top.hyp.axis}] {top.hyp.title} ({top.verdict}{score})"


def _when(e) -> str:
    if e.start == e.end:
        return f"{e.start:%m/%d %H:%M}"
    return f"{e.start:%m/%d %H:%M} ~ {e.end:%m/%d %H:%M}"


def _clip(text: str, n: int) -> str:
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _issue_overview_slide(d: _Deck, dx: Diagnosis, store: DataStore, registry: SpecRegistry, src: str):
    e, q = dx.event, dx.quant
    s = d.slide()
    d.title(s, f"이상 분석: {_prod(e.product)}{dx.phenomenon.title}", _top_cause_line(dx))
    d.box(s, 0.6, 1.75, 5.6, 3.0, d.CARD, name="Phenomenon card")
    d.text(s, 0.8, 1.88, 5.2, 0.4, "① 현상 정량화", size=15, bold=True)
    nd = q["decimals"]
    lines = [(f"{_when(e)} · {e.severity} · {e.rule_desc}", {"size": 12}),
             (f"구간 평균 {q['mean']:,.{nd}f} {q['unit']} (기준기간 {q['base_mean']:,.{nd}f}, Δ {q['delta']:+,.{nd}f})",
              {"size": 12})]
    if e.limit_value is not None:
        lines.append((f"{'최대' if e.direction == 'high' else '최소'} {e.worst_value:,.{nd}f} {q['unit']} · "
                      f"{q.get('limit_word', '기준')} {e.limit_value:,.{nd}f}", {"size": 12}))
    lines.append((_clip("종합 판단: " + dx.summary, 210), {"size": 11, "color": d.MUTED}))
    d.text(s, 0.8, 2.3, 5.2, 2.4, lines, size=12)
    d.box(s, 0.6, 4.95, 5.6, 1.9, d.CARD, name="KS card")
    d.text(s, 0.8, 5.07, 5.2, 0.4, "⑤ KS·사내기준 평가", size=15, bold=True)
    ks_lines = [(_clip(f"{r['항목']}({r['품종']}): {r['KS 판정']}", 60), {"size": 12, "bullet": True}) for r in dx.ks_eval[:3]]
    d.text(s, 0.8, 5.47, 5.2, 1.35, ks_lines or [("평가 대상 없음", {})], size=12)
    rows, fills = [], {}
    for i, r in enumerate(dx.results[:5]):
        rows.append({"순위": i + 1, "축": r.hyp.axis, "원인 가설(②)": r.hyp.title, "판정": r.verdict,
                     "근거": r.evidence_score if r.evidence_score is not None else "-"})
        fills[(i, "판정")] = _VERDICT_FILL.get(r.verdict, ("FFFFFF", d.INK))
    d.table(s, pd.DataFrame(rows), 6.5, 1.75, 6.2, [0.5, 0.8, 2.7, 1.6, 0.6], font_size=11, row_h=0.42,
            cell_fills=fills, name="Hypothesis table")
    st = [(a, {"bullet": True, "size": 12}) for a, _ in dx.short_term[:3]]
    rc = [(a, {"bullet": True, "size": 12}) for a, _ in dx.root_cause[:3]]
    d.text(s, 6.5, 4.45, 3.0, 0.4, "④ 단기 조치", size=15, bold=True)
    d.text(s, 6.5, 4.85, 3.0, 2.0, st or [("-", {})], size=12)
    d.text(s, 9.7, 4.45, 3.0, 0.4, "④ 근본 대책", size=15, bold=True)
    d.text(s, 9.7, 4.85, 3.0, 2.0, rc or [("-", {})], size=12)
    d.footer(s, VERDICT_LEGEND_SHORT)


_VERDICT_FILL = {"데이터 지지": ("C8E6C9", "1F2A30"), "부분 확인": ("DCEDC8", "1F2A30"),
                 "데이터 없음(추정)": ("ECEFF1", "1F2A30"), "근거 약함": ("FFFFFF", "5F6B73"),
                 "반증(가능성 낮음)": ("FFFFFF", "5F6B73")}
VERDICT_LEGEND_SHORT = ("판정: 데이터 지지 = 예상 변화가 측정 데이터로 확인됨(사실, 인과는 현장 확인으로 확정) · "
                        "데이터 없음 = 경험칙 추정 · 반증 = 데이터에 예상 변화 없음")


# ── 단건 이상 분석 PPT ─────────────────────────────────────────────────
def build_issue_ppt(dx: Diagnosis, store: DataStore, registry: SpecRegistry, org: str = "품질관리팀") -> bytes:
    e, q = dx.event, dx.quant
    item = registry[e.item_key]
    d = _Deck("품질 이상 분석 보고")
    src = f"출처: Blue365 QMS 자동 진단({datetime.now():%Y-%m-%d %H:%M}) · {item.ks_ref or item.basis}"
    prod = f"[{e.product}] " if e.product else ""

    s = d.slide(dark=True)
    d.text(s, 0.8, 2.1, 11.5, 0.9, "품질 이상 분석 보고", size=40, bold=True, color="FFFFFF")
    d.text(s, 0.8, 3.1, 11.5, 0.7, f"{prod}{dx.phenomenon.title}", size=24, color="FFCCBC")
    d.text(s, 0.8, 3.85, 11.5, 0.6, f"{e.start:%Y-%m-%d %H:%M} ~ {e.end:%Y-%m-%d %H:%M} · 심각도 {e.severity}", size=16,
           color="CFD8DC")
    d.text(s, 0.8, 5.9, 11.5, 0.5, f"{org} · Blue365 QMS 자동 생성 — 현장 확인 후 확정 필요", size=14, color="B0BEC5")

    # ① 현상
    s = d.slide()
    d.title(s, "① 현상 정의와 정량화", f"{e.rule_desc} — {item.name}")
    pad = pd.Timedelta(days=7)
    ser = _chart_series_for_item(store, registry, e.item_key, e.product, e.start - pad, e.end + pd.Timedelta(days=3))
    if len(ser):
        daily = TABLES[item.table].get("spc_resample") is None
        cats = [_fmt_cat(t, daily) for t in ser.index]
        lim = item.limits_for(e.product)
        lines = [{"name": item.name + ("" if daily else "(8h 평균)"), "values": ser.to_numpy(float), "color": d.SERIES[0],
                  "marker": daily}]
        for v, nm, col, dash in ((lim.usl, "사내 상한", d.WARN, True), (lim.lsl, "사내 하한", d.WARN, True),
                                 (lim.ks_max, f"{item.ks_label} 상한", d.DANGER, False),
                                 (lim.ks_min, f"{item.ks_label} 하한", d.DANGER, False)):
            if v is not None:
                lines.append({"name": f"{nm} {v:g}", "values": [v] * len(cats), "color": col, "dash": dash, "width": 1.25})
        d.line_chart(s, cats, lines, 0.6, 1.7, 7.9, 5.15, y_title=item.unit, name="Item chart")
    d.box(s, 8.8, 1.75, 3.95, 5.05, d.CARD, name="Quant card")
    nd = q["decimals"]
    facts = [("정량 요약(사실)", {"bold": True, "size": 15}),
             (f"기간: {e.start:%m/%d %H:%M} ~ {e.end:%m/%d %H:%M}", {"bullet": True}),
             (f"구간 {q['n']}점 중 {q['n_exceed']}점 {q.get('limit_word', '기준')} 이탈", {"bullet": True}),
             (f"구간 평균 {q['mean']:,.{nd}f} {q['unit']}", {"bullet": True}),
             (f"기준기간 평균 {q['base_mean']:,.{nd}f} (Δ {q['delta']:+,.{nd}f}, {q['effect']:+.1f}σ)", {"bullet": True}),
             (f"{'최대' if e.direction == 'high' else '최소'}값 {e.worst_value:,.{nd}f} {q['unit']}", {"bullet": True})]
    if e.patterns:
        facts.append((f"동반 SPC 패턴: {', '.join(e.patterns)}", {"bullet": True}))
    facts.append((f"영향: {dx.phenomenon.impact}", {"size": 12, "color": d.MUTED}))
    d.text(s, 9.0, 1.95, 3.6, 4.8, facts, size=13)
    d.footer(s, src + " · 차트: 이벤트 전 7일 ~ 후 3일")

    # ②③ 원인 가설
    s = d.slide()
    d.title(s, "② 원인 가설 우선순위 · ③ 확인 방법", _top_cause_line(dx))
    rows, fills = [], {}
    vcol = _VERDICT_FILL
    for i, r in enumerate(dx.results[:7]):
        ev_txt = next((x.text for x in r.evidences if x.support is not None and x.support >= 0.1), "")
        if not ev_txt:
            ev_txt = "측정 데이터 없음 — 현장 확인" if r.support is None else "연관 지표 변화 없음"
        rows.append({"순위": i + 1, "축": r.hyp.axis, "원인 가설": r.hyp.title, "판정": r.verdict,
                     "핵심 데이터 근거": ev_txt[:60], "확인 방법": (r.hyp.verify[0] if r.hyp.verify else "-")[:38]})
        fills[(i, "판정")] = vcol.get(r.verdict, ("FFFFFF", d.INK))
    d.table(s, pd.DataFrame(rows), 0.6, 1.75, 12.1, [0.55, 0.85, 2.6, 1.35, 4.15, 2.6], font_size=10.5, row_h=0.6,
            cell_fills=fills, name="Hypothesis table")
    d.footer(s, VERDICT_LEGEND_SHORT)

    # ④ 조치
    s = d.slide()
    d.title(s, "④ 단기 조치 vs 근본 대책", "단기 조치로 영향 확산을 막고, 근본 대책으로 재발을 방지합니다")
    for idx, (label, acts) in enumerate((("단기 조치 (수일 이내)", dx.short_term), ("근본 대책 (재발 방지)", dx.root_cause))):
        x = 0.6 + idx * 6.15
        d.box(s, x, 1.75, 5.95, 5.0, d.CARD, name=f"Action card {idx + 1}")
        d.text(s, x + 0.25, 1.92, 5.5, 0.45, label, size=17, bold=True, color=d.ACCENT if idx == 0 else d.INK)
        items = [(f"{a}", {"bullet": True, "size": 13}) for a, _ in acts[:6]] or [("-", {})]
        d.text(s, x + 0.25, 2.45, 5.5, 3.6, items, size=13)
        tags = sorted({t for _, t in acts[:6]})
        d.text(s, x + 0.25, 6.05, 5.5, 0.6, "관련 원인: " + ", ".join(tags)[:110], size=11, color=d.MUTED)
    d.footer(s, src)

    # ⑤ KS 평가
    s = d.slide()
    d.title(s, "⑤ KS 규격·사내 관리기준 대비 평가",
            "KS 부적합 발생 — 출하 판정 필요" if any("부적합" in r["KS 판정"] for r in dx.ks_eval) else "KS 규격 대비 판정 결과")
    if dx.ks_eval:
        kdf = pd.DataFrame(dx.ks_eval)[["항목", "품종", "평가 구간", "평균", "최소~최대", "KS 기준", "KS 판정", "사내 기준", "사내기준 판정"]]
        fills = {}
        for i, r in kdf.iterrows():
            if "부적합" in r["KS 판정"]:
                fills[(i, "KS 판정")] = ("FFCDD2", d.INK)
            elif "적합" in r["KS 판정"]:
                fills[(i, "KS 판정")] = ("C8E6C9", d.INK)
        d.table(s, kdf, 0.6, 1.75, 12.1, [1.7, 0.6, 1.75, 0.8, 1.35, 1.05, 2.45, 1.05, 1.35], font_size=10.5, row_h=0.55,
                cell_fills=fills, name="KS table")
    n_ks_bad = sum("부적합" in r["KS 판정"] for r in dx.ks_eval)
    n_sp_bad = sum(str(r["사내기준 판정"]).startswith("이탈") for r in dx.ks_eval)
    pending = sum("미도래" in r["KS 판정"] for r in dx.ks_eval)
    y0 = 1.75 + 0.55 * (len(dx.ks_eval) + 1) + 0.35
    cards = [("KS 규격", "부적합" if n_ks_bad else "적합", "위험" if n_ks_bad else "양호",
              f"부적합 {n_ks_bad}개 항목" if n_ks_bad else "평가 항목 모두 규격 이내"),
             ("사내 관리기준", f"이탈 {n_sp_bad}건" if n_sp_bad else "적합", "경고" if n_sp_bad else "양호",
              "관리기준 이탈 항목 수" if n_sp_bad else "관리기준 이내"),
             ("후속 판단", "출하 판정" if n_ks_bad else ("결과 대기" if pending else "공정 조치"), None,
              "품질책임자 출하 판정·고객 영향 검토" if n_ks_bad else
              ("미도래 결과(예측) 실측으로 확인" if pending else "원인 조치 후 재발 여부 모니터링"))]
    if y0 + 2.0 < 5.85:
        for k, (label, val, stt, sub) in enumerate(cards):
            d.kpi_card(s, 0.6 + k * 4.1, y0, 3.9, 2.0, val, label, sub, stt, value_size=30)
    d.text(s, 0.6, 5.9, 12.1, 0.9, [("근거: " + (item.ks_ref or item.basis), {"size": 12, "color": d.MUTED}),
                                    ("KS 값은 KS L 5201 요약(한국시멘트협회·e나라표준인증) 기준 — 최신 개정 원문 대조 권장",
                                     {"size": 12, "color": d.MUTED})], size=12)
    d.footer(s, src)
    return d.save()
